import streamlit as st
import streamlit.components.v1 as components
import sqlite3
import hashlib
import io
import copy
import requests
import os
import json
import base64
import html
import csv
import re
import unicodedata
import threading
import math
import time
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ==========================================
# CONFIGURAÇÃO DE LAYOUT E ESTILOS
# ==========================================
st.set_page_config(page_title="GAIA", page_icon="logo.png", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

    :root {
        /* ---- Marca ---- */
        --gaia-primary: #00866B;
        --gaia-primary-dark: #006B56;
        --gaia-primary-darker: #00543F;
        --gaia-primary-soft: #E6F5F0;
        --gaia-primary-tint: #F3FAF8;
        --gaia-accent: #F0A21B;
        --gaia-accent-dark: #B97900;
        --gaia-accent-soft: #FDF1DC;

        /* ---- Superfícies e texto ---- */
        --gaia-bg: #EEF2F7;
        --gaia-surface: #FFFFFF;
        --gaia-surface-2: #F8FAFC;
        --gaia-border: #E3E8EF;
        --gaia-border-strong: #C9D2DE;
        --gaia-text: #101828;
        --gaia-muted: #5A6679;
        --gaia-muted-soft: #8B97A8;

        /* ---- Estados semânticos ---- */
        --gaia-danger: #C8311B;
        --gaia-danger-dark: #A4250F;
        --gaia-danger-soft: #FCEAE6;
        --gaia-success: #16794F;
        --gaia-success-soft: #E5F5EC;
        --gaia-info: #0E7CA8;
        --gaia-info-soft: #E5F4FA;
        --gaia-warning: #B9740B;
        --gaia-warning-soft: #FDF1DC;
        --gaia-purple: #6E4FCB;
        --gaia-purple-soft: #EFEAFB;
        --gaia-neutral: #5A6679;
        --gaia-neutral-soft: #EAEDF2;

        /* ---- Sidebar ---- */
        --gaia-sidebar: #00735C;
        --gaia-sidebar-dark: #013F33;
        --gaia-sidebar-active: rgba(255,255,255,.14);
        --gaia-sidebar-soft: rgba(255,255,255,.10);

        /* ---- Elevação ---- */
        --gaia-shadow-xs: 0 1px 2px rgba(16,24,40,.05);
        --gaia-shadow-soft: 0 1px 2px rgba(16,24,40,.04), 0 2px 8px rgba(16,24,40,.05);
        --gaia-shadow: 0 4px 12px rgba(16,24,40,.07), 0 14px 32px -8px rgba(16,24,40,.13);
        --gaia-shadow-lg: 0 18px 40px -10px rgba(16,24,40,.22);
        --gaia-shadow-press: 0 1px 2px rgba(16,24,40,.10);
        --gaia-shadow-primary: 0 8px 18px -4px rgba(0,134,107,.34);
        --gaia-shadow-danger: 0 8px 18px -4px rgba(200,49,27,.30);

        /* ---- Forma e ritmo ---- */
        --gaia-radius-xl: 20px;
        --gaia-radius-lg: 16px;
        --gaia-radius-md: 12px;
        --gaia-radius-sm: 9px;
        --gaia-radius-pill: 999px;
        --gaia-control-h: 44px;
        --gaia-ease: 160ms cubic-bezier(.22,.75,.32,1);
        --gaia-ease-spring: 240ms cubic-bezier(.34,1.35,.4,1);
        --gaia-focus-ring: 0 0 0 3px rgba(0,134,107,.20), 0 0 0 1px rgba(0,134,107,.55);
    }

    /* ==========================================
       BASE
       ========================================== */
    header[data-testid="stHeader"] { display:none !important; }
    .stApp {
        background:
            radial-gradient(1100px 480px at 100% -8%, rgba(0,134,107,.05), transparent 60%),
            radial-gradient(900px 420px at -8% 8%, rgba(240,162,27,.05), transparent 55%),
            var(--gaia-bg) !important;
        color: var(--gaia-text) !important;
    }
    .block-container {
        max-width: 1440px !important;
        padding-top: 1.1rem !important;
        padding-bottom: 2.4rem !important;
        padding-left: 1.6rem !important;
        padding-right: 1.6rem !important;
    }

    html, body, h1, h2, h3, h4, h5, h6, p, label, span, div, button, input, textarea, select {
        font-family: "Inter", "Segoe UI", Arial, sans-serif !important;
        -webkit-font-smoothing: antialiased;
    }

    /* Corrige o bug em que o Streamlit exibe "keyboard_arrow_right" como texto
       nos expanders. O seletor global acima força Inter em todos os spans; aqui
       devolvemos a fonte correta dos ícones nativos do Streamlit. */
    span[data-testid="stIconMaterial"],
    span.material-symbols-rounded,
    span.material-symbols-outlined,
    span.material-icons,
    span[class*="material-symbols"] {
        font-family: "Material Symbols Rounded", "Material Symbols Outlined", "Material Icons" !important;
        font-weight: normal !important;
        font-style: normal !important;
        font-size: 1.15rem !important;
        line-height: 1 !important;
        letter-spacing: normal !important;
        text-transform: none !important;
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        direction: ltr !important;
        -webkit-font-feature-settings: "liga" !important;
        -webkit-font-smoothing: antialiased !important;
        font-feature-settings: "liga" !important;
    }
    h1, h2, h3, h4 { color: var(--gaia-text) !important; letter-spacing: -.015em; font-weight: 700 !important; }
    p, span, label, div { letter-spacing: -.001em; }
    hr { border-color: var(--gaia-border) !important; margin: 1.1rem 0 !important; }
    a { color: var(--gaia-primary-dark); }

    /* Cabeçalhos nativos (st.subheader) usados como título de tela */
    h3 {
        font-size: 1.22rem !important;
        font-weight: 700 !important;
        padding-bottom: .55rem;
        margin-bottom: .35rem !important;
        border-bottom: 1px solid var(--gaia-border);
    }

    /* Texto em negrito simples usado como pequeno rótulo de seção dentro de formulários */
    .stMarkdown strong { color: var(--gaia-text); }

    /* st.caption */
    [data-testid="stCaptionContainer"], .stCaption {
        color: var(--gaia-muted) !important;
        font-size: .83rem !important;
    }

    /* Barra de rolagem discreta, alinhada à identidade visual */
    ::-webkit-scrollbar { width: 10px; height: 10px; }
    ::-webkit-scrollbar-track { background: transparent; }
    ::-webkit-scrollbar-thumb { background: var(--gaia-border-strong); border-radius: 999px; }
    ::-webkit-scrollbar-thumb:hover { background: var(--gaia-muted-soft); }

    /* Estado de foco visível e consistente para acessibilidade de teclado */
    a:focus-visible, button:focus-visible,
    div[data-baseweb="select"]:focus-within,
    div[data-baseweb="input"]:focus-within {
        outline: none !important;
        box-shadow: var(--gaia-focus-ring) !important;
        border-radius: var(--gaia-radius-sm);
    }

    /* ==========================================
       BARRA SUPERIOR
       ========================================== */
    .gaia-topbar {
        background: var(--gaia-surface);
        border: 1px solid var(--gaia-border);
        border-radius: var(--gaia-radius-lg);
        box-shadow: var(--gaia-shadow-soft);
        padding: 18px 28px;
        min-height: 96px;
        display:flex;
        align-items:center;
        justify-content:space-between;
        gap: 16px;
        margin-bottom: 1.3rem;
        position: relative;
        overflow: hidden;
    }
    .gaia-topbar::before {
        content: "";
        position: absolute; left:0; top:0; bottom:0; width:4px;
        background: linear-gradient(180deg, var(--gaia-primary), var(--gaia-accent));
    }
    .gaia-brand-row { display:flex; align-items:center; gap:12px; min-width:0; padding-left:6px; }
    .gaia-mark {
        width: 44px; height: 44px; border-radius: 13px;
        display:flex; align-items:center; justify-content:center;
        background: linear-gradient(135deg, var(--gaia-primary), var(--gaia-primary-dark));
        color: #fff; font-weight: 800; letter-spacing:.03em;
        box-shadow: var(--gaia-shadow-primary);
        flex: 0 0 auto;
    }
    .gaia-brand-title { font-size: 1.05rem; font-weight: 800; color: var(--gaia-text); line-height: 1.05; }
    .gaia-brand-subtitle { font-size: .79rem; color: var(--gaia-muted); margin-top: 3px; }
    .gaia-top-logo-img {
        width: 280px; max-width: 38vw; height: 64px;
        object-fit: contain; object-position: left center; display:block;
    }
    .gaia-wordmark-fallback { font-size:1.28rem; font-weight:800; color:var(--gaia-primary-dark); letter-spacing:.01em; }
    .gaia-top-wordmark {
        display:flex; align-items:center; gap:10px;
    }
    .gaia-top-wordmark-symbol {
        width:42px; height:42px; border-radius:12px; display:flex; align-items:center; justify-content:center;
        background:linear-gradient(135deg, var(--gaia-primary), var(--gaia-primary-dark));
        color:#fff; font-weight:900; box-shadow:var(--gaia-shadow-primary);
    }
    .gaia-top-wordmark-title {
        font-size:1.28rem; font-weight:900; color:var(--gaia-primary-dark); line-height:1; letter-spacing:.02em;
    }
    .gaia-top-wordmark-sub {
        font-size:.72rem; color:var(--gaia-muted); margin-top:4px; font-weight:600;
    }
    .gaia-user-pill {
        display:inline-flex; align-items:center; gap:9px;
        border:1px solid #BFE4D9;
        background: linear-gradient(180deg, var(--gaia-primary-tint), var(--gaia-primary-soft));
        color: var(--gaia-primary-darker);
        padding: 7px 16px 7px 8px;
        border-radius: var(--gaia-radius-pill);
        font-size: .81rem;
        font-weight: 700;
        white-space: nowrap;
        flex: 0 0 auto;
        box-shadow: var(--gaia-shadow-xs);
    }
    .gaia-user-pill .gaia-user-avatar {
        width: 26px; height: 26px; border-radius: 50%;
        background: linear-gradient(135deg, var(--gaia-primary), var(--gaia-primary-dark));
        color: #fff; font-size: .76rem; font-weight: 800;
        display:flex; align-items:center; justify-content:center;
        flex: 0 0 auto;
    }

    /* ==========================================
       LOGIN — tela enxuta, sem hero/landing
       ========================================== */
    .gaia-login-page { max-width: 920px; margin: 1.4rem auto 0 auto; }
    .gaia-logo-img { display:block; width: 260px; max-width:100%; margin: 0 auto 6px auto; height:auto; object-fit:contain; }
    .gaia-logo-fallback { text-align:center; color:var(--gaia-primary-dark); font-size:2.05rem; font-weight:800; margin-bottom: 8px; }
    .gaia-login-heading { text-align:center; color: var(--gaia-text); font-size: 1.5rem; font-weight: 700; margin: 10px 0 4px 0; }
    .gaia-login-caption { text-align:center; color: var(--gaia-muted); font-size: .92rem; margin-bottom: 26px; }
    .gaia-login-card, div[data-testid="stForm"] {
        background: var(--gaia-surface) !important;
        border: 1px solid var(--gaia-border) !important;
        border-radius: var(--gaia-radius-lg) !important;
        box-shadow: var(--gaia-shadow) !important;
        padding: 28px 28px 22px 28px !important;
        position: relative;
    }
    div[data-testid="stForm"]::before {
        content: "";
        position: absolute; left:0; right:0; top:0; height:4px;
        border-radius: var(--gaia-radius-lg) var(--gaia-radius-lg) 0 0;
        background: linear-gradient(90deg, var(--gaia-primary), var(--gaia-accent));
        opacity: .9;
    }
    .gaia-help-text { color: var(--gaia-muted); font-size: .81rem; margin: 8px 0 12px 0; }

    /* ==========================================
       CARDS E COMPONENTES GERAIS
       ========================================== */
    .gaia-page-card {
        background: var(--gaia-surface);
        border: 1px solid var(--gaia-border);
        border-radius: var(--gaia-radius-lg);
        box-shadow: var(--gaia-shadow-soft);
        padding: 22px 24px;
        margin-bottom: 18px;
    }
    .gaia-work-title { font-size: 1.46rem; font-weight: 800; color: var(--gaia-text); margin-bottom: 4px; letter-spacing: -.02em; }
    .gaia-work-subtitle { color: var(--gaia-muted); font-size: .92rem; line-height: 1.55; }

    .gaia-section-heading, .gaia-section-title {
        font-size: 1.0rem; font-weight: 700; color: var(--gaia-text);
        margin: 22px 0 13px 0;
        padding-left: 12px;
        border-left: 3px solid var(--gaia-primary);
        letter-spacing: -.01em;
    }
    .gaia-kicker {
        display:inline-flex; align-items:center; gap:6px;
        border-radius: var(--gaia-radius-pill); padding: 5px 13px;
        background: var(--gaia-primary-soft); color: var(--gaia-primary-darker);
        border: 1px solid #C7E9DE;
        font-size: .68rem; font-weight: 800;
        letter-spacing: .09em; text-transform: uppercase;
        margin-bottom: 10px;
    }

    .gaia-tool-card {
        background: var(--gaia-surface);
        border: 1px solid var(--gaia-border);
        border-radius: var(--gaia-radius-lg);
        min-height: 128px;
        padding: 18px 18px 15px 18px;
        box-shadow: var(--gaia-shadow-soft);
        display:flex;
        flex-direction:column;
        justify-content:space-between;
        margin-bottom: 10px;
        transition: transform var(--gaia-ease-spring), box-shadow var(--gaia-ease);
    }
    .gaia-tool-card:hover { transform: translateY(-3px); box-shadow: var(--gaia-shadow); }
    .gaia-tool-card-main { border-top: 4px solid var(--gaia-primary); }
    .gaia-tool-card-accent { border-top: 4px solid var(--gaia-accent); }
    .gaia-tool-card-neutral { border-top: 4px solid #90A4B8; }
    .gaia-tool-icon {
        width: 46px; height: 46px; border-radius: 13px;
        display:flex; align-items:center; justify-content:center;
        margin-bottom: 10px; font-size: 1.35rem; font-weight: 700;
    }
    .gaia-icon-main { background: var(--gaia-primary-soft); color: var(--gaia-primary-dark); }
    .gaia-icon-accent { background: var(--gaia-accent-soft); color: #8A5B00; }
    .gaia-icon-neutral { background: #EFF3F7; color: #516173; }
    .gaia-tool-title { font-weight: 700; color: var(--gaia-text); font-size: .98rem; margin-bottom: 4px; }
    .gaia-tool-desc { color: var(--gaia-muted); font-size: .82rem; line-height: 1.4; }

    .gaia-inline-help {
        background: var(--gaia-accent-soft);
        border: 1px solid #F0D08D;
        color: #745009;
        border-radius: var(--gaia-radius-sm);
        padding: 11px 13px;
        font-size: .82rem;
        line-height: 1.45;
        margin: 6px 0 12px 0;
        display:flex; gap:8px; align-items:flex-start;
    }
    .gaia-inline-help::before { content: "💡"; flex:0 0 auto; }
    .gaia-inline-note {
        display:inline-flex; align-items:center; gap:7px;
        background: var(--gaia-info-soft);
        border: 1px solid #BCE0EE;
        color: #0a5068;
        border-radius: var(--gaia-radius-sm);
        padding: 10px 13px;
        font-size: .82rem;
        line-height: 1.45;
    }
    .gaia-prev-value {
        background: var(--gaia-surface-2); border:1px solid var(--gaia-border); border-left:4px solid var(--gaia-primary);
        border-radius: var(--gaia-radius-sm); padding:.64rem .85rem; margin:.3rem 0 .7rem 0; color:var(--gaia-text); font-size:.92rem;
    }
    .gaia-company-name { color:var(--gaia-primary-dark); font-size:1.18rem; font-weight:700; }
    .gaia-meta, .gaia-subtitle { color: var(--gaia-muted); font-size:.88rem; line-height:1.5; }

    .gaia-card, .gaia-card-compact {
        background: var(--gaia-surface);
        border: 1px solid var(--gaia-border);
        border-radius: var(--gaia-radius-md);
        box-shadow: var(--gaia-shadow-soft);
        padding: 17px 19px;
        margin: 0 0 13px 0;
        transition: box-shadow var(--gaia-ease), border-color var(--gaia-ease);
    }
    .gaia-card-compact { border-left: 4px solid var(--gaia-primary); }
    .gaia-card-compact:hover, .gaia-card:hover { box-shadow: var(--gaia-shadow-soft), 0 0 0 1px var(--gaia-border-strong) inset; }
    .gaia-page-title { color:var(--gaia-text); font-size:1.42rem; font-weight:800; margin: 4px 0 2px 0; letter-spacing: -.02em; }

    /* Listas de registros (usuários, sócios, documentos) */
    .gaia-socio-name { color: var(--gaia-text); font-size: 1rem; font-weight: 700; margin: 0 0 2px 0; }
    .gaia-socio-meta { color: var(--gaia-muted); font-size: .85rem; line-height: 1.5; }
    .gaia-separator { height: 1px; background: var(--gaia-border); margin: 14px 0; border: none; }
    .gaia-footer-actions { height: 6px; margin-top: 10px; border-top: 1px dashed var(--gaia-border); }

    .gaia-bottom-section { text-align:center; }
    .gaia-bottom-title { font-size: 1.14rem; font-weight: 800; color: var(--gaia-text); }
    @keyframes gaia-doc-pulse {
        0%   { box-shadow: 0 0 0 0 rgba(0,134,107,.35), var(--gaia-shadow-soft); }
        60%  { box-shadow: 0 0 0 10px rgba(0,134,107,0), var(--gaia-shadow-soft); }
        100% { box-shadow: 0 0 0 0 rgba(0,134,107,0), var(--gaia-shadow-soft); }
    }
    #documento-gerado {
        background: linear-gradient(135deg, #eaf6f2 0%, var(--gaia-surface) 100%);
        border: 1.5px solid rgba(0,134,107,.30);
        border-radius: var(--gaia-radius-lg);
        padding: 22px 24px !important;
        animation: gaia-doc-pulse 1.4s ease 0.2s 3;
    }

    .gaia-workbench-card {
        padding: 34px 32px !important;
        border-left: 5px solid var(--gaia-primary);
        background: linear-gradient(135deg, var(--gaia-surface) 60%, #eaf6f2 100%) !important;
    }
    .gaia-work-title { font-size: 1.72rem !important; font-weight: 800; color: var(--gaia-text); margin-bottom: 6px; letter-spacing: -.025em; }
    .gaia-work-subtitle { color: var(--gaia-muted); font-size: .96rem; line-height: 1.6; }

    /* Selo de status (histórico de documentos) */
    .gaia-status-pill {
        display:inline-flex; align-items:center; gap:6px;
        padding: 5px 13px 5px 10px; border-radius: var(--gaia-radius-pill);
        font-weight: 700; font-size: .74rem; letter-spacing:.02em;
        box-shadow: var(--gaia-shadow-xs);
        white-space: nowrap;
    }
    .gaia-status-pill::before { content:""; width:6px; height:6px; border-radius:50%; background:currentColor; flex:0 0 auto; }

    /* ==========================================
       CAMPOS DE FORMULÁRIO
       ========================================== */
    div[data-baseweb="input"],
    div[data-baseweb="textarea"],
    div[data-baseweb="select"] > div {
        background: #FFFFFF !important;
        border: 1.4px solid var(--gaia-border-strong) !important;
        border-radius: var(--gaia-radius-sm) !important;
        transition: border-color var(--gaia-ease), box-shadow var(--gaia-ease);
    }
    div[data-baseweb="input"]:focus-within,
    div[data-baseweb="textarea"]:focus-within {
        border-color: var(--gaia-primary) !important;
        box-shadow: var(--gaia-focus-ring) !important;
    }
    /* O <input> em si NÃO recebe borda/raio próprios — só o container acima
       (data-baseweb="input") desenha a borda. Ter borda nos dois níveis ao
       mesmo tempo é o que produzia a faixa retangular sobrando ao lado do
       ícone de revelar senha, pois o <input> não preenchia 100% do espaço
       interno e a borda "extra" dele ficava visível como um retalho à parte. */
    div[data-baseweb="input"] input,
    textarea,
    div[data-baseweb="select"] {
        background: transparent !important;
        border: none !important;
        min-height: var(--gaia-control-h) !important;
        color: var(--gaia-text) !important;
        font-size: .93rem !important;
    }
    textarea { padding-top: 10px !important; border: 1.4px solid var(--gaia-border-strong) !important; border-radius: var(--gaia-radius-sm) !important; background: #FFFFFF !important; }
    div[data-baseweb="input"] input::placeholder, textarea::placeholder { color: var(--gaia-muted-soft) !important; }
    div[data-baseweb="select"]:hover { border-color: var(--gaia-border-strong); }

    /* Rótulos dos campos */
    div[data-testid="stWidgetLabel"] label, label[data-testid="stWidgetLabel"] {
        font-size: .85rem !important;
        font-weight: 600 !important;
        color: var(--gaia-text) !important;
        margin-bottom: 4px !important;
    }

    /* Selectbox / multiselect — menu suspenso e marcações (tags) */
    div[data-baseweb="popover"] li[role="option"]:hover,
    div[data-baseweb="menu"] li:hover {
        background: var(--gaia-primary-tint) !important;
    }
    div[data-baseweb="select"] [data-baseweb="tag"] {
        background: var(--gaia-primary-soft) !important;
        border: 1px solid #C7E9DE !important;
        color: var(--gaia-primary-darker) !important;
        border-radius: var(--gaia-radius-sm) !important;
        font-weight: 600 !important;
    }

    /* Checkbox — usados na seleção de eventos (lista de opções) */
    div[data-testid="stCheckbox"] {
        background: var(--gaia-surface);
        border: 1px solid var(--gaia-border);
        border-radius: var(--gaia-radius-sm);
        /* padding removido do div e colocado no label para que clicar
           em qualquer ponto do card propague o clique ao checkbox */
        padding: 0;
        margin-bottom: 8px;
        transition: border-color var(--gaia-ease), background var(--gaia-ease), box-shadow var(--gaia-ease);
        cursor: pointer;
    }
    div[data-testid="stCheckbox"]:hover {
        border-color: var(--gaia-border-strong);
        background: var(--gaia-surface-2);
        box-shadow: var(--gaia-shadow-xs);
    }
    /* Label preenche todo o card — qualquer clique na área registra no input */
    div[data-testid="stCheckbox"] label {
        font-size: .92rem !important;
        font-weight: 600 !important;
        display: block !important;
        width: 100% !important;
        padding: 11px 14px !important;
        box-sizing: border-box !important;
        cursor: pointer !important;
        margin: 0 !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div:first-child {
        border-radius: 6px !important;
        border-width: 1.6px !important;
        width: 20px !important; height: 20px !important;
    }
    /* Checkbox marcado: destaque discreto no card e cor apenas no quadradinho,
       sem pintar o texto inteiro da opção. */
    div[data-testid="stCheckbox"]:has(input:checked) {
        background: rgba(0,134,107,.035) !important;
        border-color: rgba(0,134,107,.28) !important;
        box-shadow: 0 0 0 1px rgba(0,134,107,.06) inset !important;
    }
    div[data-testid="stCheckbox"]:has(input:checked) span[data-baseweb="checkbox"] > div:first-child {
        background-color: var(--gaia-primary) !important;
        border-color: var(--gaia-primary) !important;
    }
    div[data-testid="stCheckbox"]:has(input:checked) span[data-baseweb="checkbox"] > div:first-child svg {
        color: #ffffff !important;
        fill: #ffffff !important;
        stroke: #ffffff !important;
        opacity: 1 !important;
        visibility: visible !important;
    }
    div[data-testid="stCheckbox"]:has(input:checked) label,
    div[data-testid="stCheckbox"]:has(input:checked) p,
    div[data-testid="stCheckbox"]:has(input:checked) span {
        background: transparent !important;
        color: var(--gaia-text) !important;
    }

    /* Date / number input steppers */
    div[data-testid="stDateInput"] input, div[data-testid="stNumberInput"] input { font-weight: 600 !important; }

    /* File uploader */
    section[data-testid="stFileUploaderDropzone"], div[data-testid="stFileUploaderDropzone"] {
        background: var(--gaia-surface-2) !important;
        border: 1.6px dashed var(--gaia-border-strong) !important;
        border-radius: var(--gaia-radius-md) !important;
    }

    /* st.expander */
    details[data-testid="stExpander"], div[data-testid="stExpander"] {
        background: var(--gaia-surface) !important;
        border: 1px solid var(--gaia-border) !important;
        border-radius: var(--gaia-radius-md) !important;
        box-shadow: var(--gaia-shadow-xs) !important;
        overflow: hidden;
    }
    details[data-testid="stExpander"] summary, div[data-testid="stExpander"] summary {
        font-weight: 700 !important;
        font-size: .9rem !important;
        padding: 10px 14px !important;
        background: var(--gaia-surface) !important;
        color: var(--gaia-text) !important;
        border: 0 !important;
        box-shadow: none !important;
        transition: background var(--gaia-ease), color var(--gaia-ease);
    }
    details[data-testid="stExpander"] summary:hover, div[data-testid="stExpander"] summary:hover {
        background: var(--gaia-surface-2) !important;
        color: var(--gaia-text) !important;
    }
    details[data-testid="stExpander"] > div, div[data-testid="stExpander"] > div {
        background: var(--gaia-surface) !important;
    }

    /* Proteção extra para o ícone do expander: evita aparecer o texto interno
       "keyboard_arrow_right" quando a fonte de ícones não renderiza. */
    details[data-testid="stExpander"] summary span[data-testid="stIconMaterial"],
    div[data-testid="stExpander"] summary span[data-testid="stIconMaterial"],
    details[data-testid="stExpander"] summary span.material-symbols-rounded,
    div[data-testid="stExpander"] summary span.material-symbols-rounded {
        font-size: 0 !important;
        width: 18px !important;
        height: 18px !important;
        min-width: 18px !important;
        overflow: hidden !important;
        margin-right: 6px !important;
    }
    details[data-testid="stExpander"] summary span[data-testid="stIconMaterial"]::after,
    div[data-testid="stExpander"] summary span[data-testid="stIconMaterial"]::after,
    details[data-testid="stExpander"] summary span.material-symbols-rounded::after,
    div[data-testid="stExpander"] summary span.material-symbols-rounded::after {
        content: "▸";
        font-family: "Inter", "Segoe UI", Arial, sans-serif !important;
        font-size: 13px !important;
        line-height: 18px !important;
        color: var(--gaia-muted) !important;
    }
    details[data-testid="stExpander"][open] summary span[data-testid="stIconMaterial"]::after,
    details[data-testid="stExpander"][open] summary span.material-symbols-rounded::after {
        content: "▾";
    }

    /* Alertas nativos do Streamlit (st.info / st.warning / st.success / st.error) */
    div[data-testid="stAlertContainer"] {
        border-radius: var(--gaia-radius-md) !important;
        border: 1px solid transparent !important;
        box-shadow: var(--gaia-shadow-xs) !important;
        padding: 13px 16px !important;
        font-size: .89rem !important;
    }
    div[data-testid="stAlertContainer"][data-baseweb="notification-positive"],
    div[data-testid="stAlertContainer"].st-emotion-cache-1lz0u4n,
    div[data-testid="stAlertContainer"]:has(svg[data-testid="stIconMaterial"]) {
        border-color: var(--gaia-border) !important;
    }
    div[kind="info"] { background: var(--gaia-info-soft) !important; border-color: #BCE0EE !important; color: #0a5068 !important; }
    div[kind="warning"] { background: var(--gaia-warning-soft) !important; border-color: #F0D08D !important; color: #6b4a05 !important; }
    div[kind="success"] { background: var(--gaia-success-soft) !important; border-color: #BFE3CC !important; color: #0f5132 !important; }
    div[kind="error"] { background: var(--gaia-danger-soft) !important; border-color: #F3C2BA !important; color: #7a1d0f !important; }

    /* ==========================================
       BOTÕES — hierarquia clara: primário, secundário e perigo
       ========================================== */
    div[data-testid="stButton"] button,
    div[data-testid="stFormSubmitButton"] button,
    div[data-testid="stDownloadButton"] button,
    .stButton button, .stFormSubmitButton button, .stDownloadButton button,
    button[data-testid^="stBaseButton"] {
        border-radius: var(--gaia-radius-sm) !important;
        min-height: var(--gaia-control-h) !important;
        font-weight: 700 !important;
        font-size: .92rem !important;
        letter-spacing: -.005em;
        box-shadow: none !important;
        transition: background-color var(--gaia-ease), border-color var(--gaia-ease),
                    transform var(--gaia-ease), box-shadow var(--gaia-ease) !important;
        background: #FFFFFF !important;
        background-color: #FFFFFF !important;
        border: 1.4px solid var(--gaia-border-strong) !important;
        color: var(--gaia-text) !important;
    }
    div[data-testid="stButton"] button:hover,
    div[data-testid="stFormSubmitButton"] button:hover,
    div[data-testid="stDownloadButton"] button:hover,
    button[data-testid^="stBaseButton"]:hover {
        background: var(--gaia-surface-2) !important;
        background-color: var(--gaia-surface-2) !important;
        border-color: #94A6BA !important;
        color: var(--gaia-text) !important;
        transform: translateY(-1px);
        box-shadow: var(--gaia-shadow-soft) !important;
    }
    div[data-testid="stButton"] button:active,
    div[data-testid="stFormSubmitButton"] button:active,
    button[data-testid^="stBaseButton"]:active {
        transform: translateY(0px);
        box-shadow: var(--gaia-shadow-press) !important;
    }
    div[data-testid="stButton"] button:disabled,
    button[data-testid^="stBaseButton"]:disabled {
        opacity: .5 !important;
        cursor: not-allowed !important;
        transform: none !important;
        box-shadow: none !important;
    }

    /* Botões primários — verde translúcido para ações de seguir, salvar, login e avançar. */
    button[data-testid="stBaseButton-primary"],
    button[data-testid="stBaseButton-primaryFormSubmit"] {
        background: rgba(0, 134, 107, .12) !important;
        background-color: rgba(0, 134, 107, .12) !important;
        border: 1.4px solid rgba(0, 107, 86, .34) !important;
        color: var(--gaia-primary-darker) !important;
        box-shadow: none !important;
    }
    button[data-testid="stBaseButton-primary"]:hover,
    button[data-testid="stBaseButton-primaryFormSubmit"]:hover {
        background: rgba(0, 134, 107, .18) !important;
        background-color: rgba(0, 134, 107, .18) !important;
        border-color: rgba(0, 107, 86, .52) !important;
        color: var(--gaia-primary-darker) !important;
        transform: translateY(-1px);
        box-shadow: 0 8px 16px -10px rgba(0,107,86,.45) !important;
    }
    button[data-testid="stBaseButton-primary"]:active,
    button[data-testid="stBaseButton-primaryFormSubmit"]:active {
        background: rgba(0, 134, 107, .24) !important;
        transform: translateY(0px);
    }

    /* Botões secundários: contorno neutro, fundo branco (definidos pela regra base acima) */
    button[data-testid="stBaseButton-secondary"]:hover,
    button[data-testid="stBaseButton-secondaryFormSubmit"]:hover {
        color: var(--gaia-primary-dark) !important;
        border-color: var(--gaia-primary) !important;
        background: var(--gaia-primary-tint) !important;
    }

    /* Ação de destaque "Cadastrar empresa" no menu inicial recebe tom âmbar */
    div.st-key-btn_nova_empresa_home button {
        background: var(--gaia-accent-soft) !important;
        background-color: var(--gaia-accent-soft) !important;
        border: 1.4px solid #F0C36C !important;
        color: #6F4A00 !important;
    }
    div.st-key-btn_nova_empresa_home button:hover {
        background:#FCE7BC !important;
        border-color: var(--gaia-accent-dark) !important;
        color:#5F3E00 !important;
    }

    /* Ações da tela inicial como opções clicáveis (cartão + botão em um só elemento) */
    div.st-key-card_continuar_alt button,
    div.st-key-card_nova_empresa button,
    div.st-key-card_historico_docs button,
    div.st-key-empresa_card_click button {
        min-height: 156px !important;
        border-radius: var(--gaia-radius-lg) !important;
        background: #FFFFFF !important;
        background-color: #FFFFFF !important;
        color: var(--gaia-text) !important;
        justify-content: flex-start !important;
        text-align: left !important;
        align-items: flex-start !important;
        padding: 26px 24px !important;
        white-space: pre-wrap !important;
        line-height: 1.45 !important;
        font-size: .92rem !important;
        font-weight: 600 !important;
        box-shadow: var(--gaia-shadow-soft) !important;
        position: relative !important;
        overflow: hidden;
    }
    div.st-key-card_continuar_alt button::first-line,
    div.st-key-card_nova_empresa button::first-line,
    div.st-key-card_historico_docs button::first-line,
    div.st-key-empresa_card_click button::first-line {
        font-size: 2.4rem !important;
        line-height: 2.2 !important;
    }
    div.st-key-card_continuar_alt button,
    div.st-key-card_nova_empresa button,
    div.st-key-card_historico_docs button {
        min-height: 148px !important;
    }
    div.st-key-empresa_card_click button { min-height: 112px !important; padding: 20px 22px !important; }
    div.st-key-card_continuar_alt button { border: 1px solid #BFE6D8 !important; border-top: 5px solid var(--gaia-primary) !important; }
    div.st-key-card_nova_empresa button { border: 1px solid #F1D28A !important; border-top: 5px solid var(--gaia-accent) !important; }
    div.st-key-card_historico_docs button { border: 1px solid #CAD7E5 !important; border-top: 5px solid #8EA3B7 !important; }
    div.st-key-empresa_card_click button { border: 1px solid var(--gaia-border) !important; border-left: 5px solid var(--gaia-primary) !important; }
    div.st-key-card_continuar_alt button:hover,
    div.st-key-card_nova_empresa button:hover,
    div.st-key-card_historico_docs button:hover,
    div.st-key-empresa_card_click button:hover {
        transform: translateY(-4px);
        box-shadow: var(--gaia-shadow-lg) !important;
        background: var(--gaia-surface-2) !important;
        background-color: var(--gaia-surface-2) !important;
    }
    div.st-key-card_continuar_alt button:hover { border-color: var(--gaia-primary) !important; }
    div.st-key-card_nova_empresa button:hover { border-color: var(--gaia-accent) !important; }
    div.st-key-card_historico_docs button:hover { border-color: #8EA3B7 !important; }
    div.st-key-empresa_card_click button:hover { border-color: var(--gaia-primary) !important; }



    /* ==========================================
       CORES DE BOTÕES POR FUNÇÃO — v8.21
       ========================================== */
    /* Regra visual: cores translúcidas e sem brigar com a leitura da tela.
       Verde = seguir/continuar/login/salvar. Amarelo = editar/ajustar. Vermelho = excluir.
       Azul/roxo = consulta, histórico e downloads. */

    /* Cards do painel: são opções clicáveis. O fluxo principal fica verde translúcido. */
    div.st-key-card_continuar_alt button {
        background: linear-gradient(180deg, rgba(0,134,107,.12) 0%, #FFFFFF 74%) !important;
        border-color: rgba(0,107,86,.34) !important;
        border-top-color: rgba(0,107,86,.70) !important;
        color: var(--gaia-text) !important;
    }
    div.st-key-card_continuar_alt button:hover {
        background: linear-gradient(180deg, rgba(0,134,107,.18) 0%, #FFFFFF 70%) !important;
        border-color: rgba(0,107,86,.52) !important;
        box-shadow: var(--gaia-shadow-lg) !important;
    }

    div.st-key-card_nova_empresa button {
        background: linear-gradient(180deg, rgba(240,162,27,.13) 0%, #FFFFFF 74%) !important;
        border-color: rgba(240,162,27,.38) !important;
        border-top-color: rgba(240,162,27,.78) !important;
        color: var(--gaia-text) !important;
    }
    div.st-key-card_nova_empresa button:hover {
        background: linear-gradient(180deg, rgba(240,162,27,.20) 0%, #FFFFFF 70%) !important;
        border-color: rgba(185,121,0,.54) !important;
        box-shadow: var(--gaia-shadow-lg) !important;
    }

    div.st-key-card_historico_docs button {
        background: linear-gradient(180deg, rgba(14,124,168,.10) 0%, #FFFFFF 74%) !important;
        border-color: rgba(14,124,168,.26) !important;
        border-top-color: rgba(14,124,168,.58) !important;
        color: var(--gaia-text) !important;
    }
    div.st-key-card_historico_docs button:hover {
        background: linear-gradient(180deg, rgba(14,124,168,.16) 0%, #FFFFFF 70%) !important;
        border-color: rgba(14,124,168,.44) !important;
        box-shadow: var(--gaia-shadow-lg) !important;
    }

    /* Editar / senha / redefinir — amarelo translúcido. */
    div[class*="st-key-ed_s_"] button,
    div[class*="st-key-senha_self_"] button,
    div[class*="st-key-reset_senha_"] button,
    div.st-key-btn_senha_home button,
    div.st-key-side_menu_senha button {
        background: rgba(240,162,27,.12) !important;
        background-color: rgba(240,162,27,.12) !important;
        border: 1.4px solid rgba(185,121,0,.35) !important;
        color: #6F4A00 !important;
        box-shadow: none !important;
    }
    div[class*="st-key-ed_s_"] button:hover,
    div[class*="st-key-senha_self_"] button:hover,
    div[class*="st-key-reset_senha_"] button:hover,
    div.st-key-btn_senha_home button:hover,
    div.st-key-side_menu_senha button:hover {
        background: rgba(240,162,27,.20) !important;
        border-color: rgba(185,121,0,.55) !important;
        color: #5F3E00 !important;
    }

    /* Excluir (gatilho) — vermelho quase transparente. Apenas os botões que
       DISPARAM a confirmação, não os que CONFIRMAM a ação. */
    div[class*="st-key-del_user_"] button,
    div[class*="st-key-ex_s_"] button,
    div.st-key-ex_emp button,
    div[class*="st-key-hist_geral_del_doc_"] button,
    div[class*="st-key-hist_empresa_del_doc_"] button,
    div[class*="st-key-limpar_comp_"] button {
        background: rgba(200,49,27,.09) !important;
        background-color: rgba(200,49,27,.09) !important;
        border: 1.4px solid rgba(164,37,15,.25) !important;
        color: var(--gaia-danger-dark) !important;
        box-shadow: none !important;
    }
    div[class*="st-key-del_user_"] button:hover,
    div[class*="st-key-ex_s_"] button:hover,
    div.st-key-ex_emp button:hover,
    div[class*="st-key-hist_geral_del_doc_"] button:hover,
    div[class*="st-key-hist_empresa_del_doc_"] button:hover,
    div[class*="st-key-limpar_comp_"] button:hover {
        background: rgba(200,49,27,.16) !important;
        background-color: rgba(200,49,27,.16) !important;
        border-color: rgba(164,37,15,.44) !important;
        color: #7A1D0F !important;
        transform: translateY(-1px);
        box-shadow: 0 8px 16px -12px rgba(164,37,15,.35) !important;
    }

    /* ==========================================
       BOTÕES DE CONFIRMAR — verde translúcido (idêntico ao login)
       Inclui confirmações de exclusão: verde = "sim, prosseguir".
       ========================================== */
    div.st-key-confirm_del_user button,
    div.st-key-confirm_del_empresa button,
    div[class*="st-key-confirm_del_socio_"] button,
    div[class*="st-key-hist_geral_confirm_del_doc_"] button,
    div[class*="st-key-hist_empresa_confirm_del_doc_"] button {
        background: rgba(0,134,107,.12) !important;
        background-color: rgba(0,134,107,.12) !important;
        border: 1.4px solid rgba(0,107,86,.34) !important;
        color: var(--gaia-primary-darker) !important;
        box-shadow: none !important;
    }
    div.st-key-confirm_del_user button:hover,
    div.st-key-confirm_del_empresa button:hover,
    div[class*="st-key-confirm_del_socio_"] button:hover,
    div[class*="st-key-hist_geral_confirm_del_doc_"] button:hover,
    div[class*="st-key-hist_empresa_confirm_del_doc_"] button:hover {
        background: rgba(0,134,107,.19) !important;
        background-color: rgba(0,134,107,.19) !important;
        border-color: rgba(0,107,86,.52) !important;
        color: var(--gaia-primary-darker) !important;
        transform: translateY(-1px);
        box-shadow: 0 8px 16px -10px rgba(0,107,86,.40) !important;
    }

    /* ==========================================
       BOTÕES DE CANCELAR — vermelho translúcido
       Cobre: cancel_* (keys explícitas) e fsb_cancelar_* (submit de forms).
       ========================================== */
    div[class*="st-key-cancel_"] button,
    div[class*="st-key-fsb_cancelar_"] button {
        background: rgba(200,49,27,.10) !important;
        background-color: rgba(200,49,27,.10) !important;
        border: 1.4px solid rgba(164,37,15,.28) !important;
        color: var(--gaia-danger-dark) !important;
        box-shadow: none !important;
    }
    div[class*="st-key-cancel_"] button:hover,
    div[class*="st-key-fsb_cancelar_"] button:hover {
        background: rgba(200,49,27,.17) !important;
        background-color: rgba(200,49,27,.17) !important;
        border-color: rgba(164,37,15,.46) !important;
        color: #7A1D0F !important;
        transform: translateY(-1px);
        box-shadow: 0 8px 16px -12px rgba(164,37,15,.35) !important;
    }

    /* ==========================================
       BOTÕES DE VOLTAR/RETORNAR — vermelho translúcido mais claro que o Cancelar.
       Cobre: voltar_* (keys explícitas) e fsb_voltar_* (submit de forms).
       ========================================== */
    div[class*="st-key-voltar_"] button,
    div[class*="st-key-fsb_voltar_"] button {
        background: rgba(200,49,27,.055) !important;
        background-color: rgba(200,49,27,.055) !important;
        border: 1.4px solid rgba(164,37,15,.18) !important;
        color: var(--gaia-danger-dark) !important;
        box-shadow: none !important;
    }
    div[class*="st-key-voltar_"] button:hover,
    div[class*="st-key-fsb_voltar_"] button:hover {
        background: rgba(200,49,27,.10) !important;
        background-color: rgba(200,49,27,.10) !important;
        border-color: rgba(164,37,15,.32) !important;
        color: #7A1D0F !important;
        transform: translateY(-1px);
        box-shadow: 0 8px 16px -12px rgba(164,37,15,.22) !important;
    }

    /* Consulta, histórico, visualização e downloads — azul/roxo translúcido. */
    div[class*="st-key-hist_geral_ver_json_"] button,
    div[class*="st-key-hist_empresa_ver_json_"] button,
    div[class*="st-key-hist_geral_docx_"] button,
    div[class*="st-key-hist_geral_pdf_"] button,
    div[class*="st-key-hist_geral_comp_"] button,
    div[class*="st-key-hist_empresa_docx_"] button,
    div[class*="st-key-hist_empresa_pdf_"] button,
    div[class*="st-key-hist_empresa_comp_"] button {
        background: rgba(14,124,168,.10) !important;
        background-color: rgba(14,124,168,.10) !important;
        border-color: rgba(14,124,168,.26) !important;
        color: #0A5068 !important;
    }
    div[class*="st-key-hist_geral_ver_json_"] button:hover,
    div[class*="st-key-hist_empresa_ver_json_"] button:hover,
    div[class*="st-key-hist_geral_docx_"] button:hover,
    div[class*="st-key-hist_geral_pdf_"] button:hover,
    div[class*="st-key-hist_geral_comp_"] button:hover,
    div[class*="st-key-hist_empresa_docx_"] button:hover,
    div[class*="st-key-hist_empresa_pdf_"] button:hover,
    div[class*="st-key-hist_empresa_comp_"] button:hover {
        background: rgba(14,124,168,.16) !important;
        border-color: rgba(14,124,168,.44) !important;
        color: #083C4E !important;
    }

    /* Botão sair: cinza frio, discreto. */
    div.st-key-btn_logout_home button,
    div.st-key-side_menu_sair button {
        background: rgba(255,255,255,.08) !important;
        border-color: rgba(255,255,255,.20) !important;
        color: rgba(255,255,255,.92) !important;
    }

    /* ==========================================
       BARRA LATERAL — estilo portal administrativo
       ========================================== */
    [data-testid="stSidebarHeader"],
    [data-testid="stSidebarCollapseButton"],
    [data-testid="collapsedControl"],
    button[kind="header"],
    button[title="Collapse sidebar"],
    button[aria-label="Collapse sidebar"],
    button[aria-label="Close sidebar"],
    button[aria-label="Open sidebar"] {
        display: none !important;
        visibility: hidden !important;
        width: 0 !important;
        height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }
    [data-testid="stSidebarHeader"] {
        min-height: 0 !important;
        height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
    }
    section[data-testid="stSidebar"] {
        background: linear-gradient(195deg, var(--gaia-sidebar) 0%, var(--gaia-sidebar-dark) 100%) !important;
        border-right: none !important;
        box-shadow: 8px 0 26px rgba(0, 50, 40, .22);
    }
    section[data-testid="stSidebar"] .block-container {
        padding-top: 1.7rem !important;
        padding-left: 1.05rem !important;
        padding-right: 1.05rem !important;
    }
    section[data-testid="stSidebar"] * { color: rgba(255,255,255,.94) !important; }
    .gaia-side-logo {
        display:block; width: 190px; max-width: 92%; height:auto; object-fit:contain;
        margin: 0 0 16px 0; padding: 0; background: transparent;
        filter: drop-shadow(0 5px 12px rgba(0,0,0,.20));
    }
    .gaia-side-wordmark {
        display:flex; align-items:center; gap:10px; margin: 0 0 12px 0;
    }
    .gaia-side-wordmark-badge {
        width:38px; height:38px; border-radius:11px; background:rgba(255,255,255,.14);
        border:1px solid rgba(255,255,255,.18); display:flex; align-items:center; justify-content:center;
        color:#fff !important; font-weight:800;
    }
    .gaia-side-wordmark-text { color:#fff !important; font-size:1.45rem; font-weight:900; letter-spacing:.02em; }
    .gaia-side-title { font-size:1.2rem; font-weight: 800; color:white; letter-spacing:.01em; margin-bottom:2px; }
    .gaia-side-sub { font-size:.77rem; line-height:1.45; color:rgba(255,255,255,.70) !important; margin-bottom: 16px; }
    .gaia-side-user {
        display:flex; align-items:center; gap:11px;
        border: 1px solid rgba(255,255,255,.16);
        background: rgba(255,255,255,.08);
        border-radius: var(--gaia-radius-md);
        padding: 12px 13px;
        margin: 14px 0 18px 0;
    }
    .gaia-side-avatar {
        width: 38px; height: 38px; border-radius: 50%; flex: 0 0 auto;
        background: linear-gradient(135deg, var(--gaia-accent), var(--gaia-accent-dark));
        display:flex; align-items:center; justify-content:center;
        font-weight: 800; font-size: .95rem; color: #2B1B00 !important;
        box-shadow: 0 3px 8px rgba(0,0,0,.18);
    }
    .gaia-side-user b { color:white !important; font-size:.88rem; }
    .gaia-side-user span { color: rgba(255,255,255,.68) !important; font-size: .76rem; }
    .gaia-side-section {
        color:rgba(255,255,255,.60) !important; font-size:.67rem; font-weight:800;
        letter-spacing:.13em; text-transform:uppercase; margin:20px 2px 9px 2px;
    }
    section[data-testid="stSidebar"] div[data-testid="stButton"] button,
    section[data-testid="stSidebar"] .stButton button {
        background: transparent !important;
        background-color: transparent !important;
        border: 1px solid transparent !important;
        color: rgba(255,255,255,.88) !important;
        border-radius: var(--gaia-radius-sm) !important;
        text-align: left !important;
        justify-content: flex-start !important;
        min-height: 42px !important;
        padding-left: 14px !important;
        font-weight: 500 !important;
        font-size: .88rem !important;
        box-shadow: none !important;
        transition: background-color var(--gaia-ease), border-color var(--gaia-ease), padding-left var(--gaia-ease) !important;
        letter-spacing: .01em !important;
    }
    section[data-testid="stSidebar"] div[data-testid="stButton"] button:hover {
        background: var(--gaia-sidebar-soft) !important;
        border-color: rgba(255,255,255,.22) !important;
        padding-left: 17px !important;
        color: #ffffff !important;
    }
    /* Página atual: indicador sólido à esquerda em destaque âmbar */
    section[data-testid="stSidebar"] div[data-testid="stButton"] button:disabled {
        background: var(--gaia-sidebar-active) !important;
        border: 1px solid rgba(255,255,255,.20) !important;
        border-left: 3px solid var(--gaia-accent) !important;
        opacity: 1 !important;
        color: #FFFFFF !important;
        font-weight: 700 !important;
        padding-left: 13px !important;
        cursor: default !important;
    }

    @media (max-width: 980px) {
        .block-container { padding-left: .9rem !important; padding-right: .9rem !important; }
        .gaia-topbar { flex-direction:column; align-items:flex-start; }
        .gaia-login-page { margin-top: .8rem; }
        div.st-key-card_continuar_alt button,
        div.st-key-card_nova_empresa button,
        div.st-key-card_historico_docs button { min-height: 124px !important; }
    }


    /* ==========================================
       AJUSTES FINAIS — v8.22
       ========================================== */

    /* Botões de fluxo principal: verde claro, quase transparente,
       aplicado por key para evitar depender do tema interno do Streamlit. */
    div.st-key-btn_login_entrar button,
    div.st-key-btn_acessar_empresa button,
    div.st-key-btn_avancar_para_selecao button,
    div.st-key-btn_avancar_eventos button,
    div.st-key-download_docx_gerado_persistente button,
    div.st-key-download_pdf_gerado_persistente button {
        background: rgba(0, 134, 107, .12) !important;
        background-color: rgba(0, 134, 107, .12) !important;
        border: 1.4px solid rgba(0, 107, 86, .34) !important;
        color: var(--gaia-primary-darker) !important;
        box-shadow: none !important;
    }
    div.st-key-btn_login_entrar button:hover,
    div.st-key-btn_acessar_empresa button:hover,
    div.st-key-btn_avancar_para_selecao button:hover,
    div.st-key-btn_avancar_eventos button:hover,
    div.st-key-download_docx_gerado_persistente button:hover,
    div.st-key-download_pdf_gerado_persistente button:hover {
        background: rgba(0, 134, 107, .18) !important;
        background-color: rgba(0, 134, 107, .18) !important;
        border-color: rgba(0, 107, 86, .52) !important;
        color: var(--gaia-primary-darker) !important;
        box-shadow: 0 8px 16px -10px rgba(0,107,86,.45) !important;
    }

    @keyframes gaia-download-pulse {
        0%, 100% { box-shadow: 0 10px 22px -12px rgba(0,107,86,.55), 0 0 0 0 rgba(0,134,107,.18); }
        50% { box-shadow: 0 12px 26px -10px rgba(0,107,86,.65), 0 0 0 5px rgba(0,134,107,.08); }
    }
    /* Downloads finais: aparecem após gerar o documento e precisam chamar mais atenção. */
    div.st-key-download_docx_gerado_persistente button,
    div.st-key-download_pdf_gerado_persistente button {
        background: linear-gradient(180deg, rgba(0,134,107,.22), rgba(0,134,107,.14)) !important;
        background-color: rgba(0,134,107,.20) !important;
        border: 1.8px solid rgba(0,107,86,.58) !important;
        color: var(--gaia-primary-darker) !important;
        font-weight: 800 !important;
        min-height: 48px !important;
        box-shadow: 0 10px 22px -12px rgba(0,107,86,.55) !important;
        animation: gaia-download-pulse 1.8s ease-in-out 2;
    }
    div.st-key-download_docx_gerado_persistente button:hover,
    div.st-key-download_pdf_gerado_persistente button:hover {
        background: linear-gradient(180deg, rgba(0,134,107,.30), rgba(0,134,107,.20)) !important;
        border-color: rgba(0,107,86,.72) !important;
        transform: translateY(-1px) !important;
        box-shadow: 0 14px 26px -12px rgba(0,107,86,.70) !important;
    }

    /* Novo sócio: ação de apoio, azul claro no mesmo padrão translúcido. */
    div.st-key-btn_novo_socio button {
        background: rgba(14, 124, 168, .10) !important;
        background-color: rgba(14, 124, 168, .10) !important;
        border: 1.4px solid rgba(14, 124, 168, .28) !important;
        color: #0A5068 !important;
        box-shadow: none !important;
    }
    div.st-key-btn_novo_socio button:hover {
        background: rgba(14, 124, 168, .16) !important;
        background-color: rgba(14, 124, 168, .16) !important;
        border-color: rgba(14, 124, 168, .44) !important;
        color: #083C4E !important;
        box-shadow: 0 8px 16px -12px rgba(14,124,168,.35) !important;
    }

    /* Seleção de eventos: ao marcar, o feedback fica claro no quadradinho.
       O Streamlit/BaseWeb muda a estrutura interna entre versões; por isso
       combinamos seletores nativos com um check visual próprio via CSS. */
    div[data-testid="stCheckbox"] {
        background: #FFFFFF !important;
        border: 1px solid var(--gaia-border) !important;
        max-width: max-content !important;
        margin-bottom: 6px !important;
    }
    div[data-testid="stCheckbox"]:has(input:checked) {
        background: rgba(0,134,107,.045) !important;
        border-color: rgba(0,134,107,.24) !important;
    }
    /* neutraliza a regra antiga que pintava o texto inteiro quando marcado */
    div[data-testid="stCheckbox"] input:checked + div {
        background: transparent !important;
        background-color: transparent !important;
        border-color: var(--gaia-border-strong) !important;
    }
    div[data-testid="stCheckbox"] label,
    div[data-testid="stCheckbox"] label div,
    div[data-testid="stCheckbox"] label span,
    div[data-testid="stCheckbox"] label p {
        background: transparent !important;
        background-color: transparent !important;
        color: var(--gaia-text) !important;
    }
    /* Card menor e conteúdo em linha: quadradinho à esquerda, texto à direita. */
    div[data-testid="stCheckbox"] label {
        display: flex !important;
        align-items: center !important;
        gap: 9px !important;
        width: auto !important;
        min-width: unset !important;
        max-width: 100% !important;
        padding: 7px 10px !important;
        margin: 0 !important;
        box-sizing: border-box !important;
        font-size: .88rem !important;
    }
    div[data-testid="stCheckbox"] label > div,
    div[data-testid="stCheckbox"] label [data-testid="stMarkdownContainer"] {
        display: inline-flex !important;
        align-items: center !important;
        width: auto !important;
        margin: 0 !important;
    }
    div[data-testid="stCheckbox"] label p {
        margin: 0 !important;
        line-height: 1.25 !important;
    }
    div[data-testid="stCheckbox"]:has(input:checked) label div,
    div[data-testid="stCheckbox"]:has(input:checked) label span,
    div[data-testid="stCheckbox"]:has(input:checked) label p {
        background: transparent !important;
        background-color: transparent !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div:first-child {
        background: #FFFFFF !important;
        background-color: #FFFFFF !important;
        border: 1.5px solid var(--gaia-border-strong) !important;
        border-radius: 6px !important;
        width: 20px !important;
        height: 20px !important;
        position: relative !important;
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] {
        position: relative !important;
        width: 20px !important;
        height: 20px !important;
        flex: 0 0 20px !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] input:checked + div,
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div[aria-checked="true"],
    div[data-testid="stCheckbox"]:has(input:checked) span[data-baseweb="checkbox"] > div:first-child {
        background: rgba(0,134,107,.16) !important;
        background-color: rgba(0,134,107,.16) !important;
        border-color: var(--gaia-primary) !important;
        position: relative !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] input:checked + div svg,
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div[aria-checked="true"] svg {
        color: var(--gaia-primary-dark) !important;
        fill: var(--gaia-primary-dark) !important;
        stroke: var(--gaia-primary-dark) !important;
        opacity: 1 !important;
        visibility: visible !important;
        display: block !important;
    }
    /* Remove o check desenhado no contêiner, que podia cair sobre o texto. */
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"]:has(input:checked)::after,
    div[data-testid="stCheckbox"]:has(input:checked) span[data-baseweb="checkbox"]::after {
        content: none !important;
        display: none !important;
    }
    /* Check desenhado somente dentro da caixinha. */
    div[data-testid="stCheckbox"]:has(input:checked) span[data-baseweb="checkbox"] > div:first-child::after {
        content: "" !important;
        position: absolute !important;
        left: 6px !important;
        top: 3px !important;
        width: 5px !important;
        height: 10px !important;
        border: solid var(--gaia-primary-dark) !important;
        border-width: 0 2px 2px 0 !important;
        transform: rotate(45deg) !important;
        pointer-events: none !important;
        z-index: 4 !important;
    }
    div[data-testid="stCheckbox"] input[type="checkbox"] {
        accent-color: var(--gaia-primary) !important;
    }

    /* Ajuste de leitura nos cards clicáveis. */
    div.st-key-card_continuar_alt button,
    div.st-key-card_nova_empresa button,
    div.st-key-card_historico_docs button {
        white-space: normal !important;
        letter-spacing: -.01em !important;
    }


    /* ==========================================
       AJUSTE v8.23 — campo de senha limpo
       ==========================================
       O Streamlit/BaseWeb adiciona automaticamente instruções como
       "Press Enter to submit form" e um botão de revelar senha. No tema
       customizado, esses elementos estavam invadindo a borda do input e
       deixando o "olhinho" desalinhado. Aqui o campo de senha passa a usar
       uma caixa única, sem borda duplicada e sem o ícone nativo quebrado. */

    [data-testid="InputInstructions"],
    div[data-testid="InputInstructions"] {
        display: none !important;
        visibility: hidden !important;
        height: 0 !important;
        min-height: 0 !important;
        margin: 0 !important;
        padding: 0 !important;
        overflow: hidden !important;
    }

    /* ==========================================
       CORREÇÃO DEFINITIVA — faixa/retângulo sobrando após o olho
       ==========================================
       Causa real (confirmada via DevTools no navegador do usuário):
       o Streamlit moderno gera DOIS containers arredondados sobrepostos:

         div[data-baseweb="input"]      (data-testid="stTextInputRootElement")
           └─ div[data-baseweb="base-input"]   (display:flex)
                ├─ input
                └─ button (revelar senha)

       O <input> não tinha "flex: 1 1 auto", então não esticava para ocupar
       o espaço entre o texto digitado e o botão do olho. Isso deixava uma
       área "morta" dentro do base-input — pintada com o fundo/borda do
       base-input, por cima da borda colorida do elemento-raiz (visível
       como aquela faixa retangular sobrando à direita, antes do ícone).

       A correção: o elemento-raiz (stTextInputRootElement) passa a ser o
       ÚNICO responsável pela borda/raio visível; o base-input interno fica
       sem borda e sem raio próprios (apenas herda o fundo branco); e o
       input ganha flex:1 para preencher 100% do espaço disponível. */

    div[data-baseweb="input"] {
        background: #FFFFFF !important;
        border: 1.4px solid var(--gaia-border-strong) !important;
        border-radius: var(--gaia-radius-sm) !important;
        min-height: var(--gaia-control-h) !important;
        overflow: hidden !important;
        box-shadow: none !important;
    }
    div[data-baseweb="input"]:focus-within {
        border-color: var(--gaia-primary) !important;
        box-shadow: var(--gaia-focus-ring) !important;
    }

    div[data-baseweb="base-input"] {
        background: transparent !important;
        background-color: transparent !important;
        border: none !important;
        border-radius: 0 !important;
        min-height: calc(var(--gaia-control-h) - 2px) !important;
        overflow: hidden !important;
        box-shadow: none !important;
        display: flex !important;
        align-items: center !important;
        width: 100% !important;
    }

    div[data-baseweb="base-input"] input {
        flex: 1 1 auto !important;
        width: 100% !important;
        min-width: 0 !important;
        border: none !important;
        outline: none !important;
        box-shadow: none !important;
        background: transparent !important;
        background-color: transparent !important;
        min-height: calc(var(--gaia-control-h) - 2px) !important;
        padding-right: 12px !important;
    }

    div[data-baseweb="base-input"] input:focus {
        border: none !important;
        outline: none !important;
        box-shadow: none !important;
    }

    /* Botão de revelar senha: tamanho confortável, sem fundo/borda próprios,
       e sem ocupar espaço além do necessário (flex: 0 0 auto). */
    div[data-baseweb="base-input"] button {
        flex: 0 0 auto !important;
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        width: 34px !important;
        min-width: 34px !important;
        height: 34px !important;
        min-height: 34px !important;
        margin: 0 2px 0 0 !important;
        padding: 0 !important;
        border: 0 !important;
        outline: 0 !important;
        background: transparent !important;
        background-color: transparent !important;
        color: var(--gaia-muted-soft) !important;
        box-shadow: none !important;
        opacity: .65 !important;
        border-radius: 999px !important;
    }
    div[data-baseweb="base-input"] button:hover {
        background: rgba(0,134,107,.08) !important;
        color: var(--gaia-primary-dark) !important;
        opacity: 1 !important;
        transform: none !important;
        box-shadow: none !important;
    }
    div[data-baseweb="base-input"] button svg {
        display: block !important;
        visibility: visible !important;
        width: 20px !important;
        height: 20px !important;
        min-width: 20px !important;
        min-height: 20px !important;
        max-width: 20px !important;
        max-height: 20px !important;
        margin: 0 !important;
        padding: 0 !important;
        color: currentColor !important;
        fill: currentColor !important;
        stroke: currentColor !important;
    }
    div[data-baseweb="base-input"] input[type="password"],
    div[data-baseweb="base-input"] input[type="text"] {
        padding-right: 8px !important;
    }

    /* Botões de cadastro/voltar em telas de formulário seguem a mesma lógica translúcida. */
    div.st-key-btn_criar_conta_login button,
    div.st-key-btn_logout_home button,
    div.st-key-side_menu_sair button {
        background: rgba(90, 102, 121, .08) !important;
        background-color: rgba(90, 102, 121, .08) !important;
        border: 1.4px solid rgba(90, 102, 121, .24) !important;
        color: var(--gaia-text) !important;
    }
    div.st-key-btn_criar_conta_login button:hover,
    div.st-key-btn_logout_home button:hover,
    div.st-key-side_menu_sair button:hover {
        background: rgba(90, 102, 121, .13) !important;
        background-color: rgba(90, 102, 121, .13) !important;
        border-color: rgba(90, 102, 121, .38) !important;
    }

    /* Botões do cadastro de usuário: evita quebra como "CONFIRMA / R". */
    div[data-testid="stFormSubmitButton"] button {
        white-space: nowrap !important;
        word-break: keep-all !important;
        overflow-wrap: normal !important;
        text-align: center !important;
        justify-content: center !important;
        min-width: 0 !important;
        padding-left: 10px !important;
        padding-right: 10px !important;
    }

    /* ==========================================================
       OVERRIDE FINAL - CHECKBOX DE EVENTOS
       Foco: quadradinho real à esquerda + check dentro da caixa.
       ========================================================== */
    div[data-testid="stCheckbox"] {
        display: inline-block !important;
        width: auto !important;
        max-width: 100% !important;
        padding: 0 !important;
        margin: 0 0 8px 0 !important;
        background: #FFFFFF !important;
        border: 1px solid var(--gaia-border) !important;
        border-radius: 10px !important;
        box-shadow: none !important;
    }
    div[data-testid="stCheckbox"]:has(input:checked) {
        background: rgba(0,134,107,.045) !important;
        border-color: rgba(0,134,107,.28) !important;
    }
    div[data-testid="stCheckbox"] label {
        display: inline-flex !important;
        flex-direction: row !important;
        align-items: center !important;
        justify-content: flex-start !important;
        gap: 10px !important;
        width: auto !important;
        min-width: 0 !important;
        max-width: 100% !important;
        padding: 8px 12px !important;
        margin: 0 !important;
        cursor: pointer !important;
        line-height: 1.25 !important;
        background: transparent !important;
    }
    div[data-testid="stCheckbox"] label > div,
    div[data-testid="stCheckbox"] label [data-testid="stMarkdownContainer"],
    div[data-testid="stCheckbox"] label p,
    div[data-testid="stCheckbox"] label span {
        width: auto !important;
        margin: 0 !important;
        padding: 0 !important;
        background: transparent !important;
        background-color: transparent !important;
        color: var(--gaia-text) !important;
        line-height: 1.25 !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] {
        order: 0 !important;
        position: relative !important;
        display: inline-flex !important;
        align-items: center !important;
        justify-content: center !important;
        width: 18px !important;
        min-width: 18px !important;
        height: 18px !important;
        min-height: 18px !important;
        flex: 0 0 18px !important;
        margin: 0 !important;
        padding: 0 !important;
        background: transparent !important;
    }
    div[data-testid="stCheckbox"] input[type="checkbox"] {
        appearance: none !important;
        -webkit-appearance: none !important;
        position: relative !important;
        display: inline-block !important;
        width: 18px !important;
        min-width: 18px !important;
        height: 18px !important;
        min-height: 18px !important;
        margin: 0 !important;
        padding: 0 !important;
        opacity: 1 !important;
        border: 1.4px solid var(--gaia-border-strong) !important;
        border-radius: 5px !important;
        background: #FFFFFF !important;
        background-color: #FFFFFF !important;
        background-repeat: no-repeat !important;
        background-position: center !important;
        background-size: 13px 13px !important;
        box-shadow: none !important;
        cursor: pointer !important;
    }
    div[data-testid="stCheckbox"] input[type="checkbox"]:checked {
        border-color: var(--gaia-primary) !important;
        background-color: rgba(0,134,107,.16) !important;
        background-image: url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='14' height='14' viewBox='0 0 14 14'%3E%3Cpath d='M3 7.2 5.8 10 11 4' fill='none' stroke='%2300543F' stroke-width='2.1' stroke-linecap='round' stroke-linejoin='round'/%3E%3C/svg%3E") !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div,
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div:first-child {
        position: absolute !important;
        inset: 0 !important;
        width: 18px !important;
        height: 18px !important;
        min-width: 18px !important;
        min-height: 18px !important;
        margin: 0 !important;
        padding: 0 !important;
        border: 0 !important;
        background: transparent !important;
        background-color: transparent !important;
        pointer-events: none !important;
    }
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div::before,
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"] > div::after,
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"]::before,
    div[data-testid="stCheckbox"] span[data-baseweb="checkbox"]::after {
        content: none !important;
        display: none !important;
    }

    /* ==========================================================
       OVERRIDE FINAL - MENU LATERAL
       Opções sempre delimitadas, hover mais escuro e clique visível.
       ========================================================== */
    section[data-testid="stSidebar"] div[data-testid="stButton"] {
        margin-bottom: 7px !important;
    }
    section[data-testid="stSidebar"] div[data-testid="stButton"] button,
    section[data-testid="stSidebar"] .stButton button {
        background: rgba(255,255,255,.055) !important;
        background-color: rgba(255,255,255,.055) !important;
        border: 1px solid rgba(255,255,255,.16) !important;
        border-radius: 10px !important;
        color: rgba(255,255,255,.93) !important;
        min-height: 42px !important;
        padding: 0 14px !important;
        text-align: left !important;
        justify-content: flex-start !important;
        font-weight: 600 !important;
        box-shadow: 0 1px 0 rgba(255,255,255,.05) inset !important;
        transform: translateX(0) scale(1) !important;
        transition:
            background-color 150ms ease,
            border-color 150ms ease,
            transform 120ms ease,
            box-shadow 150ms ease !important;
    }
    section[data-testid="stSidebar"] div[data-testid="stButton"] button:hover,
    section[data-testid="stSidebar"] .stButton button:hover {
        background: rgba(0,0,0,.16) !important;
        background-color: rgba(0,0,0,.16) !important;
        border-color: rgba(255,255,255,.32) !important;
        color: #FFFFFF !important;
        transform: translateX(3px) !important;
        box-shadow: 0 5px 14px rgba(0,0,0,.16), 0 1px 0 rgba(255,255,255,.08) inset !important;
    }
    section[data-testid="stSidebar"] div[data-testid="stButton"] button:active,
    section[data-testid="stSidebar"] .stButton button:active {
        background: rgba(0,0,0,.24) !important;
        background-color: rgba(0,0,0,.24) !important;
        border-color: rgba(255,255,255,.42) !important;
        transform: translateX(3px) scale(.985) !important;
        box-shadow: 0 0 0 3px rgba(255,255,255,.08), 0 2px 8px rgba(0,0,0,.18) inset !important;
    }
    /* Página atual: botão desabilitado pelo Streamlit, mas visualmente ativo. */
    section[data-testid="stSidebar"] div[data-testid="stButton"] button:disabled,
    section[data-testid="stSidebar"] .stButton button:disabled {
        background: rgba(255,255,255,.18) !important;
        background-color: rgba(255,255,255,.18) !important;
        border: 1px solid rgba(255,255,255,.34) !important;
        border-left: 4px solid var(--gaia-accent) !important;
        color: #FFFFFF !important;
        opacity: 1 !important;
        font-weight: 800 !important;
        padding-left: 12px !important;
        box-shadow: 0 7px 18px rgba(0,0,0,.14), 0 0 0 1px rgba(255,255,255,.03) inset !important;
        transform: translateX(0) !important;
        cursor: default !important;
    }
    section[data-testid="stSidebar"] div[data-testid="stButton"] button:disabled::after,
    section[data-testid="stSidebar"] .stButton button:disabled::after {
        content: "";
        margin-left: auto;
        width: 7px;
        height: 7px;
        border-radius: 999px;
        background: var(--gaia-accent);
        box-shadow: 0 0 0 4px rgba(240,162,27,.14);
    }


    /* Ajuste de apresentação — botões principais bem destacados em verde. */
    div.st-key-btn_confirmar_registro button,
    div.st-key-btn_avancar_eventos button,
    div[class*="st-key-btn_seguir_wizard_"] button {
        background: rgba(0, 134, 107, .18) !important;
        background-color: rgba(0, 134, 107, .18) !important;
        border: 1.5px solid rgba(0, 107, 86, .50) !important;
        color: var(--gaia-primary-darker) !important;
        font-weight: 800 !important;
    }
    div.st-key-btn_confirmar_registro button:hover,
    div.st-key-btn_avancar_eventos button:hover,
    div[class*="st-key-btn_seguir_wizard_"] button:hover {
        background: rgba(0, 134, 107, .25) !important;
        background-color: rgba(0, 134, 107, .25) !important;
        border-color: rgba(0, 107, 86, .68) !important;
        color: var(--gaia-primary-darker) !important;
    }



    /* Botões invisíveis usados só para o Enter acionar a ação principal sem inverter a ordem visual. */
    div[class*="st-key-enter_default_"] {
        position: absolute !important;
        width: 1px !important;
        height: 1px !important;
        min-height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
        overflow: hidden !important;
        opacity: 0 !important;
        pointer-events: none !important;
    }
    div[class*="st-key-enter_default_"] button {
        width: 1px !important;
        height: 1px !important;
        min-height: 0 !important;
        padding: 0 !important;
        margin: 0 !important;
        border: 0 !important;
        box-shadow: none !important;
    }

    </style>
""", unsafe_allow_html=True)

def scroll_to_top():
    """Rola a tela para o topo."""
    components.html(
        """
        <script>
        setTimeout(function() {
            window.parent.scrollTo({top: 0, behavior: 'smooth'});
        }, 80);
        </script>
        """,
        height=0,
        width=0,
    )


def scroll_para_downloads():
    """Rola a tela até o painel de download após gerar o documento."""
    components.html(
        """
        <script>
        setTimeout(function() {
            const doc = window.parent.document;
            const el = doc.getElementById('documento-gerado');
            if (el) {
                el.scrollIntoView({behavior: 'smooth', block: 'start'});
            } else {
                window.parent.scrollTo({top: document.body.scrollHeight, behavior: 'smooth'});
            }
        }, 300);
        </script>
        """,
        height=0,
        width=0,
    )


def _limpar_markdown_msg(texto):
    """Remove marcações simples para a mensagem lateral não aparecer com asteriscos."""
    if texto is None:
        return ""
    return str(texto).replace("**", "").replace("`", "")


def alerta_lateral(tipo, texto):
    """Exibe mensagem lateral única e repetível."""
    texto = _limpar_markdown_msg(texto)
    cores = {
        'success': {'bg': '#e9f7ef', 'border': '#28a745', 'text': '#155724', 'icon': '✅'},
        'error': {'bg': '#fdecea', 'border': '#dc3545', 'text': '#721c24', 'icon': '❌'},
        'warning': {'bg': '#fff8e1', 'border': '#ffc107', 'text': '#856404', 'icon': '⚠️'},
        'info': {'bg': '#e8f4fd', 'border': '#17a2b8', 'text': '#0c5460', 'icon': 'ℹ️'},
    }
    cfg = cores.get(tipo, cores['info'])

    st.session_state['_alerta_lateral_seq'] = st.session_state.get('_alerta_lateral_seq', 0) + 1
    alerta_id = f"gaia-alerta-lateral-{st.session_state['_alerta_lateral_seq']}"

    payload = {
        'id': alerta_id,
        'texto': texto,
        'bg': cfg['bg'],
        'border': cfg['border'],
        'text': cfg['text'],
        'icon': cfg['icon'],
        'ttl': 4200,
    }
    components.html(
        f"""
        <script data-gaia-alerta="{alerta_id}">
        (function() {{
            const data = {json.dumps(payload, ensure_ascii=False)};
            const doc = window.parent.document;

            doc.querySelectorAll('.gaia-alerta-lateral').forEach(el => el.remove());

            if (!doc.getElementById('gaia-alerta-lateral-style')) {{
                const style = doc.createElement('style');
                style.id = 'gaia-alerta-lateral-style';
                style.textContent = `
                    @keyframes gaiaAlertaEntrada {{
                        from {{ opacity: 0; transform: translateX(16px); }}
                        to {{ opacity: 1; transform: translateX(0); }}
                    }}
                    @keyframes gaiaAlertaSaida {{
                        from {{ opacity: 1; transform: translateX(0); }}
                        to {{ opacity: 0; transform: translateX(16px); }}
                    }}
                `;
                doc.head.appendChild(style);
            }}

            const box = doc.createElement('div');
            box.id = data.id;
            box.className = 'gaia-alerta-lateral';
            box.dataset.expiresAt = String(Date.now() + data.ttl + 600);
            box.innerHTML = `<span style="font-size:18px; margin-right:8px;">${{data.icon}}</span><span>${{data.texto}}</span>`;
            box.style.position = 'fixed';
            box.style.top = '22px';
            box.style.right = '22px';
            box.style.maxWidth = '420px';
            box.style.minWidth = '260px';
            box.style.padding = '14px 16px';
            box.style.borderLeft = `6px solid ${{data.border}}`;
            box.style.background = data.bg;
            box.style.color = data.text;
            box.style.borderRadius = '10px';
            box.style.boxShadow = '0 8px 24px rgba(0,0,0,0.18)';
            box.style.fontFamily = 'Arial, sans-serif';
            box.style.fontSize = '14px';
            box.style.lineHeight = '1.35';
            box.style.zIndex = '2147483647';
            box.style.display = 'flex';
            box.style.alignItems = 'center';
            box.style.opacity = '0';
            box.style.transform = 'translateX(16px)';
            box.style.transition = 'opacity 250ms ease, transform 250ms ease';
            box.style.animation = 'gaiaAlertaEntrada 180ms ease forwards, gaiaAlertaSaida 300ms ease forwards 3.9s';
            doc.body.appendChild(box);
            const fechar = box.querySelector('button');
            const fecharBox = function() {{
                if (!box || box.dataset.closing === '1') return;
                box.dataset.closing = '1';
                box.style.opacity = '0';
                box.style.transform = 'translateX(16px)';
                window.parent.setTimeout(function() {{
                    const atual = doc.getElementById(data.id);
                    if (atual) atual.remove();
                }}, 300);
            }};
            if (fechar) fechar.onclick = fecharBox;

            window.parent.setTimeout(function() {{
                box.style.opacity = '1';
                box.style.transform = 'translateX(0)';
            }}, 20);

            window.parent.setTimeout(fecharBox, data.ttl);
            window.parent.setTimeout(function() {{
                const atual = doc.getElementById(data.id);
                if (atual) atual.remove();
            }}, data.ttl + 900);

            if (!window.parent.__gaiaAlertaCleaner) {{
                window.parent.__gaiaAlertaCleaner = window.parent.setInterval(function() {{
                    const agora = Date.now();
                    doc.querySelectorAll('.gaia-alerta-lateral').forEach(function(el) {{
                        const expira = Number(el.dataset.expiresAt || 0);
                        if (expira && expira < agora) el.remove();
                    }});
                }}, 1000);
            }}
        }})();
        </script>
        """,
        height=0,
        width=0,
    )


def set_flash(tipo, texto):
    """Guarda mensagem para aparecer após st.rerun()/troca de tela."""
    st.session_state['_flash_msg'] = {'tipo': tipo, 'texto': texto}


def show_flash():
    """Mostra mensagens persistentes somente no alerta lateral."""
    flash = st.session_state.pop('_flash_msg', None)
    if not flash:
        return
    alerta_lateral(flash.get('tipo', 'info'), flash.get('texto', ''))


def msg_success(texto):
    alerta_lateral('success', texto)


def msg_error(texto):
    alerta_lateral('error', texto)


def msg_warning(texto):
    alerta_lateral('warning', texto)


def msg_info(texto):
    alerta_lateral('info', texto)


def msg_api_offline(nome_api="API"):
    msg_warning(
        f"Não foi possível consultar a {nome_api}. Se o computador estiver sem internet, "
        "a busca automática não funcionará, mas você pode preencher os dados manualmente."
    )


def limpar_api(prefixo):
    for k in list(st.session_state.keys()):
        if k.startswith(prefixo):
            del st.session_state[k]

def _logo_html_base64():
    """Retorna a logo em base64.

    Procura primeiro no diretório de execução e depois na mesma pasta do arquivo.
    Isso evita perder a logo quando o Streamlit é executado a partir de outro diretório.
    """
    caminhos = [Path("logo.png")]
    try:
        caminhos.append(Path(__file__).resolve().with_name("logo.png"))
    except Exception:
        pass
    for logo_path in caminhos:
        try:
            if logo_path.exists():
                return base64.b64encode(logo_path.read_bytes()).decode("utf-8")
        except Exception:
            continue
    return None



def limpar_alertas_laterais():
    """Remove alertas laterais antigos quando a tela recarrega ou muda de estado."""
    components.html(
        """
        <script>
        (function() {
            const doc = window.parent.document;
            doc.querySelectorAll('.gaia-alerta-lateral').forEach(el => el.remove());
        })();
        </script>
        """,
        height=0,
        width=0,
    )

def voltar_menu_principal():
    """Volta ao menu principal sem limpar a sessão do usuário.

    Esta função não usa link HTML, não altera query params e não faz logout.
    Ela apenas muda a página interna do Streamlit.
    """
    st.session_state.pagina = 'menu' if st.session_state.get('user_id') else 'login'
    st.session_state.empresa_id = None
    st.session_state.socio_id = None
    st.session_state.mostrar_confirmacao_repetidos = False
    st.session_state.eventos_selecionados_lista = []
    st.rerun()


def cabecalho():
    """Cabeçalho padrão do GAIA."""
    logo_b64 = _logo_html_base64()
    usuario_logado = bool(st.session_state.get('user_id'))

    if usuario_logado:
        nome = html.escape(str(st.session_state.get('nome_usuario') or 'Usuário'))
        perfil = html.escape(str(st.session_state.get('user_role') or ''))
        if logo_b64:
            marca = f'<img src="data:image/png;base64,{logo_b64}" class="gaia-top-logo-img" alt="GAIA">'
        else:
            marca = '<div style="width:8px"></div>'
        inicial = (nome.strip()[:1] or "U").upper()
        st.markdown(
            f"""
            <div class="gaia-topbar">
                <div class="gaia-brand-row">
                    {marca}
                </div>
                <div class="gaia-user-pill">
                    <div class="gaia-user-avatar">{inicial}</div>
                    <span>{nome} <span style="opacity:.55;">•</span> {perfil}</span>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
    else:
        st.markdown('<div class="gaia-login-page">', unsafe_allow_html=True)
        if logo_b64:
            st.markdown(
                f'<img src="data:image/png;base64,{logo_b64}" class="gaia-logo-img" alt="GAIA">',
                unsafe_allow_html=True,
            )
        else:
            st.markdown("<div class='gaia-logo-fallback'>GAIA</div>", unsafe_allow_html=True)

    limpar_alertas_laterais()
    show_flash()

# ==========================================
# SPRINT 4 - Modelagem e Banco de Dados
# ==========================================
def get_db():
    conn = sqlite3.connect('gaia.db')
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn


def coluna_existe(conn, tabela, coluna):
    """Verifica se uma coluna já existe em uma tabela SQLite."""
    cols = conn.execute(f"PRAGMA table_info({tabela})").fetchall()
    return any(c[1] == coluna for c in cols)


def garantir_coluna(conn, tabela, coluna, definicao):
    """Adiciona coluna sem quebrar bancos antigos já criados."""
    if not coluna_existe(conn, tabela, coluna):
        conn.execute(f"ALTER TABLE {tabela} ADD COLUMN {coluna} {definicao}")


def sanitizar_nome_arquivo(texto):
    texto = (texto or "EMPRESA").strip().replace(' ', '_')
    permitido = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-"
    saida = ''.join(ch for ch in texto if ch in permitido)
    return saida or "EMPRESA"


def formatar_data_documento(data_valor=None):
    """Formata a data do documento no padrão usado no instrumento contratual."""
    data_valor = data_valor or date.today()
    meses = {
        1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril",
        5: "maio", 6: "junho", 7: "julho", 8: "agosto",
        9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro"
    }
    return f"{data_valor.day:02d} de {meses[data_valor.month]} de {data_valor.year}"


# ==========================================
# VALIDAÇÃO E PADRONIZAÇÃO DE DADOS
# ==========================================
UFS_VALIDAS = {
    "AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS",
    "MG", "PA", "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC",
    "SP", "SE", "TO"
}


def somente_digitos(valor):
    """Retorna somente os números de uma string."""
    return ''.join(ch for ch in str(valor or '') if ch.isdigit())


def remover_espacos_extras(texto):
    """Remove espaços duplicados e espaços nas extremidades."""
    return ' '.join(str(texto or '').strip().split())


def padronizar_texto(texto):
    """Padroniza textos cadastrais em caixa alta, sem espaços duplicados."""
    return remover_espacos_extras(texto).upper()


def padronizar_texto_livre(texto):
    """Padroniza textos longos preservando quebras de linha úteis, como listas de CNAEs."""
    linhas = [padronizar_texto(linha) for linha in str(texto or '').splitlines()]
    return '\n'.join(linha for linha in linhas if linha)


def formatar_cpf(cpf):
    nums = somente_digitos(cpf)
    if len(nums) != 11:
        return remover_espacos_extras(cpf)
    return f"{nums[:3]}.{nums[3:6]}.{nums[6:9]}-{nums[9:]}"


def validar_cpf(cpf):
    nums = somente_digitos(cpf)
    if len(nums) != 11 or nums == nums[0] * 11:
        return False

    soma = sum(int(nums[i]) * (10 - i) for i in range(9))
    digito_1 = (soma * 10) % 11
    if digito_1 == 10:
        digito_1 = 0
    if digito_1 != int(nums[9]):
        return False

    soma = sum(int(nums[i]) * (11 - i) for i in range(10))
    digito_2 = (soma * 10) % 11
    if digito_2 == 10:
        digito_2 = 0
    return digito_2 == int(nums[10])


def formatar_cnpj(cnpj):
    nums = somente_digitos(cnpj)
    if len(nums) != 14:
        return remover_espacos_extras(cnpj)
    return f"{nums[:2]}.{nums[2:5]}.{nums[5:8]}/{nums[8:12]}-{nums[12:]}"


def validar_cnpj(cnpj):
    nums = somente_digitos(cnpj)
    if len(nums) != 14 or nums == nums[0] * 14:
        return False

    pesos_1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    soma = sum(int(nums[i]) * pesos_1[i] for i in range(12))
    resto = soma % 11
    digito_1 = 0 if resto < 2 else 11 - resto
    if digito_1 != int(nums[12]):
        return False

    pesos_2 = [6] + pesos_1
    soma = sum(int(nums[i]) * pesos_2[i] for i in range(13))
    resto = soma % 11
    digito_2 = 0 if resto < 2 else 11 - resto
    return digito_2 == int(nums[13])


def formatar_cep(cep):
    nums = somente_digitos(cep)
    if len(nums) != 8:
        return remover_espacos_extras(cep)
    return f"{nums[:5]}-{nums[5:]}"


def validar_cep(cep):
    return len(somente_digitos(cep)) == 8


def consultar_cep_viacep(cep):
    cep_limpo = somente_digitos(cep)
    if len(cep_limpo) != 8:
        return None, "O CEP deve conter 8 dígitos."
    try:
        resp = requests.get(f"https://viacep.com.br/ws/{cep_limpo}/json/", timeout=5)
        if resp.status_code != 200:
            return None, "CEP não encontrado."
        dados_cep = resp.json()
        if "erro" in dados_cep:
            return None, "CEP não encontrado."
        return {
            "cep": formatar_cep(dados_cep.get("cep", "")),
            "logradouro": padronizar_texto(dados_cep.get("logradouro", "")),
            "bairro": padronizar_texto(dados_cep.get("bairro", "")),
            "cidade": padronizar_texto(dados_cep.get("localidade", "")),
            "uf": padronizar_texto(dados_cep.get("uf", "")),
        }, None
    except requests.exceptions.RequestException:
        return None, "ViaCEP indisponível. Preencha o endereço manualmente."
    except Exception:
        return None, "Erro inesperado ao consultar o CEP. Preencha o endereço manualmente."


def validar_uf(uf):
    return padronizar_texto(uf) in UFS_VALIDAS


def normalizar_data_br(data_texto):
    """Valida data em DD/MM/AAAA. Também aceita somente números DDMMAAAA."""
    texto = remover_espacos_extras(data_texto)
    nums = somente_digitos(texto)
    if not nums:
        return False, ""
    if len(nums) == 8 and "/" not in texto:
        texto = f"{nums[:2]}/{nums[2:4]}/{nums[4:]}"
    try:
        data_valida = datetime.strptime(texto, "%d/%m/%Y").date()
    except ValueError:
        return False, texto
    if data_valida > date.today():
        return False, texto
    if data_valida.year < 1900:
        return False, texto
    return True, data_valida.strftime("%d/%m/%Y")


def formatar_valor_brasileiro(valor):
    """Converte valores como 1000, 1.000,00, R$ 1.000,00 ou 1000.50 para 1.000,00."""
    texto = remover_espacos_extras(str(valor or '').replace('R$', '').replace('r$', ''))
    if not texto:
        return False, ""

    texto = texto.replace(' ', '')
    try:
        if ',' in texto:
            numero = Decimal(texto.replace('.', '').replace(',', '.'))
        elif '.' in texto:
            partes = texto.split('.')
            if len(partes) == 2 and len(partes[1]) <= 2:
                numero = Decimal(texto)
            else:
                numero = Decimal(texto.replace('.', ''))
        else:
            numero = Decimal(texto)
    except (InvalidOperation, ValueError):
        return False, texto

    if numero <= 0:
        return False, texto

    valor_formatado = f"{numero:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
    return True, valor_formatado


def cnpj_ja_cadastrado(cnpj, empresa_id_atual=None):
    """Evita duplicidade mesmo quando um CNPJ foi salvo com ou sem máscara."""
    nums = somente_digitos(cnpj)
    with get_db() as conn:
        empresas = conn.execute("SELECT id, cnpj FROM tb_empresa").fetchall()
    for emp_id, cnpj_salvo in empresas:
        if empresa_id_atual is not None and emp_id == empresa_id_atual:
            continue
        if somente_digitos(cnpj_salvo) == nums:
            return True
    return False


def validar_dados_empresa(dados, empresa_id_atual=None):
    erros = []
    campos_obrigatorios = {
        "CNPJ": dados.get('cnpj'),
        "Nome Empresarial": dados.get('nome_empresarial'),
        "Logradouro": dados.get('logradouro'),
        "Número": dados.get('numero'),
        "Bairro": dados.get('bairro'),
        "Cidade": dados.get('cidade'),
        "UF": dados.get('uf'),
        "CEP": dados.get('cep'),
    }
    for campo, valor in campos_obrigatorios.items():
        if not remover_espacos_extras(valor):
            erros.append(campo)

    if dados.get('cnpj') and not validar_cnpj(dados.get('cnpj')):
        erros.append("CNPJ inválido")
    if dados.get('cep') and not validar_cep(dados.get('cep')):
        erros.append("CEP inválido")
    if dados.get('uf') and not validar_uf(dados.get('uf')):
        erros.append("UF inválida")
    if dados.get('cnpj') and validar_cnpj(dados.get('cnpj')) and cnpj_ja_cadastrado(dados.get('cnpj'), empresa_id_atual):
        erros.append("Empresa já cadastrada com este CNPJ. Use a opção Continuar alteração")
    return erros


def validar_dados_socio(dados):
    erros = []
    campos_obrigatorios = {
        "Nome Completo": dados.get('nome'),
        "Nacionalidade": dados.get('nacionalidade'),
        "Profissão": dados.get('profissao'),
        "CPF": dados.get('cpf'),
        "Data de nascimento": dados.get('data_nasc'),
        "RG": dados.get('rg'),
        "Órgão emissor": dados.get('orgao_emissor'),
        "Estado civil": dados.get('estado_civil'),
        "Logradouro": dados.get('logradouro'),
        "Número": dados.get('numero'),
        "Bairro": dados.get('bairro'),
        "Cidade": dados.get('cidade'),
        "UF": dados.get('uf'),
        "CEP": dados.get('cep'),
    }
    for campo, valor in campos_obrigatorios.items():
        if not remover_espacos_extras(valor):
            erros.append(campo)

    if dados.get('cpf') and not validar_cpf(dados.get('cpf')):
        erros.append("CPF inválido")

    if dados.get('data_nasc'):
        data_ok, data_formatada = normalizar_data_br(dados.get('data_nasc'))
        if not data_ok:
            erros.append("Data de nascimento inválida")
        else:
            dados['data_nasc'] = data_formatada

    if dados.get('cep') and not validar_cep(dados.get('cep')):
        erros.append("CEP inválido")
    if dados.get('uf') and not validar_uf(dados.get('uf')):
        erros.append("UF inválida")
    return erros


def inicializar_banco():
    with get_db() as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS tb_perfil_acesso (id INTEGER PRIMARY KEY, nome_perfil TEXT UNIQUE NOT NULL)")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_usuario (
            id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT, login TEXT UNIQUE,
            senha_hash TEXT, id_perfil INTEGER,
            aviso_senha_padrao INTEGER DEFAULT 0,
            aviso_senha_padrao_visualizado INTEGER DEFAULT 0,
            FOREIGN KEY (id_perfil) REFERENCES tb_perfil_acesso (id))""")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_empresa (
            id INTEGER PRIMARY KEY AUTOINCREMENT, cnpj TEXT UNIQUE, nome_empresarial TEXT,
            nome_fantasia TEXT, nire TEXT, cidade TEXT, uf TEXT, cep TEXT, logradouro TEXT, numero TEXT,
            bairro TEXT, complemento TEXT)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_socio (
            id INTEGER PRIMARY KEY AUTOINCREMENT, id_empresa INTEGER, nome TEXT,
            nacionalidade TEXT, cpf TEXT, rg TEXT, orgao_emissor TEXT,
            estado_civil TEXT, data_nasc TEXT, profissao TEXT, cep TEXT,
            logradouro TEXT, numero TEXT, complemento TEXT, bairro TEXT, cidade TEXT, uf TEXT,
            FOREIGN KEY (id_empresa) REFERENCES tb_empresa (id) ON DELETE CASCADE)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_documento (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            id_empresa INTEGER,
            id_usuario INTEGER,
            data_geracao DATETIME DEFAULT CURRENT_TIMESTAMP,
            data_documento_texto TEXT,
            tipo_documento TEXT,
            status TEXT DEFAULT 'Gerado',
            eventos_json TEXT,
            clausulas_json TEXT,
            nome_arquivo_docx TEXT,
            nome_arquivo_pdf TEXT,
            nome_arquivo_relatorio TEXT,
            modelo_documento TEXT,
            protocolo TEXT,
            data_protocolo TEXT,
            docx_blob BLOB,
            pdf_blob BLOB,
            relatorio_conferencia_blob BLOB,
            comprovante_nome TEXT,
            comprovante_blob BLOB,
            observacao TEXT,
            FOREIGN KEY (id_empresa) REFERENCES tb_empresa (id) ON DELETE CASCADE,
            FOREIGN KEY (id_usuario) REFERENCES tb_usuario (id))""")

        # Migração: garante que bancos antigos recebam as novas colunas de histórico.
        garantir_coluna(conn, 'tb_documento', 'data_documento_texto', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'eventos_json', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'clausulas_json', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'nome_arquivo_docx', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'nome_arquivo_pdf', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'nome_arquivo_relatorio', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'modelo_documento', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'protocolo', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'data_protocolo', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'docx_blob', 'BLOB')
        garantir_coluna(conn, 'tb_documento', 'pdf_blob', 'BLOB')
        garantir_coluna(conn, 'tb_documento', 'relatorio_conferencia_blob', 'BLOB')
        garantir_coluna(conn, 'tb_documento', 'comprovante_nome', 'TEXT')
        garantir_coluna(conn, 'tb_documento', 'comprovante_blob', 'BLOB')
        garantir_coluna(conn, 'tb_documento', 'observacao', 'TEXT')

        # Migração: controle do aviso de senha padrão do administrador.
        garantir_coluna(conn, 'tb_usuario', 'aviso_senha_padrao', 'INTEGER DEFAULT 0')
        garantir_coluna(conn, 'tb_usuario', 'aviso_senha_padrao_visualizado', 'INTEGER DEFAULT 0')
        conn.commit()

        if conn.execute("SELECT COUNT(*) FROM tb_perfil_acesso").fetchone()[0] == 0:
            conn.execute("INSERT INTO tb_perfil_acesso (nome_perfil) VALUES ('Admin'), ('Operador')")
            conn.commit()

        if conn.execute("SELECT COUNT(*) FROM tb_usuario").fetchone()[0] == 0:
            h = hashlib.sha256("1234".encode()).hexdigest()
            conn.execute(
                """INSERT INTO tb_usuario 
                   (nome, login, senha_hash, id_perfil, aviso_senha_padrao, aviso_senha_padrao_visualizado)
                   VALUES (?,?,?,?,?,?)""",
                ("Administrador", "admin", h, 1, 1, 0)
            )
            conn.commit()

        # Para bancos já existentes: se o admin ainda usa a senha padrão e o aviso nunca foi visto,
        # exibe a recomendação de troca no próximo login.
        h_padrao = hashlib.sha256("1234".encode()).hexdigest()
        conn.execute("""
            UPDATE tb_usuario
               SET aviso_senha_padrao = 1
             WHERE login = 'admin'
               AND senha_hash = ?
               AND COALESCE(aviso_senha_padrao_visualizado, 0) = 0
        """, (h_padrao,))
        conn.commit()

inicializar_banco()

# ==========================================
# ESTADO E NAVEGAÇÃO
# ==========================================
defaults = {
    'pagina': 'login', 'empresa_id': None, 'socio_id': None, 'user_id': None,
    'nome_usuario': None, 'user_role': None, 'eventos_selecionados_lista': [],
    'mostrar_confirmacao_repetidos': False, 'msg_aviso_texto': '',
    'documento_gerado': None, '_confirm_excluir_documento': None,
    'login_usuario': None, '_reset_senha_usuario': None, '_alterar_senha_origem_aviso': False
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


def gerar_hash(s): return hashlib.sha256(s.encode()).hexdigest()


def senha_padrao_admin_ativa(login_usuario, senha_hash):
    """Identifica o acesso inicial de instalação do GAIA.

    O login admin/senha 1234 é mantido apenas para simplificar o primeiro uso
    em ambiente local/offline. O sistema recomenda a troca, mas não bloqueia o uso.
    """
    return str(login_usuario or '').strip().lower() == 'admin' and senha_hash == gerar_hash('1234')


def marcar_aviso_senha_padrao_visualizado(user_id):
    with get_db() as conn:
        conn.execute(
            "UPDATE tb_usuario SET aviso_senha_padrao=0, aviso_senha_padrao_visualizado=1 WHERE id=?",
            (user_id,)
        )
        conn.commit()


def validar_nova_senha(nova_senha, confirmar_senha, login_usuario=None):
    erros = []
    if not nova_senha or not confirmar_senha:
        erros.append("Informe e confirme a nova senha.")
    elif nova_senha != confirmar_senha:
        erros.append("As senhas não conferem.")
    elif len(nova_senha) < 4:
        erros.append("A senha deve ter pelo menos 4 caracteres.")
    elif str(login_usuario or '').strip().lower() == 'admin' and nova_senha == '1234':
        erros.append("Escolha uma senha diferente da senha padrão 1234.")
    return erros


def nav(p):
    st.session_state.pagina = p
    scroll_to_top()
    st.rerun()

def _form_socio(titulo, defaults=None):
    d = defaults or {}
    estados_civis = ["", "Solteiro(a)", "Casado(a)", "Divorciado(a)", "Viúvo(a)", "União Estável"]
    cabecalho()
    st.subheader(titulo)

    with st.form("f_socio"):
        st.markdown("**Dados Pessoais**")
        nome = st.text_input("NOME COMPLETO*", value=padronizar_texto(d.get('nome') or ''))       
        c1, c2 = st.columns(2)
        nac = c1.text_input("NACIONALIDADE*", value=padronizar_texto(d.get('nacionalidade') or 'BRASILEIRO(A)'))
        profissao = c2.text_input("PROFISSÃO*", value=padronizar_texto(d.get('profissao') or 'EMPRESÁRIO'))
        
        c3, c4 = st.columns(2)
        cpf = c3.text_input("CPF*", value=formatar_cpf(d.get('cpf') or ''))
        data_nasc = c4.text_input("DATA NASCIMENTO (DD/MM/AAAA)*", value=d.get('data_nasc') or '', placeholder='DD/MM/AAAA')
        
        c5, c6 = st.columns(2)
        rg = c5.text_input("RG*", value=d.get('rg') or '')
        orgao = c6.text_input("ÓRGÃO EMISSOR*", value=d.get('orgao_emissor') or '')
        
        ec_valor = d.get('estado_civil') or ""
        ec_idx = estados_civis.index(ec_valor) if ec_valor in estados_civis else 0
        estado_civil = st.selectbox("ESTADO CIVIL*", estados_civis, index=ec_idx)
        
        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown("**Morada**")
        
        c_cep, c_btn_cep = st.columns([3, 1], vertical_alignment="bottom")
        cep_valor = formatar_cep(st.session_state.get('api_socio_cep', d.get('cep') or ''))
        cep = c_cep.text_input("CEP*", value=cep_valor)
        btn_buscar_cep = c_btn_cep.form_submit_button("BUSCAR", type="secondary", use_container_width=True)
        
        if btn_buscar_cep:
            cep_limpo = ''.join(filter(str.isdigit, cep))
            if len(cep_limpo) == 8:
                try:
                    resp = requests.get(f"https://viacep.com.br/ws/{cep_limpo}/json/", timeout=5)
                    if resp.status_code == 200:
                        dados_cep = resp.json()
                        if "erro" not in dados_cep:
                            st.session_state['api_socio_cep'] = formatar_cep(dados_cep.get('cep', ''))
                            st.session_state['api_socio_logradouro'] = padronizar_texto(dados_cep.get('logradouro', ''))
                            st.session_state['api_socio_bairro'] = padronizar_texto(dados_cep.get('bairro', ''))
                            st.session_state['api_socio_cidade'] = padronizar_texto(dados_cep.get('localidade', ''))
                            st.session_state['api_socio_uf'] = padronizar_texto(dados_cep.get('uf', ''))
                            st.rerun()
                        else:
                            msg_error("CEP não encontrado.")
                except requests.exceptions.RequestException:
                    msg_api_offline("ViaCEP")
                except Exception:
                    msg_error("Erro inesperado ao consultar o CEP. Você pode preencher os dados manualmente.")
            else:
                msg_warning("O CEP deve conter 8 dígitos.")

        c7, c8 = st.columns([3, 1])
        v_log = st.session_state.get('api_socio_logradouro', d.get('logradouro') or '')
        logradouro = c7.text_input("LOGRADOURO*", value=v_log)
        numero = c8.text_input("NÚMERO*", value=d.get('numero') or '')
        
        c9, c10 = st.columns(2)
        complemento = c9.text_input("COMPLEMENTO", value=d.get('complemento') or '')
        v_bai = st.session_state.get('api_socio_bairro', d.get('bairro') or '')
        bairro = c10.text_input("BAIRRO*", value=v_bai)
        
        c11, c12 = st.columns([3, 1])
        v_cid = st.session_state.get('api_socio_cidade', d.get('cidade') or '')
        cidade = c11.text_input("CIDADE*", value=v_cid)
        v_uf = st.session_state.get('api_socio_uf', d.get('uf') or '')
        uf = c12.text_input("UF*", value=v_uf)
        
        st.markdown("<br>", unsafe_allow_html=True)
        btn_ok_enter = st.form_submit_button("Confirmar", type="primary", use_container_width=True, key="enter_default_socio")
        b1, b2 = st.columns(2)
        btn_v = b1.form_submit_button("← Cancelar", type="secondary", use_container_width=True, key="fsb_cancelar_socio")
        btn_ok_click = b2.form_submit_button("Confirmar", type="primary", use_container_width=True, key="btn_confirmar_socio")
        btn_ok = btn_ok_enter or btn_ok_click

    if btn_v or btn_ok:
        for k in list(st.session_state.keys()):
            if k.startswith('api_socio_'):
                del st.session_state[k]

    data_ok, data_nasc_formatada = normalizar_data_br(data_nasc) if data_nasc else (False, remover_espacos_extras(data_nasc))

    return btn_v, btn_ok, dict(
        nome=padronizar_texto(nome),
        nacionalidade=padronizar_texto(nac),
        cpf=formatar_cpf(cpf),
        rg=padronizar_texto(rg),
        orgao_emissor=padronizar_texto(orgao),
        estado_civil=estado_civil,
        data_nasc=data_nasc_formatada if data_ok else remover_espacos_extras(data_nasc),
        profissao=padronizar_texto(profissao),
        logradouro=padronizar_texto(logradouro),
        numero=remover_espacos_extras(numero),
        complemento=padronizar_texto(complemento),
        bairro=padronizar_texto(bairro),
        cidade=padronizar_texto(cidade),
        uf=padronizar_texto(uf),
        cep=formatar_cep(cep)
    )

# ==========================================
# TELAS DO SISTEMA
# ==========================================
def tela_login():
    cabecalho()
    st.markdown('<div class="gaia-login-heading">Acesso ao sistema</div>', unsafe_allow_html=True)
    st.markdown('<div class="gaia-login-caption">Entre com suas credenciais para continuar.</div>', unsafe_allow_html=True)

    col_a, col_b, col_c = st.columns([1.05, 1.9, 1.05])
    with col_b:
        with st.form("form_login"):
            login = st.text_input("Login", placeholder="Ex.: admin")
            senha = st.text_input("Senha", type="password", placeholder="Digite sua senha")
            st.markdown('<div class="gaia-help-text">Esqueceu a senha? contate o administrador.</div>', unsafe_allow_html=True)
            btn_e = st.form_submit_button("Entrar", use_container_width=True, key="btn_login_entrar")

        if st.button("Criar nova conta", use_container_width=True, key="btn_criar_conta_login"):
            nav('registrar')
            return

    st.markdown('</div>', unsafe_allow_html=True)

    if btn_e:
        with get_db() as conn:
            user = conn.execute("""
                SELECT u.id, u.nome, p.nome_perfil, u.login, u.senha_hash,
                       COALESCE(u.aviso_senha_padrao, 0) AS aviso_senha_padrao
                FROM tb_usuario u
                JOIN tb_perfil_acesso p ON u.id_perfil = p.id
                WHERE u.login=? AND u.senha_hash=?
            """, (login, gerar_hash(senha))).fetchone()
        if user:
            st.session_state.user_id = user[0]
            st.session_state.nome_usuario = user[1]
            st.session_state.user_role = user[2]
            st.session_state.login_usuario = user[3]

            if user[5] and senha_padrao_admin_ativa(user[3], user[4]):
                nav('aviso_senha_padrao')
            else:
                set_flash('success', f'Bem-vindo(a), {user[1]}!')
                nav('menu')
        else:
            msg_error("Login ou senha incorretos.")

def tela_registrar():
    cabecalho()
    col_a, col_b, col_c = st.columns([1, 1.35, 1])
    with col_b:
        st.markdown(
            """
            <div style="
                text-align:center;
                font-size:1.22rem;
                font-weight:700;
                color:var(--gaia-text);
                padding-bottom:.55rem;
                margin-bottom:.65rem;
                border-bottom:1px solid var(--gaia-border);
            ">
                Criar nova conta
            </div>
            """,
            unsafe_allow_html=True
        )
        with st.form("form_reg"):
            nome_completo = st.text_input("Nome Completo (opcional)")
            u = st.text_input("Usuário (Login)*")
            s1 = st.text_input("Senha*", type="password")
            s2 = st.text_input("Confirmar senha*", type="password")
            btn_confirmar_enter = st.form_submit_button("Confirmar", type="primary", use_container_width=True, key="enter_default_registro")
            c1, c2 = st.columns(2)
            btn_voltar = c1.form_submit_button("← Cancelar", type="secondary", use_container_width=True, key="fsb_cancelar_registro")
            btn_confirmar_click = c2.form_submit_button("Confirmar", type="primary", use_container_width=True, key="btn_confirmar_registro")
            btn_confirmar = btn_confirmar_enter or btn_confirmar_click

    if btn_voltar: nav('login')
    if btn_confirmar:
        if not u: msg_error("Informe um nome de utilizador.")
        elif s1 != s2: msg_error("As senhas não conferem.")
        elif len(s1) < 4: msg_error("A senha deve ter pelo menos 4 caracteres.")
        else:
            nome_final = nome_completo.strip() if nome_completo.strip() else u
            try:
                with get_db() as conn:
                    conn.execute("INSERT INTO tb_usuario (nome, login, senha_hash, id_perfil) VALUES (?,?,?,2)", (nome_final, u, gerar_hash(s1)))
                    conn.commit()
                set_flash("success", f"Utilizador **{u}** criado com sucesso! Faça login para continuar.")
                nav('login')
            except sqlite3.IntegrityError:
                msg_error(f"O utilizador '{u}' já existe. Escolha outro nome.")


def tela_aviso_senha_padrao():
    if not st.session_state.get('user_id'):
        nav('login')
        return

    cabecalho()
    st.subheader("Aviso de segurança")

    st.markdown(
        """
        <div class="gaia-card" style="max-width: 760px; margin: 0 auto 12px auto; border-left:4px solid var(--gaia-warning);">
            <span class="gaia-kicker" style="color:var(--gaia-warning);background:var(--gaia-warning-soft); border-color:#F0D08D;">Senha padrão</span>
            <div style="font-weight:700;color:var(--gaia-text);margin-top:4px;">Você está utilizando a senha padrão do administrador.</div>
            <div class="gaia-meta" style="margin-top:6px;">
                Por segurança, recomenda-se alterar essa senha. Não compartilhe esse acesso com terceiros.
                <br><br>
                Como o GAIA pode funcionar offline, a recuperação automática por e-mail ou SMS não foi implementada.
                Caso a senha seja esquecida e não exista outro administrador, será necessário um procedimento técnico local de restauração do acesso.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    c1, c2, c3 = st.columns([1, 1.25, 1])
    with c2:
        if st.button("🔐 Alterar senha agora", type="primary", use_container_width=True):
            st.session_state['_alterar_senha_origem_aviso'] = True
            nav('alterar_senha')
        if st.button("Continuar usando o sistema", type="secondary", use_container_width=True):
            marcar_aviso_senha_padrao_visualizado(st.session_state.user_id)
            set_flash('warning', 'Atenção: o sistema continuará usando a senha padrão do administrador.')
            nav('menu')


def tela_alterar_senha():
    if not st.session_state.get('user_id'):
        nav('login')
        return

    cabecalho()
    st.subheader("Alterar minha senha")

    origem_aviso = bool(st.session_state.get('_alterar_senha_origem_aviso'))
    if origem_aviso:
        st.info("Você acessou com a senha padrão. Cadastre uma nova senha para aumentar a segurança do sistema.")

    col_a, col_b, col_c = st.columns([1, 2.15, 1])
    with col_b:
        with st.form("form_alterar_minha_senha"):
            senha_atual = st.text_input("Senha atual*", type="password")
            nova_senha = st.text_input("Nova senha*", type="password")
            confirmar = st.text_input("Confirmar nova senha*", type="password")
            btn_salvar_enter = st.form_submit_button("SALVAR SENHA", type="primary", use_container_width=True, key="enter_default_alterar_senha")
            c1, c2 = st.columns(2)
            btn_cancelar = c1.form_submit_button("← Cancelar", type="secondary", use_container_width=True, key="fsb_cancelar_alterar_senha")
            btn_salvar_click = c2.form_submit_button("SALVAR SENHA", type="primary", use_container_width=True, key="btn_salvar_alterar_senha")
            btn_salvar = btn_salvar_enter or btn_salvar_click

    if btn_cancelar:
        if origem_aviso:
            marcar_aviso_senha_padrao_visualizado(st.session_state.user_id)
            st.session_state['_alterar_senha_origem_aviso'] = False
        nav('menu')

    if btn_salvar:
        with get_db() as conn:
            row = conn.execute("SELECT login, senha_hash FROM tb_usuario WHERE id=?", (st.session_state.user_id,)).fetchone()

        if not row:
            msg_error("Usuário não localizado.")
            return

        login_atual, senha_hash_atual = row
        if gerar_hash(senha_atual) != senha_hash_atual:
            msg_error("Senha atual incorreta.")
            return

        erros = validar_nova_senha(nova_senha, confirmar, login_atual)
        if erros:
            msg_error(" ".join(erros))
            return

        with get_db() as conn:
            conn.execute(
                """UPDATE tb_usuario
                      SET senha_hash=?, aviso_senha_padrao=0, aviso_senha_padrao_visualizado=1
                    WHERE id=?""",
                (gerar_hash(nova_senha), st.session_state.user_id)
            )
            conn.commit()

        st.session_state['_alterar_senha_origem_aviso'] = False
        set_flash("success", "Senha alterada com sucesso.")
        nav('menu')

def tela_menu():
    cabecalho()
    nome_usuario = html.escape(str(st.session_state.get('nome_usuario') or 'Usuário'))
    perfil = html.escape(str(st.session_state.get('user_role') or ''))

    st.markdown(
        f"""
        <div class="gaia-page-card gaia-workbench-card">
            <span class="gaia-kicker">Painel de trabalho</span>
            <div class="gaia-work-title">Olá, {nome_usuario}</div>
            <div class="gaia-work-subtitle">Perfil ativo: <strong>{perfil}</strong>. Escolha uma opção para iniciar.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown('<div class="gaia-section-heading">Ações principais</div>', unsafe_allow_html=True)
    ac1, ac2, ac3 = st.columns(3, gap="large")

    with ac1:
        if st.button(
            "🔄\nContinuar alteração\nBusque uma empresa cadastrada e siga para eventos e geração da minuta.",
            use_container_width=True,
            key="card_continuar_alt",
        ):
            nav('buscar_empresa')

    with ac2:
        if st.button(
            "🏢\nCadastrar empresa\nUse quando a LTDA ainda não existe no banco local do GAIA.",
            use_container_width=True,
            key="card_nova_empresa",
        ):
            nav('cad_empresa')

    with ac3:
        if st.button(
            "📂\nHistórico\nConsulte minutas e documentos gerados anteriormente.",
            use_container_width=True,
            key="card_historico_docs",
        ):
            nav('historico')

    if st.session_state.user_role == 'Admin':
        with st.expander("Administração", expanded=False):
            ad1, ad2, ad3 = st.columns(3)
            if ad1.button("Usuários", use_container_width=True, key="btn_usuarios_home"):
                nav('gestao_usuarios')
            if ad2.button("Minha senha", use_container_width=True, key="btn_senha_home"):
                nav('alterar_senha')
            if ad3.button("Sair", use_container_width=True, key="btn_logout_home"):
                for key in ['user_id', 'user_role', 'login_usuario', 'empresa_id', 'socio_id']:
                    st.session_state[key] = None
                st.session_state['_alterar_senha_origem_aviso'] = False
                nav('login')
    else:
        st.markdown('<div style="height:16px"></div>', unsafe_allow_html=True)
        if st.button("Sair", use_container_width=True, key="btn_logout_home"):
            for key in ['user_id', 'user_role', 'login_usuario', 'empresa_id', 'socio_id']:
                st.session_state[key] = None
            st.session_state['_alterar_senha_origem_aviso'] = False
            nav('login')

def tela_gestao_usuarios():
    if st.session_state.user_role != 'Admin':
        msg_error("Acesso Negado.")
        nav('menu')
        return

    cabecalho()
    st.subheader("👥 Gestão de Utilizadores")

    with get_db() as conn:
        usuarios = conn.execute("""
            SELECT u.id, u.nome, u.login, p.nome_perfil, u.id_perfil
            FROM tb_usuario u
            JOIN tb_perfil_acesso p ON u.id_perfil = p.id
            ORDER BY u.nome COLLATE NOCASE
        """).fetchall()

    # Redefinição de senha por administrador
    if st.session_state.get('_reset_senha_usuario'):
        uid, unome, ulogin = st.session_state['_reset_senha_usuario']
        st.markdown(
            f"""
            <div class="gaia-card" style="border-left:4px solid var(--gaia-info);">
                <span class="gaia-kicker" style="color:var(--gaia-info);background:var(--gaia-info-soft); border-color:#BCE0EE;">Redefinir senha</span>
                <div style="font-weight:700;color:var(--gaia-text);margin-top:4px;">Atualizar acesso do usuário</div>
                <div class="gaia-meta">Usuário: <strong>{html.escape(str(unome))}</strong><br>Login: <strong>{html.escape(str(ulogin))}</strong></div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        with st.form("form_reset_senha_usuario"):
            nova = st.text_input("Nova senha*", type="password")
            confirmar = st.text_input("Confirmar nova senha*", type="password")
            confirmar_enter = st.form_submit_button("Redefinir senha", type="primary", use_container_width=True, key="enter_default_reset_senha")
            ca, cb = st.columns(2)
            cancelar = ca.form_submit_button("Cancelar", type="secondary", use_container_width=True, key="fsb_cancelar_reset_senha")
            confirmar_click = cb.form_submit_button("Redefinir senha", type="primary", use_container_width=True, key="btn_confirmar_reset_senha")
            confirmar_btn = confirmar_enter or confirmar_click

        if cancelar:
            st.session_state.pop('_reset_senha_usuario', None)
            st.rerun()
        if confirmar_btn:
            erros = validar_nova_senha(nova, confirmar, ulogin)
            if erros:
                msg_error(" ".join(erros))
            else:
                with get_db() as conn_reset:
                    conn_reset.execute(
                        """UPDATE tb_usuario
                              SET senha_hash=?, aviso_senha_padrao=0, aviso_senha_padrao_visualizado=1
                            WHERE id=?""",
                        (gerar_hash(nova), uid)
                    )
                    conn_reset.commit()
                st.session_state.pop('_reset_senha_usuario', None)
                msg_success("Senha redefinida com sucesso.")
                st.rerun()
        st.markdown("<div class='gaia-separator'></div>", unsafe_allow_html=True)
        return

    # Confirmação de exclusão de usuário
    if st.session_state.get('_confirm_excluir_usuario'):
        uid, unome, ulogin = st.session_state['_confirm_excluir_usuario']
        st.markdown(
            f"""
            <div class="gaia-card" style="border-left:4px solid var(--gaia-danger);">
                <span class="gaia-kicker" style="color:var(--gaia-danger);background:var(--gaia-danger-soft); border-color:#F3C2BA;">Confirmação</span>
                <div style="font-weight:700;color:var(--gaia-text);margin-top:4px;">Excluir usuário</div>
                <div class="gaia-meta">Tem certeza que deseja excluir o usuário <strong>{html.escape(str(unome))}</strong>?
                <br>Login: <strong>{html.escape(str(ulogin))}</strong></div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        ca, cb = st.columns(2)
        if ca.button("Cancelar", type="secondary", use_container_width=True, key="cancel_del_user"):
            st.session_state.pop('_confirm_excluir_usuario', None)
            st.rerun()
        if cb.button("Confirmar exclusão", type="primary", use_container_width=True, key="confirm_del_user"):
            with get_db() as conn_del:
                conn_del.execute("DELETE FROM tb_usuario WHERE id=?", (uid,))
                conn_del.commit()
            st.session_state.pop('_confirm_excluir_usuario', None)
            msg_success("Usuário excluído com sucesso.")
            st.rerun()
        st.markdown("<div class='gaia-separator'></div>", unsafe_allow_html=True)
        return

    for u in usuarios:
        user_id, nome, login, perfil_nome, perfil_id = u
        with st.container():
            col1, col2, col3, col4 = st.columns([0.48, 0.24, 0.16, 0.12], vertical_alignment="center")
            col1.markdown(
                f"""
                <div class="gaia-card-compact" style="margin:0;">
                    <div class="gaia-socio-name">{html.escape(str(nome))}</div>
                    <div class="gaia-socio-meta">Login: {html.escape(str(login))} | Perfil atual: <strong>{html.escape(str(perfil_nome))}</strong></div>
                </div>
                """,
                unsafe_allow_html=True,
            )
            if user_id == st.session_state.user_id:
                col2.markdown("<div class='gaia-inline-note'>Usuário conectado</div>", unsafe_allow_html=True)
                if col3.button("🔐 Senha", key=f"senha_self_{user_id}", type="secondary", use_container_width=True):
                    nav('alterar_senha')
                col4.markdown("")
            else:
                if perfil_nome == 'Admin':
                    acao = "Tornar Operador"
                    novo_perfil_id = 2
                    key_acao = f"btn_rebaixar_operador_{user_id}"
                    msg = "Usuário alterado para Operador."
                else:
                    acao = "Promover a Admin"
                    novo_perfil_id = 1
                    key_acao = f"btn_promover_admin_{user_id}"
                    msg = "Usuário promovido a Admin."

                if col2.button(acao, key=key_acao, type="secondary", use_container_width=True):
                    with get_db() as conn_upd:
                        conn_upd.execute("UPDATE tb_usuario SET id_perfil=? WHERE id=?", (novo_perfil_id, user_id))
                        conn_upd.commit()
                    msg_success(msg)
                    st.rerun()

                if col3.button("🔐 Senha", key=f"reset_senha_{user_id}", type="secondary", use_container_width=True):
                    st.session_state['_reset_senha_usuario'] = (user_id, nome, login)
                    st.rerun()

                if col4.button("Excluir", key=f"del_user_{user_id}", type="secondary", use_container_width=True):
                    st.session_state['_confirm_excluir_usuario'] = (user_id, nome, login)
                    st.rerun()
            st.markdown("<div style='height:8px;'></div>", unsafe_allow_html=True)

    st.markdown("<div class='gaia-footer-actions'></div>", unsafe_allow_html=True)
    c1, c2, c3 = st.columns([1.4, 1, 1.4])
    if c2.button("← Voltar ao menu", type="secondary", use_container_width=True, key="voltar_gestao_usuarios"):
        nav('menu')

def tela_cad_empresa():
    cabecalho()
    st.subheader("DADOS DA EMPRESA")
    
    c_cnpj, c_btn = st.columns([3, 1], vertical_alignment="bottom")
    cnpj_busca = c_cnpj.text_input("Digite o CNPJ para buscar na Receita", placeholder="00.000.000/0000-00")
    st.markdown("<div class=\"gaia-inline-help\">A busca automática precisa de internet. Se não localizar o CNPJ, continue com o preenchimento manual.</div>", unsafe_allow_html=True)
    
    if c_btn.button("BUSCAR", use_container_width=True):
        if cnpj_busca:
            cnpj_limpo = somente_digitos(cnpj_busca)
            if len(cnpj_limpo) != 14:
                msg_warning("O CNPJ deve conter 14 dígitos.")
            elif not validar_cnpj(cnpj_limpo):
                msg_warning("CNPJ inválido. Confira os números informados.")
            elif cnpj_ja_cadastrado(cnpj_limpo):
                msg_warning("Empresa já cadastrada com este CNPJ. Use a opção Continuar alteração para acessar o cadastro existente.")
            else:
                with st.spinner("A consultar dados oficiais..."):
                    try:
                        resp = requests.get(f"https://brasilapi.com.br/api/cnpj/v1/{cnpj_limpo}", timeout=10)
                        if resp.status_code == 200:
                            d = resp.json()
                            st.session_state['api_cnpj'] = formatar_cnpj(cnpj_limpo)
                            st.session_state['api_nome'] = padronizar_texto(d.get('razao_social') or '')
                            
                            tipo_log = d.get('descricao_tipo_de_logradouro') or ''
                            log = d.get('logradouro') or ''
                            st.session_state['api_logradouro'] = padronizar_texto(f"{tipo_log} {log}".strip())
                            
                            st.session_state['api_numero'] = remover_espacos_extras(d.get('numero') or '')
                            st.session_state['api_complemento'] = padronizar_texto(d.get('complemento') or '')
                            st.session_state['api_bairro'] = padronizar_texto(d.get('bairro') or '')
                            st.session_state['api_cidade'] = padronizar_texto(d.get('municipio') or '')
                            st.session_state['api_uf'] = padronizar_texto(d.get('uf') or '')
                            st.session_state['api_cep'] = formatar_cep(d.get('cep') or '')
                            
                            qsa = d.get('qsa') or []
                            socios_encontrados = []
                            for s in qsa:
                                nome_socio = s.get("nome_socio")
                                if nome_socio:
                                    socios_encontrados.append({
                                        "nome": padronizar_texto(nome_socio),
                                        "cpf": formatar_cpf(s.get("cnpj_cpf_do_socio") or "")
                                    })
                            st.session_state['api_qsa'] = socios_encontrados
                            
                            if socios_encontrados:
                                set_flash("success", f"✅ Dados e {len(socios_encontrados)} sócio(s) encontrados!")
                            else:
                                set_flash("success", "✅ Dados encontrados! (Nota: Esta empresa não possui quadro de sócios na Receita).")
                            st.rerun()
                        else:
                            msg_error("CNPJ não encontrado na base de dados.")
                    except requests.exceptions.RequestException:
                        msg_api_offline("BrasilAPI/CNPJ")
                    except Exception as e:
                        msg_error(f"Erro inesperado ao consultar o CNPJ: {e}. Você pode preencher os dados manualmente.")
        else:
            msg_warning("Digite um CNPJ válido antes de buscar.")

    
    with st.form("f_emp"):
        cnpj = st.text_input("CNPJ*", value=formatar_cnpj(st.session_state.get('api_cnpj', '')), placeholder="00.000.000/0000-00")
        nome = st.text_input("NOME EMPRESARIAL*", value=padronizar_texto(st.session_state.get('api_nome', '')))
        nire = st.text_input(
            "NIRE (Apenas Manual)",
            help="O NIRE é o número de registro da empresa na Junta Comercial. Você pode encontrá-lo no contrato social, alteração contratual arquivada, ficha cadastral da Junta Comercial ou comprovante de registro."
        )
        st.markdown("**Morada da Sede***")
        
        c_cep_emp, c_btn_cep_emp = st.columns([3, 1], vertical_alignment="bottom")
        cep = c_cep_emp.text_input("CEP*", value=formatar_cep(st.session_state.get('api_cep', '')), placeholder="00000-000")
        btn_buscar_cep_emp = c_btn_cep_emp.form_submit_button("BUSCAR", type="secondary", use_container_width=True)
        
        if btn_buscar_cep_emp:
             cep_limpo = somente_digitos(cep)
             if len(cep_limpo) == 8:
                 try:
                     resp = requests.get(f"https://viacep.com.br/ws/{cep_limpo}/json/", timeout=5)
                     if resp.status_code == 200:
                         dados_cep = resp.json()
                         if "erro" not in dados_cep:
                             st.session_state['api_cep'] = formatar_cep(dados_cep.get('cep', ''))
                             st.session_state['api_logradouro'] = padronizar_texto(dados_cep.get('logradouro', ''))
                             st.session_state['api_bairro'] = padronizar_texto(dados_cep.get('bairro', ''))
                             st.session_state['api_cidade'] = padronizar_texto(dados_cep.get('localidade', ''))
                             st.session_state['api_uf'] = padronizar_texto(dados_cep.get('uf', ''))
                             st.rerun()
                         else:
                             msg_error("CEP não encontrado.")
                 except requests.exceptions.RequestException:
                     msg_api_offline("ViaCEP")
                 except Exception:
                     msg_error("Erro inesperado ao consultar o CEP. Você pode preencher os dados manualmente.")
             else:
                 msg_warning("O CEP deve conter 8 dígitos.")
        
        c7, c8 = st.columns([3, 1])
        logradouro = c7.text_input("LOGRADOURO*", value=padronizar_texto(st.session_state.get('api_logradouro', '')))
        numero = c8.text_input("NÚMERO*", value=remover_espacos_extras(st.session_state.get('api_numero', '')))
        c9, c10 = st.columns(2)
        complemento = c9.text_input("COMPLEMENTO", value=padronizar_texto(st.session_state.get('api_complemento', '')))
        bairro = c10.text_input("BAIRRO*", value=padronizar_texto(st.session_state.get('api_bairro', '')))
        c11, c12 = st.columns([3, 1])
        cidade = c11.text_input("CIDADE*", value=padronizar_texto(st.session_state.get('api_cidade', '')))
        uf = c12.text_input("UF*", value=padronizar_texto(st.session_state.get('api_uf', '')), max_chars=2)
        
        btn_seguir_enter = st.form_submit_button("Seguir e salvar →", type="primary", use_container_width=True, key="enter_default_cad_empresa")
        c1, c2 = st.columns(2)
        btn_voltar = c1.form_submit_button("← Voltar", type="secondary", use_container_width=True, key="fsb_voltar_cad_empresa")
        btn_seguir_click = c2.form_submit_button("Seguir e salvar →", type="primary", use_container_width=True, key="btn_seguir_cad_empresa")
        btn_seguir = btn_seguir_enter or btn_seguir_click

    if btn_voltar:
        for k in list(st.session_state.keys()):
            if k.startswith('api_'): del st.session_state[k]
        nav('menu')
        
    if btn_seguir:
        dados_empresa = {
            'cnpj': formatar_cnpj(cnpj),
            'nome_empresarial': padronizar_texto(nome),
            'nire': padronizar_texto(nire),
            'logradouro': padronizar_texto(logradouro),
            'numero': remover_espacos_extras(numero),
            'complemento': padronizar_texto(complemento),
            'bairro': padronizar_texto(bairro),
            'cidade': padronizar_texto(cidade),
            'uf': padronizar_texto(uf),
            'cep': formatar_cep(cep),
        }
        erros = validar_dados_empresa(dados_empresa)
        if erros:
            msg_error(f"Corrija antes de salvar: {', '.join(dict.fromkeys(erros))}.")
        else:
            try:
                with get_db() as conn:
                    cur = conn.cursor()
                    cur.execute("INSERT INTO tb_empresa (cnpj, nome_empresarial, nire, logradouro, numero, complemento, bairro, cidade, uf, cep) VALUES (?,?,?,?,?,?,?,?,?,?)",
                        (dados_empresa['cnpj'], dados_empresa['nome_empresarial'], dados_empresa['nire'], dados_empresa['logradouro'], dados_empresa['numero'], dados_empresa['complemento'], dados_empresa['bairro'], dados_empresa['cidade'], dados_empresa['uf'], dados_empresa['cep']))
                    st.session_state.empresa_id = cur.lastrowid
                    
                    qsa_api = st.session_state.get('api_qsa', [])
                    for s in qsa_api:
                        if s.get('nome'):
                            cur.execute("""INSERT INTO tb_socio 
                                (id_empresa, nome, cpf, nacionalidade, estado_civil, logradouro, numero, bairro, cidade, uf, cep) 
                                VALUES (?,?,?, 'BRASILEIRO(A)', '', '', '', '', '', '', '')""",
                                (st.session_state.empresa_id, padronizar_texto(s.get('nome')), formatar_cpf(s.get('cpf'))))
                    
                    conn.commit()
                
                for k in list(st.session_state.keys()):
                    if k.startswith('api_'): del st.session_state[k]
                    
                nav('lista_socios')
            except sqlite3.IntegrityError:
                msg_error(f"Já existe uma empresa com o CNPJ '{dados_empresa['cnpj']}'. Use 'Buscar empresa' para aceder à mesma.")

def tela_buscar_empresa():
    cabecalho()
    st.subheader("SELECIONE UMA EMPRESA")
    with get_db() as conn:
        empresas = conn.execute("SELECT id, cnpj, nome_empresarial FROM tb_empresa").fetchall()
    
    if not empresas:
        st.info("Nenhuma empresa cadastrada.")
        if st.button("← Voltar ao menu", type="secondary", key="voltar_buscar_empresa_vazio"): nav('menu')
    else:
        opcoes = {f"{e[1]} - {e[2]}": e[0] for e in empresas}
        
        escolha = st.selectbox("Empresas:", list(opcoes.keys()), index=None, placeholder="Selecione uma empresa na lista...")
        
        if st.button("Acessar empresa", type="primary", use_container_width=True, key="btn_acessar_empresa"):
            if escolha is None:
                msg_warning("Selecione uma empresa antes de avançar.")
            else:
                st.session_state.empresa_id = opcoes[escolha]
                nav('lista_socios')
                
        if st.button("← Voltar ao menu", type="secondary", use_container_width=True, key="voltar_buscar_empresa_com_lista"): nav('menu')

def tela_editar_empresa():
    cabecalho()
    st.subheader("EDITAR EMPRESA")
    is_admin = st.session_state.user_role == 'Admin'
    with get_db() as conn:
        e = conn.execute("SELECT cnpj, nome_empresarial, nire, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,)).fetchone()

    if not e:
        msg_error("Empresa não encontrada.")
        nav('menu')
        return

    if not is_admin: st.info("💡 Apenas Administradores podem forçar a edição manual dos dados da empresa.")

    with st.form("f_edit_emp"):
        cnpj = st.text_input("CNPJ*", value=formatar_cnpj(e[0] or ''), disabled=not is_admin, placeholder="00.000.000/0000-00")
        nome = st.text_input("NOME EMPRESARIAL*", value=padronizar_texto(e[1] or ''), disabled=not is_admin)
        nire = st.text_input(
            "NIRE",
            value=padronizar_texto(e[2] or ''),
            disabled=not is_admin,
            help="O NIRE é o número de registro da empresa na Junta Comercial. Você pode encontrá-lo no contrato social, alteração contratual arquivada, ficha cadastral da Junta Comercial ou comprovante de registro."
        )
        st.markdown("**Endereço**")
        c7, c8 = st.columns([3, 1])
        logradouro = c7.text_input("LOGRADOURO*", value=padronizar_texto(e[3] or ''), disabled=not is_admin)
        numero = c8.text_input("NÚMERO*", value=remover_espacos_extras(e[4] or ''), disabled=not is_admin)
        c9, c10 = st.columns(2)
        complemento = c9.text_input("COMPLEMENTO", value=padronizar_texto(e[5] or ''), disabled=not is_admin)
        bairro = c10.text_input("BAIRRO*", value=padronizar_texto(e[6] or ''), disabled=not is_admin)
        c11, c12, c13 = st.columns([3, 1, 2])
        cidade = c11.text_input("CIDADE*", value=padronizar_texto(e[7] or ''), disabled=not is_admin)
        uf = c12.text_input("UF*", value=padronizar_texto(e[8] or ''), disabled=not is_admin, max_chars=2)
        cep = c13.text_input("CEP*", value=formatar_cep(e[9] or ''), disabled=not is_admin, placeholder="00000-000")
        btn_ok_enter = st.form_submit_button("Salvar", type="primary", use_container_width=True, key="enter_default_editar_empresa") if is_admin else False
        c1, c2 = st.columns(2)
        btn_v = c1.form_submit_button("← Cancelar", type="secondary", use_container_width=True, key="fsb_cancelar_editar_empresa")
        btn_ok_click = c2.form_submit_button("Salvar", type="primary", use_container_width=True, key="btn_salvar_editar_empresa") if is_admin else False
        btn_ok = btn_ok_enter or btn_ok_click 

    if btn_v: nav('lista_socios')
    if btn_ok and is_admin:
        dados_empresa = {
            'cnpj': formatar_cnpj(cnpj),
            'nome_empresarial': padronizar_texto(nome),
            'nire': padronizar_texto(nire),
            'logradouro': padronizar_texto(logradouro),
            'numero': remover_espacos_extras(numero),
            'complemento': padronizar_texto(complemento),
            'bairro': padronizar_texto(bairro),
            'cidade': padronizar_texto(cidade),
            'uf': padronizar_texto(uf),
            'cep': formatar_cep(cep),
        }
        erros = validar_dados_empresa(dados_empresa, empresa_id_atual=st.session_state.empresa_id)
        if erros:
            msg_error(f"Corrija antes de salvar: {', '.join(dict.fromkeys(erros))}.")
            return
        with get_db() as conn:
            conn.execute("UPDATE tb_empresa SET cnpj=?, nome_empresarial=?, nire=?, logradouro=?, numero=?, complemento=?, bairro=?, cidade=?, uf=?, cep=? WHERE id=?",
                (dados_empresa['cnpj'], dados_empresa['nome_empresarial'], dados_empresa['nire'], dados_empresa['logradouro'], dados_empresa['numero'], dados_empresa['complemento'], dados_empresa['bairro'], dados_empresa['cidade'], dados_empresa['uf'], dados_empresa['cep'], st.session_state.empresa_id))
            conn.commit()
        set_flash("success", "Empresa atualizada com sucesso!")
        nav('lista_socios')

def tela_socios():
    cabecalho()
    is_admin = st.session_state.user_role == 'Admin'

    with get_db() as conn:
        emp = conn.execute("SELECT nome_empresarial FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,)).fetchone()
        socios = conn.execute("SELECT id, nome, cpf FROM tb_socio WHERE id_empresa=?", (st.session_state.empresa_id,)).fetchall()

    if not emp:
        msg_error("Empresa não encontrada.")
        nav('menu')
        return

    nome_empresa = html.escape(emp[0] or "EMPRESA NÃO INFORMADA")

    if st.session_state.get('_confirm_excluir_empresa') and is_admin:
        st.markdown(
            f"<div class='gaia-card-compact' style='border-left-color:var(--gaia-danger); background:var(--gaia-danger-soft);'>⚠️ Tem certeza que deseja excluir a empresa <strong>{nome_empresa}</strong> e todos os seus sócios? Esta ação não pode ser desfeita.</div>",
            unsafe_allow_html=True,
        )
        ca, cb = st.columns(2)
        if ca.button("Cancelar", type="secondary", use_container_width=True, key="cancel_excluir_empresa"):
            st.session_state.pop('_confirm_excluir_empresa', None)
            st.rerun()
        if cb.button("Confirmar exclusão", type="primary", use_container_width=True, key="confirm_del_empresa"):
            with get_db() as conn:
                conn.execute("DELETE FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,))
                conn.commit()
            st.session_state.empresa_id = None
            st.session_state.pop('_confirm_excluir_empresa', None)
            set_flash('success', 'Empresa excluída com sucesso.')
            nav('menu')
        return

    if st.session_state.get('_confirm_excluir_socio') and is_admin:
        sid, snome = st.session_state['_confirm_excluir_socio']
        st.markdown(
            f"<div class='gaia-card-compact' style='border-left-color:var(--gaia-danger); background:var(--gaia-danger-soft);'>⚠️ Tem certeza que deseja excluir o sócio <strong>{html.escape(str(snome))}</strong>?</div>",
            unsafe_allow_html=True,
        )
        ca, cb = st.columns(2)
        if ca.button("Cancelar", type="secondary", use_container_width=True, key="cancel_excluir_socio"):
            st.session_state.pop('_confirm_excluir_socio', None)
            st.rerun()
        if cb.button("Confirmar exclusão", type="primary", use_container_width=True, key=f"confirm_del_socio_{sid}"):
            with get_db() as conn:
                conn.execute("DELETE FROM tb_socio WHERE id=?", (sid,))
                conn.commit()
            st.session_state.pop('_confirm_excluir_socio', None)
            set_flash('success', 'Sócio excluído com sucesso.')
            st.rerun()
        return

    # A própria caixa da empresa funciona como botão para abrir/ver os dados.
    # O botão de exclusão fica ao lado, no mesmo padrão das ações do quadro societário.
    texto_empresa_card = (
        "EMPRESA SELECIONADA\n"
        f"{emp[0] or 'EMPRESA NÃO INFORMADA'}\n"
        "Clique para visualizar ou editar os dados cadastrais."
    )
    if is_admin:
        col_emp, col_ex = st.columns([5.2, 1.05], vertical_alignment="center")
    else:
        col_emp, = st.columns([1])

    with col_emp:
        if st.button(texto_empresa_card, key="empresa_card_click", type="secondary", use_container_width=True):
            nav('editar_empresa')

    if is_admin:
        if col_ex.button("Excluir", key="ex_emp", type="secondary", use_container_width=True):
            st.session_state['_confirm_excluir_empresa'] = True
            st.rerun()

    st.markdown("<div class='gaia-section-title'>Quadro Societário</div>", unsafe_allow_html=True)

    if not socios:
        st.markdown("<div class='gaia-card-compact' style='text-align:center;'>Nenhum sócio cadastrado para esta empresa.</div>", unsafe_allow_html=True)

    for s in socios:
        socio_id, socio_nome, socio_cpf = s
        socio_nome_html = html.escape(socio_nome or "SÓCIO NÃO INFORMADO")
        socio_cpf_html = html.escape(socio_cpf or "Não informado")
        col_info, col_edit, col_del = st.columns([5.2, 1.05, 1.05], vertical_alignment="center")
        with col_info:
            st.markdown(
                f"""
                <div class="gaia-card-compact" style="margin: 3px 0;">
                    <p class="gaia-socio-name">{socio_nome_html}</p>
                    <div class="gaia-socio-meta">CPF: {socio_cpf_html}</div>
                </div>
                """,
                unsafe_allow_html=True,
            )
        if col_edit.button("Editar", key=f"ed_s_{socio_id}", type="secondary", use_container_width=True):
            st.session_state.socio_id = socio_id
            nav('editar_socio')
        if is_admin:
            if col_del.button("Excluir", key=f"ex_s_{socio_id}", type="secondary", use_container_width=True):
                st.session_state['_confirm_excluir_socio'] = (socio_id, socio_nome)
                st.rerun()
        else:
            col_del.markdown("")

    c_n1, c_n2, c_n3 = st.columns([2.2, 1.6, 2.2])
    with c_n2:
        if st.button("Novo sócio", type="secondary", use_container_width=True, key="btn_novo_socio"): nav('cad_socio')

    st.markdown("<div class='gaia-bottom-section'><div class='gaia-bottom-title'>Documentos e continuidade</div></div>", unsafe_allow_html=True)
    b1, b2, b3 = st.columns(3)
    if b1.button("Histórico da empresa", type="secondary", use_container_width=True):
        nav('historico_empresa')
    if b2.button("← Voltar ao menu", type="secondary", use_container_width=True, key="voltar_lista_socios"):
        st.session_state.empresa_id = None
        nav('menu')
    if b3.button("Avançar para seleção →", type="primary", use_container_width=True, key="btn_avancar_para_selecao"):
        nav('orquestrador')

def tela_cad_socio():
    btn_v, btn_ok, dados = _form_socio("CADASTRAR NOVO SÓCIO")
    if btn_v: nav('lista_socios')
    if btn_ok:
        erros = validar_dados_socio(dados)
        if erros:
            msg_error(f"Corrija antes de salvar: {', '.join(dict.fromkeys(erros))}.")
            return
        with get_db() as conn:
            conn.execute("""INSERT INTO tb_socio (id_empresa, nome, nacionalidade, cpf, rg, orgao_emissor, estado_civil, data_nasc, profissao, logradouro, numero, complemento, bairro, cidade, uf, cep) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (st.session_state.empresa_id, dados['nome'], dados['nacionalidade'], dados['cpf'], dados['rg'], dados['orgao_emissor'], dados['estado_civil'], dados['data_nasc'], dados['profissao'], dados['logradouro'], dados['numero'], dados['complemento'], dados['bairro'], dados['cidade'], dados['uf'], dados['cep']))
            conn.commit()
        set_flash("success", f"Sócio **{dados['nome']}** cadastrado com sucesso!")
        nav('lista_socios')

def tela_editar_socio():
    with get_db() as conn:
        s = conn.execute("SELECT nome, nacionalidade, cpf, rg, orgao_emissor, estado_civil, data_nasc, profissao, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_socio WHERE id=?", (st.session_state.socio_id,)).fetchone()
    if not s:
        msg_error("Sócio não encontrado.")
        nav('lista_socios')
        return
    defaults = { 'nome': s[0], 'nacionalidade': s[1], 'cpf': s[2], 'rg': s[3], 'orgao_emissor': s[4], 'estado_civil': s[5], 'data_nasc': s[6], 'profissao': s[7], 'logradouro': s[8], 'numero': s[9], 'complemento': s[10], 'bairro': s[11], 'cidade': s[12], 'uf': s[13], 'cep': s[14] }
    btn_v, btn_ok, dados = _form_socio("EDITAR SÓCIO", defaults)
    if btn_v: nav('lista_socios')
    if btn_ok:
        erros = validar_dados_socio(dados)
        if erros:
            msg_error(f"Corrija antes de salvar: {', '.join(dict.fromkeys(erros))}.")
            return
        with get_db() as conn:
            conn.execute("""UPDATE tb_socio SET nome=?, nacionalidade=?, cpf=?, rg=?, orgao_emissor=?, estado_civil=?, data_nasc=?, profissao=?, logradouro=?, numero=?, complemento=?, bairro=?, cidade=?, uf=?, cep=? WHERE id=?""",
                (dados['nome'], dados['nacionalidade'], dados['cpf'], dados['rg'], dados['orgao_emissor'], dados['estado_civil'], dados['data_nasc'], dados['profissao'], dados['logradouro'], dados['numero'], dados['complemento'], dados['bairro'], dados['cidade'], dados['uf'], dados['cep'], st.session_state.socio_id))
            conn.commit()
        set_flash("success", "Sócio atualizado com sucesso!")
        nav('lista_socios')

def resumo_eventos(eventos_json):
    try:
        eventos = json.loads(eventos_json or "[]")
        nomes = {
            "220": "Alteração de nome empresarial",
            "221": "Alteração de nome fantasia",
            "211": "Endereço no mesmo município",
            "209": "Endereço entre municípios",
            "210": "Endereço entre estados",
            "248": "Tipo de unidade",
            "244": "Atividades econômicas/CNAEs",
            "249": "Forma de atuação",
            "225": "Transformação",
            "052": "Reativação",
            "999": "Licenciamento legado",
            "capital": "Capital social",
        }
        tipos = [nomes.get(e.get("tipo", ""), e.get("tipo", "")) for e in eventos if isinstance(e, dict)]
        return ", ".join(tipos) if tipos else "Não informado"
    except Exception:
        return "Não informado"


def _json_pretty(valor_json, fallback="Não informado"):
    try:
        dados = json.loads(valor_json or "[]")
        if not dados:
            return fallback
        return json.dumps(dados, ensure_ascii=False, indent=2)
    except Exception:
        return fallback


def excluir_documento_historico(doc_id):
    """Exclui documento do histórico. Ação restrita ao Admin pela interface."""
    with get_db() as conn:
        conn.execute("DELETE FROM tb_documento WHERE id=?", (doc_id,))
        conn.commit()


def atualizar_status_documento(doc_id, status, protocolo=None, data_protocolo=None, observacao=None):
    with get_db() as conn:
        conn.execute(
            """UPDATE tb_documento
               SET status=?, protocolo=?, data_protocolo=?, observacao=?
               WHERE id=?""",
            (
                remover_espacos_extras(status),
                remover_espacos_extras(protocolo),
                remover_espacos_extras(data_protocolo),
                remover_espacos_extras(observacao),
                doc_id,
            ),
        )
        conn.commit()


def anexar_comprovante_documento(doc_id, arquivo):
    if not arquivo:
        return
    with get_db() as conn:
        conn.execute(
            "UPDATE tb_documento SET comprovante_nome=?, comprovante_blob=? WHERE id=?",
            (arquivo.name, sqlite3.Binary(arquivo.getvalue()), doc_id),
        )
        conn.commit()


def limpar_comprovante_documento(doc_id):
    with get_db() as conn:
        conn.execute(
            "UPDATE tb_documento SET comprovante_nome=NULL, comprovante_blob=NULL WHERE id=?",
            (doc_id,),
        )
        conn.commit()


def _opcoes_status_documento(status_atual):
    opcoes = [
        "Gerado",
        "Revisado",
        "Protocolado",
        "Arquivado",
        "Cancelado",
        "Gerado parcialmente",
        "Gerado em DOCX e PDF",
    ]
    status_atual = status_atual or "Gerado"
    if status_atual not in opcoes:
        opcoes.insert(0, status_atual)
    return opcoes


def renderizar_card_documento(doc, prefixo="geral"):
    (
        doc_id, data_geracao, data_documento_texto, tipo_documento, empresa_nome, usuario_nome,
        status, eventos_json, nome_docx, nome_pdf, docx_blob, pdf_blob, observacao,
        modelo_documento, nome_relatorio, relatorio_blob, protocolo, data_protocolo,
        comprovante_nome, comprovante_blob
    ) = doc

    status_exibicao = status or "Gerado"
    cor_status = {
        "Gerado": ("var(--gaia-success-soft)", "var(--gaia-success)"),
        "Gerado em DOCX e PDF": ("var(--gaia-success-soft)", "var(--gaia-success)"),
        "Revisado": ("var(--gaia-info-soft)", "var(--gaia-info)"),
        "Protocolado": ("var(--gaia-purple-soft)", "var(--gaia-purple)"),
        "Arquivado": ("var(--gaia-neutral-soft)", "var(--gaia-neutral)"),
        "Cancelado": ("var(--gaia-danger-soft)", "var(--gaia-danger)"),
        "Gerado parcialmente": ("var(--gaia-warning-soft)", "var(--gaia-warning)"),
    }.get(status_exibicao, ("var(--gaia-success-soft)", "var(--gaia-success)"))
    bg_status, fg_status = cor_status

    with st.container():
        st.markdown(
            f"""
            <div class="gaia-card-compact" style="margin-bottom:8px;">
                <div style="display:flex; gap:8px; align-items:center; justify-content:space-between; flex-wrap:wrap;">
                    <div>
                        <div class="gaia-socio-name">{html.escape(str(empresa_nome or 'Empresa não informada'))}</div>
                        <div class="gaia-socio-meta">
                            {html.escape(str(tipo_documento or 'Documento'))}
                            | Gerado por: {html.escape(str(usuario_nome or 'Usuário não informado'))}
                        </div>
                        <div class="gaia-socio-meta">
                            Geração: {html.escape(str(data_geracao or ''))}
                            {f' | Data do documento: {html.escape(str(data_documento_texto))}' if data_documento_texto else ''}
                        </div>
                    </div>
                    <div class="gaia-status-pill" style="background:{bg_status}; color:{fg_status};">
                        {html.escape(status_exibicao)}
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        st.caption(f"Eventos registrados: {resumo_eventos(eventos_json)}")
        if protocolo or data_protocolo:
            st.caption(f"Protocolo: {protocolo or 'Não informado'} | Data do protocolo: {data_protocolo or 'Não informada'}")
        if observacao:
            st.caption(f"Observação: {observacao}")

        botoes = st.columns(3)
        if docx_blob:
            botoes[0].download_button(
                "⬇️ 📘 WORD (DOCX)",
                data=docx_blob,
                file_name=nome_docx or f"documento_{doc_id}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"{prefixo}_docx_{doc_id}",
                use_container_width=True,
            )
        else:
            botoes[0].button("📘 WORD (DOCX) indisponível", disabled=True, use_container_width=True, key=f"{prefixo}_docx_off_{doc_id}")

        if pdf_blob:
            botoes[1].download_button(
                "⬇️ 📕 PDF",
                data=pdf_blob,
                file_name=nome_pdf or f"documento_{doc_id}.pdf",
                mime="application/pdf",
                key=f"{prefixo}_pdf_{doc_id}",
                use_container_width=True,
            )
        else:
            botoes[1].button("📕 PDF indisponível", disabled=True, use_container_width=True, key=f"{prefixo}_pdf_off_{doc_id}")

        if comprovante_blob:
            botoes[2].download_button(
                "Comprovante",
                data=comprovante_blob,
                file_name=comprovante_nome or f"comprovante_{doc_id}",
                mime="application/octet-stream",
                key=f"{prefixo}_comp_{doc_id}",
                use_container_width=True,
            )
        else:
            botoes[2].button("Sem comprovante", disabled=True, use_container_width=True, key=f"{prefixo}_comp_off_{doc_id}")

        with st.expander("Acompanhamento do processo", expanded=False):
            c_status, c_prot, c_data = st.columns([1.2, 1.2, 1])
            opcoes_status = _opcoes_status_documento(status_exibicao)
            novo_status = c_status.selectbox(
                "Status",
                opcoes_status,
                index=opcoes_status.index(status_exibicao) if status_exibicao in opcoes_status else 0,
                key=f"{prefixo}_status_{doc_id}",
            )
            novo_protocolo = c_prot.text_input("Nº protocolo", value=protocolo or "", key=f"{prefixo}_protocolo_{doc_id}")
            nova_data_protocolo = c_data.text_input("Data protocolo", value=data_protocolo or "", placeholder="DD/MM/AAAA", key=f"{prefixo}_data_prot_{doc_id}")
            nova_observacao = st.text_area("Observação interna", value=observacao or "", key=f"{prefixo}_obs_{doc_id}")

            up1, up2 = st.columns([2, 1])
            arquivo = up1.file_uploader(
                "Anexar/atualizar comprovante de protocolo",
                type=["pdf", "png", "jpg", "jpeg", "doc", "docx"],
                key=f"{prefixo}_upload_comp_{doc_id}",
            )
            if up2.button("💾 Salvar controle", type="primary", use_container_width=True, key=f"{prefixo}_salvar_ctrl_{doc_id}"):
                if nova_data_protocolo:
                    data_ok, data_fmt = normalizar_data_br(nova_data_protocolo)
                    if not data_ok:
                        msg_error("Data de protocolo inválida. Use DD/MM/AAAA.")
                        return
                    nova_data_protocolo = data_fmt
                atualizar_status_documento(doc_id, novo_status, novo_protocolo, nova_data_protocolo, nova_observacao)
                if arquivo:
                    anexar_comprovante_documento(doc_id, arquivo)
                set_flash("success", "Histórico atualizado com sucesso.")
                st.rerun()

            c_limpar, c_eventos = st.columns(2)
            if comprovante_blob and c_limpar.button("Remover comprovante", type="secondary", use_container_width=True, key=f"{prefixo}_limpar_comp_{doc_id}"):
                limpar_comprovante_documento(doc_id)
                set_flash("success", "Comprovante removido.")
                st.rerun()

            with c_eventos:
                if st.button("Ver eventos/cláusulas", type="secondary", use_container_width=True, key=f"{prefixo}_ver_json_{doc_id}"):
                    st.session_state[f"{prefixo}_mostrar_json_{doc_id}"] = not st.session_state.get(f"{prefixo}_mostrar_json_{doc_id}", False)

            if st.session_state.get(f"{prefixo}_mostrar_json_{doc_id}"):
                st.markdown("**Eventos JSON:**")
                st.code(_json_pretty(eventos_json), language="json")

        if st.session_state.user_role == 'Admin':
            if st.session_state.get('_confirm_excluir_documento') == doc_id:
                st.warning("Confirma excluir este documento do histórico? Esta ação remove DOCX, PDF e comprovante salvos.")
                ca, cb = st.columns(2)
                if ca.button("Cancelar exclusão", type="secondary", use_container_width=True, key=f"{prefixo}_cancel_del_doc_{doc_id}"):
                    st.session_state['_confirm_excluir_documento'] = None
                    st.rerun()
                if cb.button("Confirmar exclusão", type="primary", use_container_width=True, key=f"{prefixo}_confirm_del_doc_{doc_id}"):
                    excluir_documento_historico(doc_id)
                    st.session_state['_confirm_excluir_documento'] = None
                    set_flash("success", "Documento excluído do histórico.")
                    st.rerun()
            else:
                if st.button("Excluir documento", type="secondary", use_container_width=True, key=f"{prefixo}_del_doc_{doc_id}"):
                    st.session_state['_confirm_excluir_documento'] = doc_id
                    st.rerun()

        st.divider()


def _montar_filtros_historico(escopo_empresa_id=None):
    with get_db() as conn:
        empresas = conn.execute("SELECT id, nome_empresarial FROM tb_empresa ORDER BY nome_empresarial COLLATE NOCASE").fetchall()
        usuarios = conn.execute("SELECT id, nome FROM tb_usuario ORDER BY nome COLLATE NOCASE").fetchall()
        tipos = [r[0] for r in conn.execute("SELECT DISTINCT tipo_documento FROM tb_documento WHERE tipo_documento IS NOT NULL ORDER BY tipo_documento").fetchall()]
        status_lista = [r[0] for r in conn.execute("SELECT DISTINCT status FROM tb_documento WHERE status IS NOT NULL ORDER BY status").fetchall()]

    with st.expander("Filtros do histórico", expanded=True):
        c1, c2, c3 = st.columns(3)
        empresa_escolhida = None
        if escopo_empresa_id is None:
            mapa_empresas = {"Todas": None}
            mapa_empresas.update({nome or f"Empresa {eid}": eid for eid, nome in empresas})
            empresa_nome = c1.selectbox("Empresa", list(mapa_empresas.keys()), key="hist_filtro_empresa")
            empresa_escolhida = mapa_empresas[empresa_nome]
        else:
            empresa_escolhida = escopo_empresa_id
            c1.markdown("**Empresa:** filtro aplicado")

        mapa_usuarios = {"Todos": None}
        mapa_usuarios.update({nome or f"Usuário {uid}": uid for uid, nome in usuarios})
        usuario_nome = c2.selectbox("Usuário", list(mapa_usuarios.keys()), key=f"hist_filtro_usuario_{escopo_empresa_id or 'geral'}")
        usuario_escolhido = mapa_usuarios[usuario_nome]

        tipo_doc = c3.selectbox("Tipo", ["Todos"] + tipos, key=f"hist_filtro_tipo_{escopo_empresa_id or 'geral'}")

        c4, c5, c6 = st.columns(3)
        status_doc = c4.selectbox("Status", ["Todos"] + status_lista, key=f"hist_filtro_status_{escopo_empresa_id or 'geral'}")
        data_ini = c5.date_input("Gerado a partir de", value=None, format="DD/MM/YYYY", key=f"hist_filtro_ini_{escopo_empresa_id or 'geral'}")
        data_fim = c6.date_input("Gerado até", value=None, format="DD/MM/YYYY", key=f"hist_filtro_fim_{escopo_empresa_id or 'geral'}")

    where = []
    params = []
    if empresa_escolhida:
        where.append("d.id_empresa=?")
        params.append(empresa_escolhida)
    if usuario_escolhido:
        where.append("d.id_usuario=?")
        params.append(usuario_escolhido)
    if tipo_doc != "Todos":
        where.append("d.tipo_documento=?")
        params.append(tipo_doc)
    if status_doc != "Todos":
        where.append("d.status=?")
        params.append(status_doc)
    if data_ini:
        where.append("date(d.data_geracao) >= date(?)")
        params.append(data_ini.isoformat())
    if data_fim:
        where.append("date(d.data_geracao) <= date(?)")
        params.append(data_fim.isoformat())

    return where, params


def _buscar_documentos_historico(where, params):
    sql = """
        SELECT d.id, d.data_geracao, d.data_documento_texto, d.tipo_documento, e.nome_empresarial, u.nome, d.status,
               d.eventos_json, d.nome_arquivo_docx, d.nome_arquivo_pdf, d.docx_blob, d.pdf_blob, d.observacao,
               d.modelo_documento, d.nome_arquivo_relatorio, d.relatorio_conferencia_blob, d.protocolo, d.data_protocolo,
               d.comprovante_nome, d.comprovante_blob
        FROM tb_documento d
        JOIN tb_empresa e ON d.id_empresa = e.id
        LEFT JOIN tb_usuario u ON d.id_usuario = u.id
    """
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY d.data_geracao DESC, d.id DESC"
    with get_db() as conn:
        return conn.execute(sql, params).fetchall()


def tela_historico():
    cabecalho()
    st.subheader("📄 Histórico Geral de Documentos")
    where, params = _montar_filtros_historico()
    docs = _buscar_documentos_historico(where, params)
    st.caption(f"Total encontrado: {len(docs)} documento(s).")
    if not docs:
        st.info("Nenhum documento encontrado para os filtros selecionados.")
    else:
        for doc in docs:
            renderizar_card_documento(doc, prefixo="hist_geral")
    if st.button("← Voltar ao menu", type="secondary", key="voltar_historico_geral"):
        nav('menu')


def tela_historico_empresa():
    cabecalho()
    st.subheader("📄 Histórico da Empresa")
    with get_db() as conn:
        emp = conn.execute("SELECT nome_empresarial FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,)).fetchone()
    st.markdown(f"**Empresa:** {emp[0] if emp else 'Empresa não localizada'}")
    where, params = _montar_filtros_historico(escopo_empresa_id=st.session_state.empresa_id)
    docs = _buscar_documentos_historico(where, params)
    st.caption(f"Total encontrado: {len(docs)} documento(s).")
    if not docs:
        st.info("Nenhum documento encontrado para esta empresa nos filtros selecionados.")
    else:
        for doc in docs:
            renderizar_card_documento(doc, prefixo="hist_empresa")
    if st.button("← Voltar", type="secondary", key="voltar_historico_empresa"):
        nav('lista_socios')

# ======================================================================
# ARQUITETURA PROTOTYPE + CÓDIGOS DBE/JUCEC
# ======================================================================
class ClausulaPrototipo:
    def clone(self): return copy.deepcopy(self)
    def aplicar_dados(self, dados): pass

class ClausulaTransformacao(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Transformar o tipo jurídico para Sociedade Empresária Limitada, adotando o nome empresarial {novo_nome}."
    def aplicar_dados(self, dados): return self.texto_base.format(novo_nome=dados.get("novo_nome", ""))

class ClausulaNome(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passará a denominar-se: {novo_nome}."
    def aplicar_dados(self, dados): return self.texto_base.format(novo_nome=dados.get("novo_nome", ""))

class ClausulaNomeFantasia(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passará a ter o TÍTULO DO ESTABELECIMENTO (NOME FANTASIA): {nome_fantasia}."
    def aplicar_dados(self, dados): return self.texto_base.format(nome_fantasia=dados.get("nome_fantasia", ""))

class ClausulaEndereco(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Alterar o endereço da sociedade, que passa a localizar-se na(o) {logradouro}, número {numero}{complemento}, bairro {bairro}, município {cidade} - {uf}, CEP: {cep}."
    def aplicar_dados(self, dados):
        comp = f", complemento {dados.get('complemento')}" if dados.get("complemento") else ""
        return self.texto_base.format(
            logradouro=dados.get("logradouro", ""), numero=dados.get("numero", ""),
            complemento=comp, bairro=dados.get("bairro", ""),
            cidade=dados.get("cidade", ""), uf=dados.get("uf", ""), cep=dados.get("cep", "")
        )

class ClausulaTipoUnidade(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Fica alterado o tipo de unidade do estabelecimento, passando a ser classificado como: {tipo_unidade}."
    def aplicar_dados(self, dados): return self.texto_base.format(tipo_unidade=dados.get("tipo_unidade", ""))

class ClausulaAtividadesEcono(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passa a ter por objeto, o exercício das seguintes atividades econômicas: {objeto_descricao}. \nA sociedade passa a ter as seguintes atividades econômicas (CNAES) vinculadas ao seu objeto social: {cnaes_lista}."
    def aplicar_dados(self, dados):
        return self.texto_base.format(
            objeto_descricao=dados.get("objeto_descricao", ""),
            cnaes_lista=dados.get("cnaes_lista", "")
        )

class ClausulaFormaAtuacao(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passará a exercer suas atividades sob a seguinte forma de atuação: {forma_atuacao}."
    def aplicar_dados(self, dados): return self.texto_base.format(forma_atuacao=dados.get("forma_atuacao", ""))

class ClausulaReativacao(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade, que se encontrava com seu registro inativo, reativa neste ato as suas atividades, em conformidade com o Artigo 60 da Lei 8.934/94."
    def aplicar_dados(self, dados): return self.texto_base

class ClausulaLicenciamento(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Fica registrado o licenciamento do estabelecimento anteriormente registrado (Legado), para todos os efeitos legais e administrativos inerentes."
    def aplicar_dados(self, dados): return self.texto_base

class ClausulaCapital(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "O capital social é de R$ {novo_capital}, dividido em quotas de valor nominal de R$ 1,00 cada, totalmente subscrito e integralizado neste ato."
    def aplicar_dados(self, dados): return self.texto_base.format(novo_capital=dados.get("novo_capital", ""))

CATALOGO_CLAUSULAS = {
    "225": ClausulaTransformacao(), "220": ClausulaNome(), "221": ClausulaNomeFantasia(),
    "211": ClausulaEndereco(), "209": ClausulaEndereco(), "210": ClausulaEndereco(),
    "248": ClausulaTipoUnidade(), "244": ClausulaAtividadesEcono(), "249": ClausulaFormaAtuacao(),
    "052": ClausulaReativacao(), "999": ClausulaLicenciamento(), "capital": ClausulaCapital()
}

# ======================================================================
# APOIO PARA CNAE - BASE LOCAL OFICIAL
# ======================================================================
STOPWORDS_BUSCA = {
    "a", "as", "o", "os", "um", "uma", "uns", "umas", "de", "da", "das", "do", "dos",
    "e", "em", "no", "na", "nos", "nas", "para", "por", "com", "sem", "ao", "aos",
    "que", "quero", "queria", "fazer", "trabalhar", "abrir", "empresa", "atividade",
    "atividades", "servico", "servicos", "produto", "produtos",
}

TERMOS_REMOCAO_OBJETO = [
    r"\bn[aã]o especificad[oa]s? anteriormente\b",
    r"\boutros?\b",
    r"\boutras?\b",
    r"\bdemais\b",
    r"\betc\.?\b",
]

def _normalizar_busca(txt):
    txt = str(txt or "").lower()
    txt = unicodedata.normalize("NFD", txt)
    txt = "".join(ch for ch in txt if unicodedata.category(ch) != "Mn")
    txt = re.sub(r"[^a-z0-9/\-\s]", " ", txt)
    return re.sub(r"\s+", " ", txt).strip()

def _tokens_busca(txt):
    return [t for t in _normalizar_busca(txt).split() if len(t) >= 3 and t not in STOPWORDS_BUSCA]

def _tem_termo(texto_normalizado, termos):
    return any(re.search(rf"\b{re.escape(termo)}\b", texto_normalizado) for termo in termos)

def _intencao_varejo(texto_normalizado):
    return _tem_termo(texto_normalizado, [
        "venda", "vender", "vende", "loja", "comercio", "comercializacao",
        "revenda", "varejo", "varejista", "consumidor"
    ])

def _intencao_atacado(texto_normalizado):
    return _tem_termo(texto_normalizado, [
        "atacado", "atacadista", "distribuidora", "distribuicao",
        "fornecedor", "representante", "representacao"
    ])

def _termo_bloqueado_cnae(texto_normalizado):
    return _tem_termo(texto_normalizado, [
        "droga", "drogas", "entorpecente", "entorpecentes",
        "narcotico", "narcoticos", "maconha", "cocaina", "crack"
    ])

def _caminhos_cnaes():
    base = Path(__file__).resolve().parent
    return [
        base / "dados" / "cnaes.csv",
        base / "cnaes.csv",
        Path.cwd() / "dados" / "cnaes.csv",
        Path.cwd() / "cnaes.csv",
    ]

@st.cache_data(show_spinner=False)
def carregar_cnaes_local():
    """Carrega a base CNAE local em CSV.

    O arquivo esperado é dados/cnaes.csv com as colunas:
    codigo;descricao_oficial
    """
    caminho = next((p for p in _caminhos_cnaes() if p.exists()), None)
    if not caminho:
        return []

    with open(caminho, "r", encoding="utf-8-sig", newline="") as f:
        amostra = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(amostra, delimiters=";,")
            separador = dialect.delimiter
        except csv.Error:
            separador = ";"

        reader = csv.DictReader(f, delimiter=separador)
        cnaes = []
        for row in reader:
            codigo = (row.get("codigo") or row.get("código") or "").strip()
            descricao = (row.get("descricao_oficial") or row.get("descrição final") or row.get("descricao") or "").strip()
            if not codigo or not descricao:
                continue
            if not re.match(r"^\d{4}-\d/\d{2}$", codigo):
                continue
            cnaes.append({
                "codigo": codigo,
                "descricao": descricao,
                "busca": _normalizar_busca(f"{codigo} {descricao}"),
            })
    return cnaes

MODELO_EMBEDDING_CNAE = os.getenv("GAIA_EMBEDDING_MODEL", "nomic-embed-text")
MODELO_EMBEDDING_CNAE_FALLBACK = os.getenv("GAIA_EMBEDDING_MODEL_FALLBACK", "bge-m3")
OLLAMA_BASE_URL = "http://localhost:11434"
OLLAMA_EMBED_URL = f"{OLLAMA_BASE_URL}/api/embed"

# Nesta versão, a "IA local" do evento CNAE é o modelo de embeddings rodando no Ollama.
# Não há chamada para modelo generativo, não há prompt livre e não há regras por exemplo.
# A busca inteligente combina: BM25 local + embedding local + fusão de ranking.

NIVEIS_EMBEDDING_CNAE = {
    "Preciso": {
        "limite_embedding": 24,
        "limite_texto": 35,
        "limite_bm25": 35,
        "peso_embedding": 0.86,
        "peso_texto": 1.10,
        "peso_bm25": 1.25,
        "bonus_convergencia": 20.0,
        "limite_final": 8,
        "descricao": "Prioriza resultados mais próximos da atividade informada.",
    },
    "Equilibrado": {
        "limite_embedding": 36,
        "limite_texto": 40,
        "limite_bm25": 40,
        "peso_embedding": 1.00,
        "peso_texto": 1.05,
        "peso_bm25": 1.15,
        "bonus_convergencia": 15.0,
        "limite_final": 10,
        "descricao": "Equilibra precisão e variedade de sugestões oficiais.",
    },
    "Amplo": {
        "limite_embedding": 52,
        "limite_texto": 45,
        "limite_bm25": 45,
        "peso_embedding": 1.18,
        "peso_texto": 0.98,
        "peso_bm25": 1.05,
        "bonus_convergencia": 10.0,
        "limite_final": 12,
        "descricao": "Mostra mais possibilidades quando a descrição é curta, popular ou ambígua.",
    },
}

STOPWORDS_BM25_CNAE = STOPWORDS_BUSCA | {
    "cnae", "cnpj", "mei", "ltda", "eireli", "simples", "nacional",
    "comercio", "comercial", "comercializacao", "varejo", "varejista",
    "atacado", "atacadista", "venda", "vender", "vende", "vendendo",
    "loja", "negocio", "ramo", "area", "oficial", "base", "afins",
    "semelhante", "semelhantes", "geral", "gerais", "nao", "especificado",
    "especificada", "especificados", "especificadas", "anteriormente",
}

TERMOS_PRODUCAO_CNAE = {
    "fabricar", "fabrico", "fabricacao", "produzir", "producao", "industrializar", "industria",
    "cultivar", "cultivo", "plantar", "criacao", "criar", "extrair", "extracao", "beneficiar",
}

TERMOS_SERVICO_CNAE = {
    "servico", "servicos", "prestar", "prestacao", "instalar", "instalacao", "manutencao",
    "conserto", "reparo", "assistencia", "consultoria", "treinamento", "curso", "aula",
}


def _nivel_embedding_atual():
    nivel = st.session_state.get("cnae_nivel_embedding", "Preciso")
    return nivel if nivel in NIVEIS_EMBEDDING_CNAE else "Preciso"


def _config_embedding_atual():
    return NIVEIS_EMBEDDING_CNAE[_nivel_embedding_atual()]


def _origem_cnae_legivel(origem):
    origem = str(origem or "").lower()
    partes = []
    if "texto" in origem or "direta" in origem:
        partes.append("texto")
    if "bm25" in origem:
        partes.append("BM25")
    if "embedding" in origem:
        partes.append("semântico")
    return " + ".join(partes)


def _cosine_similarity(v1, v2):
    if not v1 or not v2:
        return 0.0
    dot = sum(a * b for a, b in zip(v1, v2))
    norm1 = math.sqrt(sum(a * a for a in v1))
    norm2 = math.sqrt(sum(b * b for b in v2))
    if not norm1 or not norm2:
        return 0.0
    return dot / (norm1 * norm2)


@st.cache_data(show_spinner=False, ttl=30)
def ollama_disponivel():
    """Verifica apenas se o serviço local do Ollama está online, sem carregar modelo."""
    try:
        resp = requests.get(f"{OLLAMA_BASE_URL}/api/tags", timeout=2)
        return resp.status_code == 200
    except Exception:
        return False


def _modelos_embedding_preferidos():
    modelos = []
    for modelo in [MODELO_EMBEDDING_CNAE, MODELO_EMBEDDING_CNAE_FALLBACK]:
        modelo = str(modelo or "").strip()
        if modelo and modelo not in modelos:
            modelos.append(modelo)
    return modelos


def _chamar_ollama_embed(textos, timeout=25, modelos=None):
    textos = [str(t or "").strip() for t in (textos or []) if str(t or "").strip()]
    if not textos:
        return [], "Texto vazio.", None

    modelos = modelos or _modelos_embedding_preferidos()
    ultimo_erro = None
    for modelo in modelos:
        try:
            payload = {
                "model": modelo,
                "input": textos,
                "keep_alive": "10m",
            }
            resp = requests.post(OLLAMA_EMBED_URL, json=payload, timeout=timeout)
            resp.raise_for_status()
            dados = resp.json()
            embeddings = dados.get("embeddings", [])
            if len(embeddings) == len(textos):
                return embeddings, None, modelo
            ultimo_erro = "Resposta de embeddings incompleta."
        except Exception as e:
            ultimo_erro = str(e)
    return [], ultimo_erro or "Embedding indisponível.", None


def _limpar_fragmento_cnae(texto):
    texto = remover_espacos_extras(texto)
    texto = re.sub(r"\bn[aã]o especificad[oa]s? anteriormente\b", "", texto, flags=re.I)
    texto = re.sub(r"\bsemelhantes\b", "", texto, flags=re.I)
    texto = re.sub(r"\boutr[oa]s?\b", "", texto, flags=re.I)
    texto = re.sub(r"\bdemais\b", "", texto, flags=re.I)
    texto = re.sub(r"\s+,", ",", texto)
    texto = re.sub(r",\s*,", ",", texto)
    return remover_espacos_extras(texto.strip(" ,;-"))


def _fragmentos_semanticos_cnae(descricao):
    """Cria diferentes visões da descrição oficial para embedding.

    Não traduz produto específico e não contém exemplos. Apenas quebra a descrição
    oficial em partes úteis: descrição completa, núcleo depois de prefixos
    econômicos e itens separados por vírgula.
    """
    descricao = remover_espacos_extras(descricao)
    if not descricao:
        return []

    fragmentos = []

    def add(txt):
        txt = _limpar_fragmento_cnae(txt)
        if len(txt) >= 4 and txt.lower() not in {f.lower() for f in fragmentos}:
            fragmentos.append(txt)

    add(descricao)

    padroes_prefixo = [
        r"^com[eé]rcio varejista de\s+",
        r"^com[eé]rcio atacadista de\s+",
        r"^fabrica[cç][aã]o de\s+",
        r"^confec[cç][aã]o de\s+",
        r"^servi[cç]os? de\s+",
        r"^atividades? de\s+",
        r"^instala[cç][aã]o de\s+",
        r"^manuten[cç][aã]o e repara[cç][aã]o de\s+",
        r"^repara[cç][aã]o e manuten[cç][aã]o de\s+",
        r"^aluguel de\s+",
        r"^loca[cç][aã]o de\s+",
        r"^cultivo de\s+",
        r"^cria[cç][aã]o de\s+",
        r"^produ[cç][aã]o de\s+",
    ]
    nucleo = descricao
    for padrao in padroes_prefixo:
        novo = re.sub(padrao, "", nucleo, flags=re.I)
        if novo != nucleo:
            nucleo = novo
            add(nucleo)
            break

    # Se houver uma enumeração oficial, cada item vira uma visão semântica.
    for parte in re.split(r",|;|\be\b|\bou\b", nucleo, flags=re.I):
        parte = _limpar_fragmento_cnae(parte)
        if len(parte) >= 5 and _normalizar_busca(parte) not in STOPWORDS_BM25_CNAE:
            add(parte)

    return fragmentos[:8]


def _caminho_cache_embeddings_cnae(modelo):
    """Arquivo persistente do índice semântico.

    O cache em arquivo evita que o GAIA gere todos os embeddings a cada
    execução do Streamlit. A assinatura inclui o modelo e a base CNAE local.
    """
    try:
        cnaes = carregar_cnaes_local()
        assinatura = "\n".join(f"{c.get('codigo','')}|{c.get('descricao','')}" for c in cnaes)
        digest = hashlib.sha256((str(modelo) + "\n" + assinatura).encode("utf-8")).hexdigest()[:16]
        pasta = Path(".gaia_cache")
        pasta.mkdir(exist_ok=True)
        return pasta / f"cnae_embeddings_{_normalizar_busca(modelo).replace(' ', '_')}_{digest}.json"
    except Exception:
        return None


def _carregar_cache_embeddings_cnae():
    """Carrega índice semântico já pronto, sem chamar o Ollama."""
    for modelo in _modelos_embedding_preferidos():
        caminho = _caminho_cache_embeddings_cnae(modelo)
        if not caminho or not caminho.exists():
            continue
        try:
            dados = json.loads(caminho.read_text(encoding="utf-8"))
            base = dados.get("base", [])
            modelo_cache = dados.get("modelo") or modelo
            if base:
                return base, None, modelo_cache
        except Exception:
            continue
    return [], "Índice semântico ainda não foi criado.", None


def _salvar_cache_embeddings_cnae(base, modelo):
    caminho = _caminho_cache_embeddings_cnae(modelo)
    if not caminho:
        return
    try:
        payload = {
            "modelo": modelo,
            "gerado_em": datetime.now().isoformat(timespec="seconds"),
            "base": base,
        }
        caminho.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass


def _registros_para_embedding_cnae():
    """Gera poucos textos por CNAE para manter a busca rápida.

    A versão anterior criava vários fragmentos para cada CNAE e tentava
    indexar tudo na hora, o que podia travar por vários minutos. Agora o
    índice usa no máximo duas visões oficiais por CNAE: descrição completa
    e núcleo limpo da própria descrição. Não há exemplos virando regra.
    """
    registros = []
    for cnae in carregar_cnaes_local():
        fragmentos = _fragmentos_semanticos_cnae(cnae["descricao"])
        usados = []
        for idx, fragmento in enumerate(fragmentos[:2]):
            frag_norm = _normalizar_busca(fragmento)
            if not frag_norm or frag_norm in usados:
                continue
            usados.append(frag_norm)
            registros.append({
                "codigo": cnae["codigo"],
                "descricao": cnae["descricao"],
                "fragmento": fragmento,
                "peso_fragmento": 1.03 if idx == 0 else 1.0,
            })
    return registros


def _gerar_embeddings_cnaes_rapido(max_seconds=38):
    """Cria o índice semântico com limite de tempo.

    Se passar do limite, retorna erro e o GAIA segue com busca textual/BM25.
    Isso evita a tela presa por minutos quando o Ollama/modelo está lento.
    """
    if not ollama_disponivel():
        return [], "Ollama não está respondendo em localhost:11434.", None

    registros = _registros_para_embedding_cnae()
    if not registros:
        return [], "Base CNAE vazia.", None

    inicio = time.time()
    textos = [r["fragmento"] for r in registros]
    embeddings = []
    modelo_usado = None
    batch_size = 24

    for i in range(0, len(textos), batch_size):
        if time.time() - inicio > max_seconds:
            return [], "Índice semântico ainda está sendo preparado; busca local usada nesta tentativa.", modelo_usado

        lote = textos[i:i + batch_size]
        if modelo_usado:
            vetores, erro, modelo = _chamar_ollama_embed(lote, timeout=18, modelos=[modelo_usado])
        else:
            vetores, erro, modelo = _chamar_ollama_embed(lote, timeout=18)

        if erro or len(vetores) != len(lote):
            return [], erro or "Resposta de embeddings incompleta.", modelo_usado

        modelo_usado = modelo or modelo_usado
        embeddings.extend(vetores)

    base = []
    for registro, embedding in zip(registros, embeddings):
        base.append({**registro, "embedding": embedding, "modelo": modelo_usado})

    if base and modelo_usado:
        _salvar_cache_embeddings_cnae(base, modelo_usado)
    return base, None, modelo_usado


def _embeddings_fragmentos_cnaes_local():
    """Obtém a base de embeddings sem travar a aplicação."""
    base, erro_cache, modelo_cache = _carregar_cache_embeddings_cnae()
    if base:
        return base, None, modelo_cache
    return _gerar_embeddings_cnaes_rapido(max_seconds=38)


def _lock_preaquecimento_embeddings_cnae():
    try:
        pasta = Path(".gaia_cache")
        pasta.mkdir(exist_ok=True)
        return pasta / "cnae_embeddings_preheat.lock"
    except Exception:
        return None


def _indice_semantico_cnae_pronto():
    """Confere se o índice semântico já existe em cache, sem gerar nada novo."""
    base, _, modelo = _carregar_cache_embeddings_cnae()
    return bool(base), modelo


def _preaquecimento_semantico_em_andamento():
    lock = _lock_preaquecimento_embeddings_cnae()
    if not lock or not lock.exists():
        return False
    try:
        # Evita lock eterno se o processo foi fechado no meio da indexação.
        idade = time.time() - lock.stat().st_mtime
        if idade > 60 * 30:
            lock.unlink(missing_ok=True)
            return False
        return True
    except Exception:
        return True


def status_indice_semantico_cnae():
    """Texto curto para a interface sobre o preparo da busca por embeddings."""
    pronto, modelo = _indice_semantico_cnae_pronto()
    if pronto:
        return f"✅ Índice semântico pronto ({modelo or MODELO_EMBEDDING_CNAE})"
    if _preaquecimento_semantico_em_andamento():
        return "⏳ Índice semântico sendo preparado em segundo plano"
    return "⚪ Índice semântico será preparado automaticamente"


def _caminho_cache_aquecimento_cnae():
    """Arquivo com o resultado das buscas internas de aquecimento.

    Esse arquivo serve como prova simples de que o motor semântico já rodou
    uma consulta real, usando a mesma base oficial de CNAEs.
    """
    try:
        pasta = Path(".gaia_cache")
        pasta.mkdir(exist_ok=True)
        return pasta / "cnae_warmup_result.json"
    except Exception:
        return None


def _salvar_resultado_aquecimento_cnae(payload):
    caminho = _caminho_cache_aquecimento_cnae()
    if not caminho:
        return
    try:
        caminho.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


def _carregar_resultado_aquecimento_cnae():
    caminho = _caminho_cache_aquecimento_cnae()
    if not caminho or not caminho.exists():
        return None
    try:
        return json.loads(caminho.read_text(encoding="utf-8"))
    except Exception:
        return None


def _rankear_embedding_cnae_com_base(texto, base_embeddings, modelo_base, limite=2):
    """Executa uma busca semântica interna sem depender da interface.

    É usada no aquecimento para simular consultas reais, como comércio e
    serviço, depois que o índice de embeddings já existe. Não adiciona regras
    de produto nem decide CNAE; apenas testa o motor de recuperação.
    """
    texto = str(texto or "").strip()
    if not texto or not base_embeddings:
        return []

    consultas = _consultas_embedding_cnae(texto)
    vetores_texto, erro_texto, modelo_busca = _chamar_ollama_embed(
        consultas,
        timeout=18,
        modelos=[modelo_base] if modelo_base else None,
    )
    if erro_texto or not vetores_texto:
        return []

    por_codigo = {}
    for item in base_embeddings:
        melhor_score = 0.0
        for vetor_busca in vetores_texto:
            score = _cosine_similarity(vetor_busca, item["embedding"]) * 100 * item.get("peso_fragmento", 1.0)
            if score > melhor_score:
                melhor_score = score
        codigo = item["codigo"]
        atual = por_codigo.get(codigo)
        if atual is None or melhor_score > atual["score"]:
            por_codigo[codigo] = {
                "codigo": codigo,
                "descricao": item["descricao"],
                "score": round(melhor_score, 2),
                "fragmento_semantico": item.get("fragmento", ""),
                "modelo": modelo_busca or modelo_base,
            }

    return sorted(por_codigo.values(), key=lambda item: item["score"], reverse=True)[:limite]


def executar_buscas_padrao_aquecimento_cnae():
    """Faz consultas reais de aquecimento após o índice estar disponível.

    As consultas são genéricas — comércio e serviço — e não representam regras
    de CNAE. Elas apenas forçam o caminho completo da busca semântica a rodar
    para que a próxima busca do usuário não seja a primeira chamada pesada.
    """
    base, erro, modelo = _carregar_cache_embeddings_cnae()
    if erro or not base:
        return None

    consultas_padrao = ["comércio varejista", "prestação de serviço"]
    resultados = {}
    for consulta in consultas_padrao:
        resultados[consulta] = _rankear_embedding_cnae_com_base(consulta, base, modelo, limite=2)

    payload = {
        "status": "ok",
        "modelo": modelo,
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "consultas": resultados,
    }
    _salvar_resultado_aquecimento_cnae(payload)
    return payload


def status_aquecimento_cnae():
    """Resumo curto do aquecimento real do motor de busca."""
    payload = _carregar_resultado_aquecimento_cnae()
    if not payload or payload.get("status") != "ok":
        return "⚪ Busca semântica padrão ainda não testada"

    consultas = payload.get("consultas") or {}
    total = sum(1 for itens in consultas.values() if itens)
    modelo = payload.get("modelo") or MODELO_EMBEDDING_CNAE
    if total >= 2:
        return f"✅ Busca semântica aquecida e testada ({modelo})"
    return f"⏳ Busca semântica aquecendo ({modelo})"


def detalhes_aquecimento_cnae():
    """Retorna texto legível com dois resultados padrão do aquecimento."""
    payload = _carregar_resultado_aquecimento_cnae()
    if not payload or payload.get("status") != "ok":
        return ""
    linhas = []
    for consulta, itens in (payload.get("consultas") or {}).items():
        if not itens:
            continue
        primeiro = itens[0]
        linhas.append(f"{consulta}: {primeiro.get('codigo')} - {primeiro.get('descricao')}")
    return " | ".join(linhas[:2])


def preaquecer_indice_semantico_cnae_em_segundo_plano():
    """Cria o índice de embeddings assim que o GAIA abre.

    A primeira criação pode demorar alguns minutos, mas roda em uma thread
    separada para não travar a navegação. Depois de salvo em .gaia_cache,
    as próximas buscas reaproveitam o índice e ficam praticamente imediatas.
    """
    if st.session_state.get("_cnae_preheat_indice_semantico_iniciado"):
        return
    st.session_state["_cnae_preheat_indice_semantico_iniciado"] = True

    pronto, _ = _indice_semantico_cnae_pronto()
    if pronto or _preaquecimento_semantico_em_andamento():
        return

    lock = _lock_preaquecimento_embeddings_cnae()

    def _worker():
        try:
            if lock:
                try:
                    lock.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
                except Exception:
                    pass
            # Tempo maior porque aqui é pré-aquecimento em segundo plano, não busca do usuário.
            _gerar_embeddings_cnaes_rapido(max_seconds=60 * 8)
            executar_buscas_padrao_aquecimento_cnae()
        except Exception:
            pass
        finally:
            if lock:
                try:
                    lock.unlink(missing_ok=True)
                except Exception:
                    pass

    try:
        threading.Thread(target=_worker, daemon=True).start()
    except Exception:
        pass


def _tokens_bm25(texto):
    tokens = []
    for token in _normalizar_busca(texto).split():
        if len(token) < 3 or token in STOPWORDS_BM25_CNAE:
            continue
        tokens.extend(_variacoes_token_cnae(token))
    return list(dict.fromkeys(tokens))


def _variacoes_token_cnae(token):
    """Variações linguísticas genéricas, sem vocabulário por exemplo."""
    token = _normalizar_busca(token)
    variacoes = {token}
    if len(token) > 4 and token.endswith("s"):
        variacoes.add(token[:-1])
    if len(token) > 5 and token.endswith("es"):
        variacoes.add(token[:-2])
    # Sufixos comuns de estabelecimentos: sorveteria, papelaria, açaiteria etc.
    # A regra é morfológica; não traduz o produto.
    for sufixo in ["teria", "aria", "eria"]:
        if len(token) > len(sufixo) + 2 and token.endswith(sufixo):
            variacoes.add(token[:-len(sufixo)])
    return [v for v in variacoes if len(v) >= 3 and v not in STOPWORDS_BM25_CNAE]


@st.cache_data(show_spinner=False)
def _indice_bm25_cnaes():
    docs = []
    df = {}
    total_len = 0
    for cnae in carregar_cnaes_local():
        fragmentos = _fragmentos_semanticos_cnae(cnae["descricao"])
        texto_doc = " ".join([cnae["descricao"], cnae["busca"], " ".join(fragmentos)])
        tokens = _tokens_bm25(texto_doc)
        total_len += len(tokens)
        freqs = {}
        for tok in tokens:
            freqs[tok] = freqs.get(tok, 0) + 1
        for tok in freqs:
            df[tok] = df.get(tok, 0) + 1
        docs.append({
            "codigo": cnae["codigo"],
            "descricao": cnae["descricao"],
            "tokens": tokens,
            "freqs": freqs,
            "len": len(tokens) or 1,
        })
    avgdl = (total_len / len(docs)) if docs else 1
    return docs, df, avgdl, len(docs)


def _bm25_score(query_tokens, doc, df, avgdl, total_docs, k1=1.5, b=0.75):
    score = 0.0
    if not query_tokens or not total_docs:
        return 0.0
    for tok in query_tokens:
        tf = doc["freqs"].get(tok, 0)
        if tf <= 0:
            continue
        # IDF Okapi BM25 clássico com suavização.
        idf = math.log(1 + ((total_docs - df.get(tok, 0) + 0.5) / (df.get(tok, 0) + 0.5)))
        denom = tf + k1 * (1 - b + b * (doc["len"] / avgdl))
        score += idf * ((tf * (k1 + 1)) / denom)
    return score


def sugerir_cnaes_por_bm25(texto, limite=30):
    query_tokens = _tokens_bm25(texto)
    if not query_tokens:
        return []
    docs, df, avgdl, total_docs = _indice_bm25_cnaes()
    resultados = []
    for doc in docs:
        score = _bm25_score(query_tokens, doc, df, avgdl, total_docs)
        if score > 0:
            resultados.append({
                "codigo": doc["codigo"],
                "descricao": doc["descricao"],
                "score": score * 12,
                "score_bm25": score,
                "origem": "bm25",
            })
    resultados.sort(key=lambda item: item["score"], reverse=True)
    return resultados[:limite]


@st.cache_data(show_spinner=False)
def _frequencia_termos_cnae():
    contagem = {}
    total = 0
    for cnae in carregar_cnaes_local():
        total += 1
        for palavra in set(cnae["busca"].split()):
            contagem[palavra] = contagem.get(palavra, 0) + 1
    return contagem, total


def _peso_termo(token, frequencias, total):
    if not total:
        return 1.0
    proporcao = frequencias.get(token, 0) / total
    return max(0.15, min(1.0, 1.0 - (proporcao * 6)))


def _parece_item_curto_sem_acao(texto_normalizado):
    tokens = [t for t in _tokens_busca(texto_normalizado) if t not in STOPWORDS_BM25_CNAE]
    if not tokens or len(tokens) > 3:
        return False
    if _intencao_varejo(texto_normalizado) or _intencao_atacado(texto_normalizado):
        return False
    if any(t in texto_normalizado for t in TERMOS_PRODUCAO_CNAE | TERMOS_SERVICO_CNAE):
        return False
    return True


def _consultas_embedding_cnae(texto):
    texto = remover_espacos_extras(texto)
    texto_norm = _normalizar_busca(texto)
    if not texto:
        return []
    consultas = [texto]

    # Templates genéricos de intenção econômica. Não traduzem o produto nem
    # mapeiam exemplos; apenas fornecem contexto operacional para o embedding.
    if _intencao_varejo(texto_norm) or _parece_item_curto_sem_acao(texto_norm):
        consultas.append(f"comércio varejista de {texto}")
        consultas.append(f"venda ao consumidor de {texto}")
    if _intencao_atacado(texto_norm):
        consultas.append(f"comércio atacadista de {texto}")
        consultas.append(f"distribuição de {texto}")
    if _tem_termo(texto_norm, ["fabricar", "fabrico", "fabricacao", "produzir", "producao", "industria"]):
        consultas.append(f"fabricação e produção de {texto}")
    if _tem_termo(texto_norm, ["conserto", "reparo", "manutencao", "assistencia", "instalar", "instalacao"]):
        consultas.append(f"serviço de manutenção reparação instalação de {texto}")
    if _tem_termo(texto_norm, ["curso", "aula", "ensino", "treinamento", "escola"]):
        consultas.append(f"ensino treinamento curso de {texto}")
    if _tem_termo(texto_norm, ["alugar", "aluguel", "locacao", "locar"]):
        consultas.append(f"aluguel locação de {texto}")

    consultas.append(f"atividade econômica relacionada a {texto}")
    return list(dict.fromkeys([c for c in consultas if c.strip()]))[:6]


def sugerir_cnaes_por_embedding(texto, limite=30):
    texto = str(texto or "").strip()
    if not texto:
        return [], "Texto vazio."

    base_embeddings, erro_base, modelo_base = _embeddings_fragmentos_cnaes_local()
    if erro_base or not base_embeddings:
        return [], erro_base or "Base de embeddings indisponível."

    consultas = _consultas_embedding_cnae(texto)
    vetores_texto, erro_texto, modelo_busca = _chamar_ollama_embed(consultas, timeout=18, modelos=[modelo_base])
    if erro_texto or not vetores_texto:
        return [], erro_texto or "Embedding da busca indisponível."
    st.session_state["cnae_modelo_embedding_usado"] = modelo_busca or modelo_base or MODELO_EMBEDDING_CNAE

    por_codigo = {}
    for item in base_embeddings:
        melhor_score = 0.0
        for vetor_busca in vetores_texto:
            score = _cosine_similarity(vetor_busca, item["embedding"]) * 100 * item.get("peso_fragmento", 1.0)
            if score > melhor_score:
                melhor_score = score
        codigo = item["codigo"]
        atual = por_codigo.get(codigo)
        if atual is None or melhor_score > atual["score"]:
            por_codigo[codigo] = {
                "codigo": codigo,
                "descricao": item["descricao"],
                "score": melhor_score,
                "score_embedding": melhor_score,
                "fragmento_semantico": item.get("fragmento", ""),
                "origem": "embedding",
            }

    resultados = sorted(por_codigo.values(), key=lambda item: item["score"], reverse=True)
    return resultados[:limite], None


def _pontuar_cnae(texto_normalizado, tokens, cnae):
    score = 0
    busca = cnae["busca"]
    descricao_tokens = set(busca.split())
    intencao_varejo = _intencao_varejo(texto_normalizado)
    intencao_atacado = _intencao_atacado(texto_normalizado)
    frequencias, total_cnaes_freq = _frequencia_termos_cnae()

    if texto_normalizado:
        if texto_normalizado in descricao_tokens:
            score += 8
        elif len(texto_normalizado) > 3 and re.search(rf"\b{re.escape(texto_normalizado)}\b", busca):
            score += 8

    for token in tokens:
        peso = _peso_termo(token, frequencias, total_cnaes_freq)
        if token in descricao_tokens:
            score += 4 * peso
        elif len(token) > 3 and re.search(rf"\b{re.escape(token)}\b", busca):
            score += 2 * peso

    # Ajuste genérico por intenção operacional. Não contém exemplos de produto.
    if intencao_varejo and not intencao_atacado:
        if "comercio varejista" in busca:
            score += 10
        if "comercio atacadista" in busca:
            score -= 10
        if not any(t in texto_normalizado for t in TERMOS_PRODUCAO_CNAE):
            if any(t in busca for t in ["cultivo", "criacao", "extracao", "fabricacao"]):
                score -= 4

    if intencao_atacado:
        if "comercio atacadista" in busca or "representantes comerciais" in busca or "agentes do comercio" in busca:
            score += 10
        if "comercio varejista" in busca:
            score -= 6

    return score


def sugerir_cnaes_por_texto(texto, limite=10):
    texto_normalizado = _normalizar_busca(texto)
    tokens = _tokens_busca(texto)
    if not texto_normalizado or not tokens:
        return []

    resultados = []
    for cnae in carregar_cnaes_local():
        score = _pontuar_cnae(texto_normalizado, tokens, cnae)
        if score > 0:
            item = {"codigo": cnae["codigo"], "descricao": cnae["descricao"], "score": score, "origem": "texto"}
            resultados.append(item)

    resultados.sort(key=lambda x: (x["score"], x["codigo"]), reverse=True)
    return resultados[:limite]


def buscar_cnaes_conhecido(consulta, limite=30):
    """Busca direta para quando o usuário já sabe o CNAE, parte do código
    ou uma palavra-chave da descrição oficial.

    Este modo não usa IA nem embeddings: a fonte é somente a base CNAE local.
    """
    consulta_original = str(consulta or "").strip()
    consulta_norm = _normalizar_busca(consulta_original)
    consulta_digitos = somente_digitos(consulta_original)
    tokens = _tokens_busca(consulta_original)

    if not consulta_norm and not consulta_digitos:
        return []

    resultados = []
    for cnae in carregar_cnaes_local():
        codigo = cnae["codigo"]
        descricao = cnae["descricao"]
        busca = cnae["busca"]
        codigo_digitos = somente_digitos(codigo)
        descricao_tokens = set(busca.split())
        score = 0.0

        if consulta_digitos:
            if consulta_digitos == codigo_digitos:
                score += 120
            elif codigo_digitos.startswith(consulta_digitos):
                score += 90 - min(25, len(codigo_digitos) - len(consulta_digitos))
            elif consulta_digitos in codigo_digitos:
                score += 55

        if consulta_norm:
            if consulta_norm == busca:
                score += 95
            elif re.search(rf"\b{re.escape(consulta_norm)}\b", busca):
                score += 70
            elif consulta_norm in busca:
                score += 42

        if tokens:
            frequencias, total_cnaes_freq = _frequencia_termos_cnae()
            for token in tokens:
                peso = _peso_termo(token, frequencias, total_cnaes_freq)
                if token in descricao_tokens:
                    score += 14 * peso
                elif len(token) > 3 and re.search(rf"\b{re.escape(token)}\b", busca):
                    score += 8 * peso
                elif len(token) > 4 and token in busca:
                    score += 3 * peso

        if score > 0:
            resultados.append({
                "codigo": codigo,
                "descricao": descricao,
                "score": score,
                "origem": "direta",
            })

    resultados.sort(key=lambda item: (item["score"], item["codigo"]), reverse=True)
    return resultados[:limite]


def _bonus_intencao_generica(texto_usuario, item):
    texto_norm = _normalizar_busca(texto_usuario)
    busca = _normalizar_busca(item.get("descricao", ""))
    bonus = 0.0

    if _intencao_varejo(texto_norm) and not _intencao_atacado(texto_norm):
        if "comercio varejista" in busca:
            bonus += 8.0
        if "comercio atacadista" in busca:
            bonus -= 7.0
        if not any(t in texto_norm for t in TERMOS_PRODUCAO_CNAE):
            if any(t in busca for t in ["cultivo", "criacao", "extracao", "fabricacao"]):
                bonus -= 4.0

    if _intencao_atacado(texto_norm):
        if "comercio atacadista" in busca or "representantes comerciais" in busca:
            bonus += 8.0
        if "comercio varejista" in busca:
            bonus -= 5.0

    if any(t in texto_norm for t in TERMOS_PRODUCAO_CNAE):
        if any(t in busca for t in ["fabricacao", "producao", "cultivo", "criacao", "extracao"]):
            bonus += 8.0

    if any(t in texto_norm for t in TERMOS_SERVICO_CNAE):
        if any(t in busca for t in ["servicos", "manutencao", "reparacao", "instalacao", "ensino", "treinamento"]):
            bonus += 6.0

    # Quando o usuário digita só um item curto, o sistema não força um contexto,
    # mas dá leve preferência a atividades de comercialização, pois é o caso mais
    # comum no evento de CNAE sem exigir pergunta extra ao usuário.
    if _parece_item_curto_sem_acao(texto_norm):
        if "comercio varejista" in busca:
            bonus += 3.0

    return bonus


def _combinar_candidatos_cnae(*listas, limite=20, config=None, texto_usuario=""):
    config = config or _config_embedding_atual()
    pesos = {
        "texto": config.get("peso_texto", 1.05),
        "direta": config.get("peso_texto", 1.05),
        "bm25": config.get("peso_bm25", 1.15),
        "embedding": config.get("peso_embedding", 1.00),
    }

    combinados = {}
    for lista in listas:
        for rank, item in enumerate(lista or [], start=1):
            codigo = item.get("codigo")
            if not codigo:
                continue
            origem = item.get("origem") or "texto"
            score_bruto = float(item.get("score", 0) or 0)
            origem_principal = "embedding" if "embedding" in origem else "bm25" if "bm25" in origem else "texto"
            peso = pesos.get(origem_principal, 1.0)

            # Combina score bruto com Reciprocal Rank Fusion para não depender
            # de uma única escala de pontuação.
            score_ponderado = (score_bruto * peso) + (95.0 * peso / (50 + rank))

            if codigo not in combinados:
                combinados[codigo] = {
                    "codigo": codigo,
                    "descricao": item.get("descricao", ""),
                    "score": 0.0,
                    "score_texto": 0.0,
                    "score_bm25": 0.0,
                    "score_embedding": 0.0,
                    "origens": set(),
                }

            combinados[codigo]["score"] += score_ponderado
            combinados[codigo]["origens"].add(origem_principal)
            if origem_principal == "embedding":
                combinados[codigo]["score_embedding"] = max(combinados[codigo]["score_embedding"], score_bruto)
            elif origem_principal == "bm25":
                combinados[codigo]["score_bm25"] = max(combinados[codigo]["score_bm25"], score_bruto)
            else:
                combinados[codigo]["score_texto"] = max(combinados[codigo]["score_texto"], score_bruto)

    resultados = []
    for item in combinados.values():
        origens = item.pop("origens")
        if len(origens) >= 2:
            item["score"] += config.get("bonus_convergencia", 15.0)
        item["score"] += _bonus_intencao_generica(texto_usuario, item)
        item["origem"] = "+".join(sorted(origens))
        resultados.append(item)

    resultados.sort(
        key=lambda item: (
            len(item.get("origem", "").split("+")),
            item.get("score", 0),
            item.get("score_bm25", 0),
            item.get("score_texto", 0),
            item.get("score_embedding", 0),
        ),
        reverse=True,
    )
    return resultados[:limite]


def _progresso_cnae(progresso, valor, texto):
    if progresso is not None:
        progresso.progress(valor, text=texto)


def _registrar_origem_busca_cnae(origem, detalhe=""):
    st.session_state["cnae_ultima_origem_busca"] = origem
    st.session_state["cnae_ultima_origem_detalhe"] = detalhe


def _mostrar_origem_busca_cnae():
    """Mantém a interface limpa para o usuário final.

    Detalhes técnicos da busca ficam internos; na apresentação, o usuário só vê
    as sugestões oficiais e mensagens de validação necessárias.
    """
    origem = st.session_state.get("cnae_ultima_origem_busca")
    if origem == "bloqueado":
        st.warning("Busca bloqueada por termo incompatível com a finalidade do sistema.")


def sugerir_cnaes_inteligente(texto_usuario, limite=10, progresso=None):
    texto_usuario = str(texto_usuario or "").strip()
    _progresso_cnae(progresso, 5, "Preparando busca...")
    if _termo_bloqueado_cnae(_normalizar_busca(texto_usuario)):
        _registrar_origem_busca_cnae("bloqueado")
        _progresso_cnae(progresso, 100, "Busca finalizada.")
        return []

    config = _config_embedding_atual()
    limite_texto = config.get("limite_texto", 40)
    limite_bm25 = config.get("limite_bm25", 40)
    limite_embedding = config.get("limite_embedding", 36)

    _registrar_origem_busca_cnae("hibrida")
    _progresso_cnae(progresso, 25, "Analisando descrições oficiais...")
    candidatos_texto = sugerir_cnaes_por_texto(texto_usuario, limite=limite_texto)
    candidatos_bm25 = sugerir_cnaes_por_bm25(texto_usuario, limite=limite_bm25)

    _progresso_cnae(progresso, 55, "Comparando atividades semelhantes...")
    candidatos_embedding, erro_embedding = sugerir_cnaes_por_embedding(texto_usuario, limite=limite_embedding)
    if erro_embedding:
        st.session_state["cnae_ultima_origem_detalhe_embedding"] = erro_embedding
        st.session_state["cnae_modelo_embedding_usado"] = "indisponível"
    else:
        st.session_state["cnae_ultima_origem_detalhe_embedding"] = ""
        st.session_state["cnae_modelo_embedding_usado"] = st.session_state.get("cnae_modelo_embedding_usado") or MODELO_EMBEDDING_CNAE

    _progresso_cnae(progresso, 80, "Organizando melhores sugestões...")
    candidatos = _combinar_candidatos_cnae(
        candidatos_texto,
        candidatos_bm25,
        candidatos_embedding,
        limite=50,
        config=config,
        texto_usuario=texto_usuario,
    )

    # Fallback: se o embedding falhou e os demais não acharam nada, usa busca direta
    # ampliada. Não cria resposta falsa, mas evita tela vazia quando houver termo oficial.
    if not candidatos:
        _registrar_origem_busca_cnae("local", erro_embedding or "Embedding sem resultados; usando fallback textual.")
        candidatos = buscar_cnaes_conhecido(texto_usuario, limite=30)

    limite_saida = min(limite, int(config.get("limite_final", limite)))
    _progresso_cnae(progresso, 100, "Busca finalizada.")
    return candidatos[:limite_saida]

def limpar_descricao_objeto_social(descricao):
    texto = str(descricao or "")
    texto = re.sub(r"\s*-\s*", ", ", texto)
    for termo in TERMOS_REMOCAO_OBJETO:
        texto = re.sub(termo, "", texto, flags=re.IGNORECASE)
    texto = re.sub(r"\s+,", ",", texto)
    texto = re.sub(r",\s*,+", ",", texto)
    texto = re.sub(r"\s+", " ", texto).strip(" ,.;")
    return texto

def montar_objeto_social_sugerido(descricao, sugestoes):
    if not sugestoes:
        return limpar_descricao_objeto_social(descricao).upper()

    descricoes_limpas = []
    for item in sugestoes:
        desc = limpar_descricao_objeto_social(item.get("descricao", ""))
        if desc and desc.lower() not in [d.lower() for d in descricoes_limpas]:
            descricoes_limpas.append(desc)

    objeto = ", ".join(descricoes_limpas)
    return objeto.upper()

def _label_cnae(item):
    if not item:
        return ""
    return f"{item['codigo']} - {item['descricao']}"

def detectar_multiplas_atividades_principal(texto):
    """Detecta quando o usuário está tentando misturar atividades econômicas distintas.

    A regra anterior tratava qualquer "e" como sinal de múltiplas atividades. Isso
    bloqueava frases naturais como "vender comidas, bebidas e afins", que na prática
    descrevem uma mesma intenção econômica ampla. Agora o bloqueio fica restrito a
    separadores fortes e conectores que realmente indiquem inclusão de outra atividade.
    """
    texto_original = str(texto or "")
    texto_limpo = _normalizar_busca(texto_original)
    if not texto_limpo:
        return False

    # Expressões como "e afins", "e similares" e "e semelhantes" não indicam,
    # sozinhas, uma segunda atividade; elas apenas ampliam a descrição do mesmo ramo.
    texto_para_analise = re.sub(
        r"\be\s+(afins|similares|semelhantes|correlatos|relacionados)\b",
        " ",
        texto_limpo,
    )

    separadores_multiplos = [";", "\n", "+", " / "]
    if any(sep in texto_original for sep in separadores_multiplos):
        return True

    # Vírgula só bloqueia quando há indício claro de lista de atividades diferentes.
    # Em descrição de produto, ela pode aparecer apenas como enumeração do mesmo ramo.
    if "," in texto_original and any(
        re.search(padrao, texto_para_analise)
        for padrao in [r"\be tambem\b", r"\btambem\b", r"\balem de\b", r"\bjunto com\b"]
    ):
        return True

    padroes_multiplos = [
        r"\be tambem\b",
        r"\btambem\b",
        r"\balem de\b",
        r"\bjunto com\b",
        r"\bprestacao de .*\be\s+comercio de\b",
        r"\bcomercio de .*\be\s+prestacao de\b",
    ]
    return any(re.search(padrao, texto_para_analise) for padrao in padroes_multiplos)

def montar_objeto_social_por_cnaes(cnae_principal, cnaes_secundarios):
    selecionados = []
    if cnae_principal:
        selecionados.append(cnae_principal)
    selecionados.extend(cnaes_secundarios or [])
    return montar_objeto_social_sugerido("", selecionados)

def montar_texto_cnaes_classificados(cnae_principal, cnaes_secundarios):
    linhas = []
    if cnae_principal:
        linhas.append(f"CNAE principal: {_label_cnae(cnae_principal)}")
    if cnaes_secundarios:
        linhas.append("CNAEs secundários:")
        linhas.extend([_label_cnae(item) for item in cnaes_secundarios])
    return "\n".join(linhas)

def _mesmo_cnae(a, b):
    return bool(a and b and a.get("codigo") == b.get("codigo"))

def _lista_contem_cnae(lista, item):
    return any(_mesmo_cnae(existente, item) for existente in (lista or []))

def _adicionar_cnae_unico(lista, item):
    lista = list(lista or [])
    if item and not _lista_contem_cnae(lista, item):
        lista.append(item)
    return lista

def _remover_cnae_por_codigo(lista, codigo):
    return [item for item in (lista or []) if item.get("codigo") != codigo]

def _ordinal(n):
    nomes = {1: "Primeira", 2: "Segunda", 3: "Terceira", 4: "Quarta", 5: "Quinta", 6: "Sexta", 7: "Sétima", 8: "Oitava", 9: "Nona", 10: "Décima"}
    return nomes.get(n, str(n) + "ª")

def processar_clausulas_prototype(eventos_selecionados):
    clausulas_prontas = []
    for i, evento in enumerate(eventos_selecionados, start=1):
        prototipo = CATALOGO_CLAUSULAS[evento["tipo"]].clone()
        texto_preenchido = prototipo.aplicar_dados(evento["dados"])
        clausulas_prontas.append(f"Cláusula {_ordinal(i)} – {texto_preenchido}")
    return clausulas_prontas

# ======================================================================
# GERAÇÃO DO DOCUMENTO DOCX / PDF
# ======================================================================
def _add_page_number_docx(paragraph):
    """Adiciona campo PAGE no DOCX para numeração automática."""
    run = paragraph.add_run()
    fld_char_1 = OxmlElement('w:fldChar')
    fld_char_1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char_2 = OxmlElement('w:fldChar')
    fld_char_2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def _configurar_numeracao_paginas_docx(doc):
    """Mantém somente a numeração de página, sem cabeçalho, logo ou rodapé personalizado."""
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.text = ""
        run = p.add_run("Página ")
        run.font.name = 'Arial'
        run.font.size = Pt(9)
        _add_page_number_docx(p)


def _formatar_qualificacao_socio(s):
    partes = [s.get('nome', '')]
    if s.get('nacionalidade'): partes.append(f"nacionalidade {s['nacionalidade']}")
    if s.get('estado_civil'): partes.append(s['estado_civil'])
    if s.get('data_nasc'): partes.append(f"nascido em {s['data_nasc']}")
    if s.get('profissao'): partes.append(f"profissão: {s['profissao']}")
    if s.get('cpf'): partes.append(f"inscrito no CPF sob o nº {s['cpf']}")
    if s.get('rg'): partes.append(f"portador da identidade nº {s['rg']}")
    if s.get('orgao_emissor'): partes.append(f"órgão expedidor {s['orgao_emissor']}")

    end_partes = []
    if s.get('logradouro'): end_partes.append(s['logradouro'])
    if s.get('numero'): end_partes.append(f"número {s['numero']}")
    if s.get('complemento'): end_partes.append(s['complemento'])
    if s.get('bairro'): end_partes.append(f"bairro {s['bairro']}")
    if s.get('cidade') and s.get('uf'): end_partes.append(f"município {s['cidade']} - {s['uf']}")
    if s.get('cep'): end_partes.append(f"CEP {s['cep']}")
    if end_partes:
        partes.append(f"residente e domiciliado no(a) {', '.join(end_partes)}")
    return ', '.join([p for p in partes if p]) + '.'


def _formatar_preambulo_empresa(empresa):
    emp_partes = [f"Sócio(s) da sociedade {empresa.get('nome_empresarial', '')}"]
    end_emp = []
    if empresa.get('logradouro'): end_emp.append(empresa['logradouro'])
    if empresa.get('numero'): end_emp.append(f"número {empresa['numero']}")
    if empresa.get('complemento'): end_emp.append(empresa['complemento'])
    if empresa.get('bairro'): end_emp.append(f"bairro {empresa['bairro']}")
    if empresa.get('cidade') and empresa.get('uf'): end_emp.append(f"município {empresa['cidade']} - {empresa['uf']}")
    if empresa.get('cep'): end_emp.append(f"CEP {empresa['cep']}")
    if end_emp:
        emp_partes.append(f"com sede no(a) {', '.join(end_emp)}")
    if empresa.get('nire'):
        emp_partes.append(f"com ato constitutivo arquivado na Junta Comercial sob NIRE {empresa['nire']}")
    else:
        emp_partes.append("com ato constitutivo arquivado na Junta Comercial")
    if empresa.get('cnpj'):
        emp_partes.append(f"inscrita no CNPJ sob o nº {empresa['cnpj']}")
    return ', '.join(emp_partes) + ", resolvem alterar o contrato social, mediante as cláusulas e condições seguintes:"


def gerar_docx_modular(empresa, socios, clausulas_prontas, data_documento_texto=None):
    """Gera o instrumento de alteração contratual em DOCX.

    Esta versão mantém a geração objetiva, sem cabeçalho, rodapé, logo, minuta,
    testemunhas ou relatório de conferência. A única melhoria visual mantida é a
    numeração de página.
    """
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)

    _configurar_numeracao_paginas_docx(doc)

    def add_paragrafo(texto, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(texto)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = 'Arial'
        return p

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(12)
    run_titulo = titulo.add_run(f"ALTERAÇÃO CONTRATUAL DA SOCIEDADE {empresa.get('nome_empresarial', '')}")
    run_titulo.bold = True
    run_titulo.font.size = Pt(14)
    run_titulo.font.name = 'Arial'

    for s in socios:
        add_paragrafo(_formatar_qualificacao_socio(s))

    add_paragrafo(_formatar_preambulo_empresa(empresa))

    for clausula in clausulas_prontas:
        add_paragrafo(clausula)

    add_paragrafo(
        f"Cláusula {_ordinal(len(clausulas_prontas) + 1)} – Permanecem inalteradas e em pleno vigor todas as demais cláusulas e condições do contrato social que não foram expressamente modificadas pelo presente instrumento."
    )
    add_paragrafo("E, por estarem assim justos e contratados, assinam o presente instrumento para que produza seus efeitos legais.")
    add_paragrafo("")

    cidade_empresa = empresa.get('cidade') or '_____________'
    data_final = data_documento_texto or formatar_data_documento()
    add_paragrafo(f"{cidade_empresa}, {data_final}.", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragrafo("")

    for s in socios:
        add_paragrafo("_" * 50, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        add_paragrafo(f"{s.get('nome', '')}: Sócio/Administrador", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _desenhar_numero_pagina_pdf(canvas, doc):
    """Desenha somente a numeração de página no PDF."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    largura, _ = A4
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(largura - 3 * cm, 1.35 * cm, f"Página {doc.page}")
    canvas.restoreState()


def gerar_pdf_modular(empresa, socios, clausulas_prontas, data_documento_texto=None):
    """Gera o PDF diretamente em Python, sem depender de internet ou LibreOffice."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import cm
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except Exception as e:
        return None, f"Biblioteca ReportLab não instalada. Instale com: pip install reportlab. Detalhe: {e}"

    try:
        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=A4,
            leftMargin=3 * cm,
            rightMargin=3 * cm,
            topMargin=2.5 * cm,
            bottomMargin=2.5 * cm,
            title=f"Alteração Contratual {empresa.get('nome_empresarial', '')}",
            author="GAIA - Geração Automatizada de Instrumentos de Alteração",
        )

        styles = getSampleStyleSheet()
        titulo_style = ParagraphStyle('GAIA_Titulo', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=14, leading=18, alignment=TA_CENTER, spaceAfter=14)
        normal_style = ParagraphStyle('GAIA_Normal', parent=styles['Normal'], fontName='Helvetica', fontSize=11, leading=15, alignment=TA_JUSTIFY, firstLineIndent=0, spaceAfter=8)
        centro_style = ParagraphStyle('GAIA_Centro', parent=normal_style, alignment=TA_CENTER, spaceAfter=10)
        assinatura_style = ParagraphStyle('GAIA_Assinatura', parent=normal_style, alignment=TA_CENTER, spaceAfter=12)

        def esc(texto):
            return html.escape(str(texto or '')).replace('\n', '<br/>')

        story = []
        story.append(Paragraph(esc(f"ALTERAÇÃO CONTRATUAL DA SOCIEDADE {empresa.get('nome_empresarial', '')}"), titulo_style))

        for s in socios:
            story.append(Paragraph(esc(_formatar_qualificacao_socio(s)), normal_style))

        story.append(Paragraph(esc(_formatar_preambulo_empresa(empresa)), normal_style))

        for clausula in clausulas_prontas:
            story.append(Paragraph(esc(clausula), normal_style))

        story.append(Paragraph(esc(f"Cláusula {_ordinal(len(clausulas_prontas) + 1)} – Permanecem inalteradas e em pleno vigor todas as demais cláusulas e condições do contrato social que não foram expressamente modificadas pelo presente instrumento."), normal_style))
        story.append(Paragraph(esc("E, por estarem assim justos e contratados, assinam o presente instrumento para que produza seus efeitos legais."), normal_style))
        story.append(Spacer(1, 16))

        cidade_empresa = empresa.get('cidade') or '_____________'
        data_final = data_documento_texto or formatar_data_documento()
        story.append(Paragraph(esc(f"{cidade_empresa}, {data_final}."), centro_style))
        story.append(Spacer(1, 24))

        for s in socios:
            story.append(Paragraph("_" * 46, assinatura_style))
            story.append(Paragraph(esc(f"{s.get('nome', '')}: Sócio/Administrador"), assinatura_style))
            story.append(Spacer(1, 8))

        doc.build(story, onFirstPage=_desenhar_numero_pagina_pdf, onLaterPages=_desenhar_numero_pagina_pdf)
        buf.seek(0)
        return buf, None
    except Exception as e:
        return None, f"Erro ao gerar PDF diretamente pelo Python: {e}"


# Mantém compatibilidade com chamadas antigas do código. Como a opção escolhida foi
# geração direta em Python, esta função não converte DOCX; ela apenas informa que a
# chamada antiga não deve mais ser usada.
def converter_docx_para_pdf(docx_bytes):
    return None, "Conversão DOCX para PDF desativada. O GAIA agora gera PDF diretamente em Python."

# ======================================================================
# SELEÇÃO E PREENCHIMENTO DOS EVENTOS
# ======================================================================
def tela_orquestrador():
    cabecalho()
    st.subheader("Seleção de Eventos")
    st.markdown("Selecione os eventos de alteração desejados e clique em Avançar.")
    
    with st.form("form_eventos_selecao"):
        alt_220 = st.checkbox("220 - Alteração do nome empresarial")
        alt_221 = st.checkbox("221 - Alteração do título do estabelecimento (nome fantasia)")
        alt_211 = st.checkbox("211 - Alteração de endereço dentro do mesmo município")
        alt_209 = st.checkbox("209 - Alteração de endereço entre municípios diferentes")
        alt_210 = st.checkbox("210 - Alteração de endereço entre estados diferentes")
        alt_248 = st.checkbox("248 - Alteração do tipo de unidade")
        alt_244 = st.checkbox("244 - Alteração de atividades econômicas / objeto social")
        alt_249 = st.checkbox("249 - Alteração da forma de atuação")
        alt_225 = st.checkbox("225 - Alteração da natureza jurídica (transformação)")
        alt_capital = st.checkbox("Alteração de capital social")
        alt_052 = st.checkbox("052 - Reativação (Artigo 60 da Lei nº 8.934/1994)")
        alt_999 = st.checkbox("999 - Licenciamento de estabelecimento legado")
        
        btn_avancar_enter = st.form_submit_button("Avançar →", type="primary", use_container_width=True, key="enter_default_eventos")
        c1, c2 = st.columns(2)
        btn_voltar = c1.form_submit_button("← Voltar", type="secondary", use_container_width=True, key="fsb_voltar_orquestrador")
        btn_avancar_click = c2.form_submit_button("Avançar →", type="primary", use_container_width=True, key="btn_avancar_eventos")
        btn_avancar = btn_avancar_enter or btn_avancar_click
        
    if btn_voltar:
        st.session_state.mostrar_confirmacao_repetidos = False
        nav('lista_socios')
        
    if btn_avancar:
        selecionados = []
        enderecos_selecionados = []
        if alt_220: selecionados.append("220")
        if alt_221: selecionados.append("221")
        if alt_211: enderecos_selecionados.append("211")
        if alt_209: enderecos_selecionados.append("209")
        if alt_210: enderecos_selecionados.append("210")
        if alt_248: selecionados.append("248")
        if alt_244: selecionados.append("244")
        if alt_249: selecionados.append("249")
        if alt_225: selecionados.append("225")
        if alt_capital: selecionados.append("capital")
        if alt_052: selecionados.append("052")
        if alt_999: selecionados.append("999")
        
        if len(enderecos_selecionados) > 1:
            msg_error("Selecione apenas uma opção de alteração de endereço.")
            return
        selecionados.extend(enderecos_selecionados)
        
        if not selecionados:
            msg_error("Selecione pelo menos um evento para continuar.")
        else:
            st.session_state.eventos_selecionados_lista = selecionados
            st.session_state.mostrar_confirmacao_repetidos = False
            st.session_state.documento_gerado = None
            nav('orquestrador_dados')


def renderizar_documento_gerado():
    """Mostra o painel de download do documento gerado e mantém os arquivos disponíveis após reruns."""
    doc = st.session_state.get("documento_gerado")
    if not doc:
        return

    st.markdown(
        """
        <div class="gaia-bottom-section" id="documento-gerado" style="margin-top: 24px; margin-bottom: 14px;">
            <div class="gaia-bottom-title">📥 Documento pronto para download</div>
            <div style="text-align:center; color:var(--gaia-muted); font-size:0.94rem; margin-top:6px;">
                Clique em <strong>Baixar WORD (DOCX)</strong> ou <strong>Baixar PDF</strong> para salvar o documento no seu computador.
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    ac1, ac2, ac3 = st.columns(3)
    if ac1.button("← Voltar ao menu", type="secondary", use_container_width=True, key="voltar_doc_menu_top"):
        st.session_state.empresa_id = None
        st.session_state.socio_id = None
        st.session_state.eventos_selecionados_lista = []
        st.session_state.mostrar_confirmacao_repetidos = False
        st.session_state.documento_gerado = None
        nav('menu')

    ac2.download_button(
        label="⬇️ 📘 Baixar WORD (DOCX)",
        data=doc.get("docx_bytes", b""),
        file_name=doc.get("nome_arquivo_docx", "Alteracao_Contratual.docx"),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
        use_container_width=True,
        key="download_docx_gerado_persistente",
    )

    pdf_bytes = doc.get("pdf_bytes")
    if pdf_bytes:
        ac3.download_button(
            label="⬇️ 📕 Baixar PDF",
            data=pdf_bytes,
            file_name=doc.get("nome_arquivo_pdf", "Alteracao_Contratual.pdf"),
            mime="application/pdf",
            type="primary",
            use_container_width=True,
            key="download_pdf_gerado_persistente",
        )
    else:
        ac3.button("PDF indisponível", type="secondary", use_container_width=True, disabled=True, key="pdf_indisponivel_persistente")


    if doc.get("erro_pdf"):
        st.markdown(
            f"<div class='gaia-inline-note' style='margin-top:8px;'>PDF não gerado neste ambiente: {html.escape(str(doc.get('erro_pdf')))}</div>",
            unsafe_allow_html=True,
        )


    clausulas = doc.get("clausulas_prontas") or []
    if clausulas:
        with st.expander("Ver cláusulas geradas", expanded=False):
            for c in clausulas:
                st.info(c)

    # Rola automaticamente até o painel de download na primeira vez em que
    # o documento é gerado (flag definida em gerar_documento antes do rerun).
    if st.session_state.pop('_rolar_para_downloads', False):
        components.html(
            """
            <script>
            (function() {
                // Aguarda o DOM estar completamente pronto antes de rolar.
                function tentarRolar(tentativas) {
                    const doc = window.parent.document;
                    const el = doc.getElementById('documento-gerado');
                    if (el) {
                        el.scrollIntoView({behavior: 'smooth', block: 'start'});
                    } else if (tentativas > 0) {
                        window.parent.setTimeout(function() { tentarRolar(tentativas - 1); }, 150);
                    } else {
                        window.parent.scrollTo({top: window.parent.document.body.scrollHeight, behavior: 'smooth'});
                    }
                }
                window.parent.setTimeout(function() { tentarRolar(6); }, 250);
            })();
            </script>
            """,
            height=0,
            width=0,
        )

def _nome_evento_wizard(codigo):
    nomes = {
        "220": "Alteração do nome empresarial",
        "221": "Alteração do título do estabelecimento (nome fantasia)",
        "211": "Alteração de endereço dentro do mesmo município",
        "209": "Alteração de endereço entre municípios diferentes",
        "210": "Alteração de endereço entre estados diferentes",
        "248": "Alteração do tipo de unidade",
        "244": "Alteração de atividades econômicas / objeto social",
        "249": "Alteração da forma de atuação",
        "225": "Alteração da natureza jurídica (transformação)",
        "capital": "Alteração de capital social",
        "052": "Reativação (Artigo 60 Lei 8.934/94)",
        "999": "Licenciamento de estabelecimento legado",
    }
    return nomes.get(codigo, codigo)


def _wizard_key_eventos(selecionados):
    return "|".join(selecionados or [])


def _inicializar_wizard_eventos(selecionados):
    chave = _wizard_key_eventos(selecionados)
    if st.session_state.get("wizard_eventos_chave") != chave:
        st.session_state["wizard_eventos_chave"] = chave
        st.session_state["wizard_evento_idx"] = 0
        st.session_state["wizard_eventos_dados"] = {}
        st.session_state.mostrar_confirmacao_repetidos = False


def _render_progresso_wizard(selecionados, idx):
    total = len(selecionados)
    progresso = (idx + 1) / max(total, 1)
    st.progress(progresso)
    st.caption(f"Etapa {idx + 1} de {total}: {_nome_evento_wizard(selecionados[idx])}")


def _evento_wizard_salvo(codigo):
    return st.session_state.get("wizard_eventos_dados", {}).get(codigo, {})


def _salvar_evento_wizard(codigo, evento):
    dados = dict(st.session_state.get("wizard_eventos_dados", {}))
    dados[codigo] = evento
    st.session_state["wizard_eventos_dados"] = dados


def _eventos_processados_wizard(selecionados):
    dados = st.session_state.get("wizard_eventos_dados", {})
    return [dados[codigo] for codigo in selecionados if codigo in dados]


def _render_resumo_evento_wizard(evento):
    dados = evento.get("dados", {}) or {}
    tipo = evento.get("tipo")
    st.markdown(f"**{tipo} - {_nome_evento_wizard(tipo)}**")
    if not dados:
        st.caption("Evento automático, sem campos manuais.")
        return
    for campo, valor in dados.items():
        rotulo = campo.replace("_", " ").capitalize()
        if isinstance(valor, str) and "\n" in valor:
            st.markdown(f"**{rotulo}:**")
            st.text(valor)
        else:
            st.markdown(f"- **{rotulo}:** {valor}")


def _limpar_checks_cnae(prefixo=None):
    for chave in list(st.session_state.keys()):
        chave_str = str(chave)
        if prefixo:
            if chave_str.startswith(prefixo):
                del st.session_state[chave]
        elif chave_str.startswith("check_cnae_principal_") or chave_str.startswith("check_cnae_secundario_"):
            del st.session_state[chave]


def _set_modo_busca_cnae(modo):
    """Define o modo de busca do evento 244 sem limpar a seleção final já feita."""
    if st.session_state.get("cnae_modo_busca") != modo:
        st.session_state["cnae_modo_busca"] = modo
        st.session_state["cnae_principal_sugestoes"] = []
        st.session_state["cnaes_secundarios_sugestoes"] = []
        _limpar_checks_cnae()
        set_flash("info", "Modo de busca de CNAE alterado.")
        st.rerun()


def _render_cards_modo_busca_cnae():
    modo_atual = st.session_state.get("cnae_modo_busca")
    st.markdown("#### Como deseja localizar o CNAE?")
    col_direta, col_inteligente = st.columns(2)

    with col_direta:
        selecionado = modo_atual == "conhecido"
        borda = "#00866B" if selecionado else "#E3E8EF"
        fundo = "#F3FAF8" if selecionado else "#FFFFFF"
        st.markdown(
            f"""
            <div class="gaia-card" style="min-height: 150px; border: 1.6px solid {borda}; background:{fundo};">
                <div style="font-size:1.55rem; margin-bottom:8px;">🔎</div>
                <div style="font-weight:800; font-size:1.02rem; margin-bottom:5px;">Buscar CNAE conhecido</div>
                <div style="color:#5A6679; font-size:.86rem; line-height:1.45;">
                    Use quando já souber o código, parte da descrição ou uma palavra-chave.
                    Ideal para localizar rapidamente opções oficiais.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        if st.button("Usar busca direta", key="btn_modo_cnae_conhecido", use_container_width=True, type="primary" if selecionado else "secondary"):
            _set_modo_busca_cnae("conhecido")

    with col_inteligente:
        selecionado = modo_atual == "descobrir"
        borda = "#00866B" if selecionado else "#E3E8EF"
        fundo = "#F3FAF8" if selecionado else "#FFFFFF"
        st.markdown(
            f"""
            <div class="gaia-card" style="min-height: 150px; border: 1.6px solid {borda}; background:{fundo};">
                <div style="font-size:1.55rem; margin-bottom:8px;">✨</div>
                <div style="font-weight:800; font-size:1.02rem; margin-bottom:5px;">Descobrir CNAE pela atividade</div>
                <div style="color:#5A6679; font-size:.86rem; line-height:1.45;">
                    Use quando não souber o CNAE. Descreva a atividade com suas palavras.
                    O GAIA busca sugestões oficiais automaticamente.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        if st.button("Usar busca inteligente", key="btn_modo_cnae_descobrir", use_container_width=True, type="primary" if selecionado else "secondary"):
            _set_modo_busca_cnae("descobrir")


def _buscar_sugestoes_cnae_por_modo(texto, finalidade, modo, progresso=None):
    """Retorna sugestões para principal/secundário conforme o modo escolhido."""
    if modo == "conhecido":
        sugestoes = buscar_cnaes_conhecido(texto, limite=30)
        _registrar_origem_busca_cnae("direta")
        return sugestoes

    return sugerir_cnaes_inteligente(texto, progresso=progresso)


def _render_form_busca_cnae_principal(modo):
    if modo == "conhecido":
        titulo = "1. CNAE principal — busca direta"
        label = "Digite o código, parte da descrição ou palavra-chave do CNAE principal"
        placeholder = "Ex.: 5611, restaurante, limpeza, comércio varejista..."
        botao = "BUSCAR"
    else:
        titulo = "1. CNAE principal — busca inteligente"
        label = "Descreva apenas uma atividade principal"
        placeholder = "Ex.: loja de açaí, vender jujuba, instalar purificador de água..."
        botao = "BUSCAR SUGESTÕES"

    st.markdown(f"#### {titulo}")
    with st.form(f"form_busca_cnae_principal_wizard_{modo}", clear_on_submit=False):
        atividade_principal_txt = st.text_input(
            label,
            placeholder=placeholder,
            key=f"atividade_principal_cnae_{modo}"
        )
        buscar_principal = st.form_submit_button(botao, type="secondary", use_container_width=True)

    if buscar_principal:
        if not atividade_principal_txt.strip():
            msg_warning("Informe uma atividade, código ou palavra-chave para buscar o CNAE principal.")
            return

        if modo == "descobrir" and detectar_multiplas_atividades_principal(atividade_principal_txt):
            st.session_state["cnae_principal_sugestoes"] = []
            msg_warning("Informe somente a atividade principal neste campo. Expressões como 'e afins' são aceitas; atividades totalmente diferentes devem ser adicionadas como secundárias.")
            return

        progresso = None
        if modo == "descobrir":
            progresso = st.progress(0, text="Preparando sugestões. Na primeira busca, isso pode levar alguns instantes...")

        sugestoes = _buscar_sugestoes_cnae_por_modo(atividade_principal_txt, "principal", modo, progresso=progresso)
        st.session_state["cnae_principal_sugestoes"] = sugestoes
        _mostrar_origem_busca_cnae()
        _limpar_checks_cnae("check_cnae_principal_")

        if sugestoes:
            msg_success("Sugestões para atividade principal encontradas. Marque apenas uma.")
        else:
            msg_warning("Nenhum CNAE principal encontrado. Tente buscar por outro termo ou use o outro modo de busca.")


def _render_form_busca_cnaes_secundarios(modo):
    if modo == "conhecido":
        titulo = "2. CNAEs secundários — busca direta"
        label = "Digite código, descrição ou palavra-chave para buscar CNAEs secundários"
        placeholder = "Ex.: delivery, bebidas, cursos, manutenção..."
        botao = "BUSCAR"
    else:
        titulo = "2. CNAEs secundários — busca inteligente"
        label = "Descreva uma ou mais atividades secundárias"
        placeholder = "Ex.: entrega por delivery, venda de bebidas, cursos de dança..."
        botao = "BUSCAR SUGESTÕES"

    st.markdown(f"#### {titulo}")
    with st.form(f"form_busca_cnaes_secundarios_wizard_{modo}", clear_on_submit=False):
        atividade_secundaria_txt = st.text_input(
            label,
            placeholder=placeholder,
            key=f"atividade_secundaria_cnae_{modo}"
        )
        buscar_secundarios = st.form_submit_button(botao, type="secondary", use_container_width=True)

    if buscar_secundarios:
        if not atividade_secundaria_txt.strip():
            msg_warning("Informe uma atividade, código ou palavra-chave para buscar CNAEs secundários.")
            return

        progresso = None
        if modo == "descobrir":
            progresso = st.progress(0, text="Preparando sugestões. Na primeira busca, isso pode levar alguns instantes...")

        sugestoes = _buscar_sugestoes_cnae_por_modo(atividade_secundaria_txt, "secundario", modo, progresso=progresso)
        st.session_state["cnaes_secundarios_sugestoes"] = sugestoes
        _mostrar_origem_busca_cnae()
        _limpar_checks_cnae("check_cnae_secundario_")

        if sugestoes:
            msg_success("Sugestões para atividades secundárias encontradas. Marque uma ou mais opções.")
        else:
            msg_warning("Nenhum CNAE secundário encontrado. Tente buscar por outro termo ou use o outro modo de busca.")


def _render_sugestoes_cnae_principal():
    sugestoes_principal = st.session_state.get("cnae_principal_sugestoes", [])
    if not sugestoes_principal:
        return

    st.caption("Marque o CNAE principal. Se marcar outro, ele substitui o anterior.")
    principal_atual = st.session_state.get("cnae_principal_selecionado")
    secundarios_atuais = st.session_state.get("cnaes_secundarios_selecionados", [])
    modo = st.session_state.get("cnae_modo_busca") or "sem_modo"
    for idx, item in enumerate(sugestoes_principal, start=1):
        codigo_principal_atual = principal_atual.get("codigo") if principal_atual else "nenhum"
        chave = f"check_cnae_principal_{modo}_{idx}_{item['codigo']}_{codigo_principal_atual}"
        marcado = st.checkbox(
            f"{item['codigo']} — {item['descricao']}",
            value=_mesmo_cnae(principal_atual, item),
            key=chave
        )
        if marcado and not _mesmo_cnae(principal_atual, item):
            if _lista_contem_cnae(secundarios_atuais, item):
                secundarios_atuais = _remover_cnae_por_codigo(secundarios_atuais, item["codigo"])
                st.session_state["cnaes_secundarios_selecionados"] = secundarios_atuais
                set_flash("warning", "Esse CNAE estava como secundário. Ele foi removido das secundárias e definido como principal.")
            st.session_state["cnae_principal_selecionado"] = item
            st.rerun()
        if not marcado and _mesmo_cnae(principal_atual, item):
            st.session_state["cnae_principal_selecionado"] = None
            st.rerun()


def _render_sugestoes_cnaes_secundarios():
    sugestoes_secundarias = st.session_state.get("cnaes_secundarios_sugestoes", [])
    if not sugestoes_secundarias:
        return

    cnae_principal_atual = st.session_state.get("cnae_principal_selecionado")
    secundarios_selecionados = list(st.session_state.get("cnaes_secundarios_selecionados", []))
    modo = st.session_state.get("cnae_modo_busca") or "sem_modo"
    st.caption("Marque os CNAEs secundários que devem entrar no objeto social.")
    for idx, item in enumerate(sugestoes_secundarias, start=1):
        if _mesmo_cnae(cnae_principal_atual, item):
            st.warning(f"{item['codigo']} já está selecionado como CNAE principal. Não pode ser adicionado como secundário.")
            continue
        marcado = st.checkbox(
            f"{item['codigo']} — {item['descricao']}",
            value=_lista_contem_cnae(secundarios_selecionados, item),
            key=f"check_cnae_secundario_{modo}_{idx}_{item['codigo']}"
        )
        if marcado and not _lista_contem_cnae(secundarios_selecionados, item):
            secundarios_selecionados = _adicionar_cnae_unico(secundarios_selecionados, item)
            st.session_state["cnaes_secundarios_selecionados"] = secundarios_selecionados
            st.rerun()
        if not marcado and _lista_contem_cnae(secundarios_selecionados, item):
            secundarios_selecionados = _remover_cnae_por_codigo(secundarios_selecionados, item["codigo"])
            st.session_state["cnaes_secundarios_selecionados"] = secundarios_selecionados
            st.rerun()
    st.session_state["cnaes_secundarios_selecionados"] = secundarios_selecionados


def _render_cnaes_selecionados():
    cnae_principal_final = st.session_state.get("cnae_principal_selecionado")
    cnaes_secundarios_final = st.session_state.get("cnaes_secundarios_selecionados", [])

    st.markdown("#### 3. CNAEs selecionados")
    if cnae_principal_final:
        col_principal, col_remover_principal = st.columns([5, 1])
        col_principal.markdown(f"**Principal:** {cnae_principal_final['codigo']} — {cnae_principal_final['descricao']}")
        if col_remover_principal.button("Remover", type="secondary", key="remover_cnae_principal"):
            st.session_state["cnae_principal_selecionado"] = None
            _limpar_checks_cnae("check_cnae_principal_")
            st.rerun()
    else:
        st.info("Nenhum CNAE principal selecionado.")

    if cnaes_secundarios_final:
        st.markdown("**Secundários:**")
        for item in cnaes_secundarios_final:
            col_secundario, col_remover_secundario = st.columns([5, 1])
            col_secundario.markdown(f"- **{item['codigo']}** — {item['descricao']}")
            if col_remover_secundario.button("Remover", type="secondary", key=f"remover_cnae_secundario_{item['codigo']}"):
                st.session_state["cnaes_secundarios_selecionados"] = _remover_cnae_por_codigo(cnaes_secundarios_final, item["codigo"])
                for chave in list(st.session_state.keys()):
                    if str(chave).startswith("check_cnae_secundario_") and item["codigo"] in str(chave):
                        del st.session_state[chave]
                st.rerun()
    else:
        st.caption("Nenhum CNAE secundário selecionado.")

    objeto = montar_objeto_social_por_cnaes(cnae_principal_final, cnaes_secundarios_final)
    cnaes_texto = montar_texto_cnaes_classificados(cnae_principal_final, cnaes_secundarios_final)
    st.session_state["ia_objeto_sugerido"] = objeto
    st.session_state["ia_cnaes_texto_classificado"] = cnaes_texto
    return objeto, cnaes_texto


def _render_busca_cnaes_wizard():
    st.markdown("### 🔎 Localização de CNAE")
    total_cnaes = len(carregar_cnaes_local())
    if not total_cnaes:
        msg_warning("Base de CNAEs não encontrada. Coloque o arquivo em dados/cnaes.csv.")

    _render_cards_modo_busca_cnae()

    modo = st.session_state.get("cnae_modo_busca")
    if not modo:
        return _render_cnaes_selecionados()

    if modo == "descobrir":
        nivel_atual = _nivel_embedding_atual()
        opcoes_nivel = ["Preciso", "Equilibrado", "Amplo"]
        nivel_escolhido = st.radio(
            "Nível de abrangência da busca",
            opcoes_nivel,
            index=opcoes_nivel.index(nivel_atual),
            horizontal=True,
            key="radio_nivel_embedding_cnae",
            help="Ajusta a quantidade de sugestões. Use Preciso quando quiser menos opções; use Amplo quando a atividade for curta ou difícil de descrever."
        )
        st.session_state["cnae_nivel_embedding"] = nivel_escolhido
        cfg = NIVEIS_EMBEDDING_CNAE[nivel_escolhido]
        st.caption(
            f"{cfg['descricao']} Resultados finais: até {cfg['limite_final']}."
        )

    _render_form_busca_cnae_principal(modo)
    _render_sugestoes_cnae_principal()

    st.markdown("---")

    _render_form_busca_cnaes_secundarios(modo)
    _render_sugestoes_cnaes_secundarios()

    st.markdown("---")

    return _render_cnaes_selecionados()


def _gerar_documento_wizard(eventos_para_processar, emp_row, data_documento):
    clausulas_prontas = processar_clausulas_prototype(eventos_para_processar)

    with get_db() as conn:
        socios_rows = conn.execute(
            "SELECT nome, nacionalidade, cpf, rg, orgao_emissor, estado_civil, data_nasc, profissao, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_socio WHERE id_empresa=?",
            (st.session_state.empresa_id,)
        ).fetchall()

    empresa = {
        'cnpj': emp_row[0] if emp_row else '', 'nome_empresarial': emp_row[1] if emp_row else '',
        'nome_fantasia': emp_row[2] if emp_row else '', 'nire': emp_row[3] if emp_row else '',
        'logradouro': emp_row[4] if emp_row else '', 'numero': emp_row[5] if emp_row else '',
        'complemento': emp_row[6] if emp_row else '', 'bairro': emp_row[7] if emp_row else '',
        'cidade': emp_row[8] if emp_row else '', 'uf': emp_row[9] if emp_row else '',
        'cep': emp_row[10] if emp_row else ''
    }
    socios = [{'nome': r[0], 'nacionalidade': r[1], 'cpf': r[2], 'rg': r[3], 'orgao_emissor': r[4], 'estado_civil': r[5], 'data_nasc': r[6], 'profissao': r[7], 'logradouro': r[8], 'numero': r[9], 'complemento': r[10], 'bairro': r[11], 'cidade': r[12], 'uf': r[13], 'cep': r[14]} for r in socios_rows]

    data_documento_texto = formatar_data_documento(data_documento)
    doc_bytes = gerar_docx_modular(empresa, socios, clausulas_prontas, data_documento_texto)
    nome_emp = sanitizar_nome_arquivo(empresa.get('nome_empresarial', 'EMPRESA'))
    nome_arquivo_docx = f"Alteracao_Contratual_{nome_emp}.docx"
    nome_arquivo_pdf = f"Alteracao_Contratual_{nome_emp}.pdf"
    pdf_bytes, erro_pdf = gerar_pdf_modular(empresa, socios, clausulas_prontas, data_documento_texto)

    observacao_hist = None if pdf_bytes else erro_pdf
    with get_db() as conn:
        conn.execute("""
            INSERT INTO tb_documento (
                id_empresa, id_usuario, data_documento_texto, tipo_documento, status, eventos_json, clausulas_json,
                nome_arquivo_docx, nome_arquivo_pdf, nome_arquivo_relatorio, modelo_documento,
                docx_blob, pdf_blob, relatorio_conferencia_blob, observacao
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            st.session_state.empresa_id,
            st.session_state.user_id,
            data_documento_texto,
            "Alteração Contratual Múltipla",
            "Gerado" if pdf_bytes else "Gerado parcialmente",
            json.dumps(eventos_para_processar, ensure_ascii=False),
            json.dumps(clausulas_prontas, ensure_ascii=False),
            nome_arquivo_docx,
            nome_arquivo_pdf if pdf_bytes else None,
            None,
            "Padrão Junta Comercial",
            sqlite3.Binary(doc_bytes.getvalue()),
            sqlite3.Binary(pdf_bytes.getvalue()) if pdf_bytes else None,
            None,
            observacao_hist
        ))
        conn.commit()

    st.session_state.documento_gerado = {
        "empresa_id": st.session_state.empresa_id,
        "eventos": [e.get("tipo") for e in eventos_para_processar],
        "clausulas_prontas": clausulas_prontas,
        "data_documento_texto": data_documento_texto,
        "nome_arquivo_docx": nome_arquivo_docx,
        "nome_arquivo_pdf": nome_arquivo_pdf if pdf_bytes else None,
        "docx_bytes": doc_bytes.getvalue(),
        "pdf_bytes": pdf_bytes.getvalue() if pdf_bytes else None,
        "erro_pdf": erro_pdf if not pdf_bytes else None,
    }
    if pdf_bytes:
        set_flash("success", "✅ Documento gerado! Baixe o WORD (DOCX) ou o PDF abaixo.")
    else:
        set_flash("warning", "Documento WORD (DOCX) gerado. O PDF não foi gerado neste ambiente. Baixe o WORD (DOCX) abaixo.")
    st.session_state['_rolar_para_downloads'] = True
    st.rerun()


def tela_orquestrador_dados():
    cabecalho()
    st.subheader("Preenchimento dos Dados")

    selecionados = st.session_state.get('eventos_selecionados_lista', [])
    if not selecionados:
        msg_error("Nenhum evento selecionado.")
        if st.button("← Voltar", type="secondary", key="voltar_orquestrador_dados_vazio"):
            nav('orquestrador')
        return

    _inicializar_wizard_eventos(selecionados)
    idx = int(st.session_state.get("wizard_evento_idx", 0))
    total = len(selecionados)

    with get_db() as conn:
        emp_row = conn.execute(
            "SELECT cnpj, nome_empresarial, nome_fantasia, nire, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_empresa WHERE id=?",
            (st.session_state.empresa_id,)
        ).fetchone()

    if idx >= total:
        st.markdown("### Revisão final")
        eventos_para_processar = _eventos_processados_wizard(selecionados)
        for evento in eventos_para_processar:
            with st.expander(f"{evento.get('tipo')} - {_nome_evento_wizard(evento.get('tipo'))}", expanded=True):
                _render_resumo_evento_wizard(evento)
        data_documento = st.date_input("Data do documento*", value=date.today(), format="DD/MM/YYYY")
        c1, c2, c3 = st.columns(3)
        if c1.button("← Voltar evento", type="secondary", use_container_width=True):
            st.session_state["wizard_evento_idx"] = max(total - 1, 0)
            st.rerun()
        if c2.button("Editar seleção de eventos", type="secondary", use_container_width=True):
            nav('orquestrador')
        if c3.button("Gerar documento", type="primary", use_container_width=True):
            if len(eventos_para_processar) != total:
                msg_error("Há eventos selecionados sem preenchimento. Volte e conclua todas as etapas.")
                return
            _gerar_documento_wizard(eventos_para_processar, emp_row, data_documento)
        renderizar_documento_gerado()
        return

    codigo = selecionados[idx]
    _render_progresso_wizard(selecionados, idx)
    st.markdown(f"### {_nome_evento_wizard(codigo)}")
    salvo = _evento_wizard_salvo(codigo)

    def botoes_wizard():
        seguir = st.form_submit_button("Seguir →", type="primary", use_container_width=True, key=f"btn_seguir_wizard_{codigo}")
        voltar = False
        return voltar, seguir

    if codigo == "220":
        st.markdown(f"<div class='gaia-prev-value'><strong>Nome anterior:</strong> {html.escape(emp_row[1] or 'Não informado')}</div>", unsafe_allow_html=True)
        with st.form(f"form_wizard_{codigo}"):
            novo_nome = padronizar_texto(st.text_input("Nome posterior*", value=salvo.get("dados", {}).get("novo_nome", "")))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not novo_nome:
                msg_error("Informe o nome posterior.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "220", "dados": {"novo_nome": novo_nome}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo == "221":
        st.markdown(f"<div class='gaia-prev-value'><strong>Nome fantasia anterior:</strong> {html.escape(emp_row[2] or 'Não informado')}</div>", unsafe_allow_html=True)
        with st.form(f"form_wizard_{codigo}"):
            nome_fantasia = padronizar_texto(st.text_input("Nome fantasia posterior*", value=salvo.get("dados", {}).get("nome_fantasia", "")))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not nome_fantasia:
                msg_error("Informe o nome fantasia posterior.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "221", "dados": {"nome_fantasia": nome_fantasia}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo in ["211", "209", "210"]:
        endereco_anterior = ", ".join([x for x in [
            emp_row[4] or "", f"nº {emp_row[5]}" if emp_row[5] else "", emp_row[6] or "", emp_row[7] or "",
            f"{emp_row[8]} - {emp_row[9]}" if emp_row[8] or emp_row[9] else "", f"CEP {emp_row[10]}" if emp_row[10] else "",
        ] if x]) or "Não informado"
        st.markdown(f"<div class='gaia-prev-value'><strong>Endereço anterior:</strong> {html.escape(endereco_anterior)}</div>", unsafe_allow_html=True)
        dados_salvos = salvo.get("dados", {})
        prefixo_cep = f"wizard_endereco_{codigo}"
        with st.form(f"form_busca_cep_{codigo}", clear_on_submit=False):
            c_cep_busca, c_btn_busca = st.columns([3, 1], vertical_alignment="bottom")
            cep_busca = c_cep_busca.text_input(
                "Buscar endereço por CEP",
                value=st.session_state.get(f"{prefixo_cep}_cep", dados_salvos.get("cep", "")),
                placeholder="00000-000",
                key=f"{prefixo_cep}_cep_busca"
            )
            buscar_cep = c_btn_busca.form_submit_button("BUSCAR CEP", type="secondary", use_container_width=True)

        if buscar_cep:
            dados_cep, erro_cep = consultar_cep_viacep(cep_busca)
            if erro_cep:
                if "ViaCEP indisponível" in erro_cep:
                    msg_api_offline("ViaCEP")
                else:
                    msg_warning(erro_cep)
            else:
                st.session_state[f"{prefixo_cep}_cep"] = dados_cep.get("cep", "")
                st.session_state[f"{prefixo_cep}_logradouro"] = dados_cep.get("logradouro", "")
                st.session_state[f"{prefixo_cep}_bairro"] = dados_cep.get("bairro", "")
                st.session_state[f"{prefixo_cep}_cidade"] = dados_cep.get("cidade", "")
                st.session_state[f"{prefixo_cep}_uf"] = dados_cep.get("uf", "")
                set_flash("success", "Endereço encontrado pelo CEP. Confira o número antes de seguir.")
                st.rerun()

        with st.form(f"form_wizard_{codigo}"):
            c1, c2 = st.columns([3, 1])
            logradouro = padronizar_texto(c1.text_input("Logradouro posterior*", value=st.session_state.get(f"{prefixo_cep}_logradouro", dados_salvos.get("logradouro", ""))))
            numero = remover_espacos_extras(c2.text_input("Número posterior*", value=dados_salvos.get("numero", "")))
            c3, c4 = st.columns(2)
            complemento = padronizar_texto(c3.text_input("Complemento posterior", value=dados_salvos.get("complemento", "")))
            bairro = padronizar_texto(c4.text_input("Bairro posterior*", value=st.session_state.get(f"{prefixo_cep}_bairro", dados_salvos.get("bairro", ""))))
            c5, c6, c7 = st.columns([2, 1, 2])
            cidade = padronizar_texto(c5.text_input("Cidade posterior*", value=st.session_state.get(f"{prefixo_cep}_cidade", dados_salvos.get("cidade", ""))))
            uf = padronizar_texto(c6.text_input("UF posterior*", value=st.session_state.get(f"{prefixo_cep}_uf", dados_salvos.get("uf", "")), max_chars=2))
            cep = formatar_cep(c7.text_input("CEP posterior*", value=st.session_state.get(f"{prefixo_cep}_cep", dados_salvos.get("cep", "")), placeholder="00000-000"))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not logradouro or not numero or not bairro or not cidade or not uf or not cep:
                msg_error("Preencha todos os campos obrigatórios do endereço.")
            elif not validar_uf(uf):
                msg_error("UF posterior inválida.")
            elif not validar_cep(cep):
                msg_error("CEP posterior inválido.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": codigo, "dados": {"logradouro": logradouro, "numero": numero, "complemento": complemento, "bairro": bairro, "cidade": cidade, "uf": uf, "cep": formatar_cep(cep)}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo == "244":
        st.markdown("<div class='gaia-prev-value'><strong>Atividades anteriores:</strong> consulte o contrato vigente ou o cadastro da empresa antes de confirmar.</div>", unsafe_allow_html=True)
        objeto_sugerido, cnaes_texto = _render_busca_cnaes_wizard()
        dados_salvos = salvo.get("dados", {})
        with st.form(f"form_wizard_{codigo}"):
            objeto = padronizar_texto_livre(st.text_area("Objeto social posterior*", value=dados_salvos.get("objeto_descricao") or objeto_sugerido))
            cnaes = padronizar_texto_livre(st.text_area("CNAEs posteriores*", value=dados_salvos.get("cnaes_lista") or cnaes_texto))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not st.session_state.get("cnae_principal_selecionado"):
                msg_error("Selecione um CNAE principal.")
            elif not objeto or not cnaes:
                msg_error("Preencha o objeto social e os CNAEs posteriores.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "244", "dados": {"objeto_descricao": objeto, "cnaes_lista": cnaes}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo == "248":
        with st.form(f"form_wizard_{codigo}"):
            tipo_unidade = padronizar_texto(st.text_input("Tipo de unidade posterior*", value=salvo.get("dados", {}).get("tipo_unidade", "")))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not tipo_unidade:
                msg_error("Informe o tipo de unidade posterior.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "248", "dados": {"tipo_unidade": tipo_unidade}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo == "249":
        with st.form(f"form_wizard_{codigo}"):
            forma_atuacao = padronizar_texto(st.text_input("Forma de atuação posterior*", value=salvo.get("dados", {}).get("forma_atuacao", "")))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not forma_atuacao:
                msg_error("Informe a forma de atuação posterior.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "249", "dados": {"forma_atuacao": forma_atuacao}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo == "225":
        st.markdown(f"<div class='gaia-prev-value'><strong>Empresa anterior:</strong> {html.escape(emp_row[1] or 'Não informado')}</div>", unsafe_allow_html=True)
        with st.form(f"form_wizard_{codigo}"):
            novo_nome = padronizar_texto(st.text_input("Nome da sociedade posterior*", value=salvo.get("dados", {}).get("novo_nome", "")))
            voltar, seguir = botoes_wizard()
        if seguir:
            if not novo_nome:
                msg_error("Informe o nome da sociedade posterior.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "225", "dados": {"novo_nome": novo_nome}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo == "capital":
        with st.form(f"form_wizard_{codigo}"):
            capital = st.text_input("Capital posterior (R$)*", value=salvo.get("dados", {}).get("novo_capital", ""))
            voltar, seguir = botoes_wizard()
        if seguir:
            valor_ok, valor_formatado = formatar_valor_brasileiro(capital)
            if not capital:
                msg_error("Informe o capital social posterior.")
            elif not valor_ok:
                msg_error("Capital social inválido.")
            else:
                _salvar_evento_wizard(codigo, {"tipo": "capital", "dados": {"novo_capital": valor_formatado}})
                st.session_state["wizard_evento_idx"] = idx + 1
                st.rerun()

    elif codigo in ["052", "999"]:
        st.info("Este evento é processado automaticamente. Clique em seguir para incluir no documento.")
        with st.form(f"form_wizard_{codigo}"):
            voltar, seguir = botoes_wizard()
        if seguir:
            _salvar_evento_wizard(codigo, {"tipo": codigo, "dados": {}})
            st.session_state["wizard_evento_idx"] = idx + 1
            st.rerun()

    if st.button("← Voltar", type="secondary", use_container_width=True, key=f"btn_voltar_wizard_externo_{idx}_{codigo}"):
        if idx <= 0:
            nav('orquestrador')
        else:
            st.session_state["wizard_evento_idx"] = idx - 1
            st.rerun()

    renderizar_documento_gerado()

def renderizar_sidebar():
    """Menu lateral com navegação persistente para a área autenticada."""
    if not st.session_state.get('user_id'):
        return

    nome = html.escape(str(st.session_state.get('nome_usuario') or 'Usuário'))
    perfil = html.escape(str(st.session_state.get('user_role') or ''))

    with st.sidebar:
        inicial = (nome.strip()[:1] or "U").upper()
        st.markdown(
            f'<div class="gaia-side-user"><div class="gaia-side-avatar">{inicial}</div>'
            f'<div><b>{nome}</b><br><span>{perfil}</span></div></div>',
            unsafe_allow_html=True,
        )

        st.markdown('<div class="gaia-side-section">Principal</div>', unsafe_allow_html=True)
        if st.button('🏠  Início', use_container_width=True, key='side_menu_home', disabled=st.session_state.pagina == 'menu'):
            nav('menu')
        if st.button('↪  Continuar alteração', use_container_width=True, key='side_menu_continuar', disabled=st.session_state.pagina == 'buscar_empresa'):
            nav('buscar_empresa')
        if st.button('＋  Nova empresa', use_container_width=True, key='side_menu_nova', disabled=st.session_state.pagina == 'cad_empresa'):
            nav('cad_empresa')

        st.markdown('<div class="gaia-side-section">Documentos</div>', unsafe_allow_html=True)
        if st.button('📋  Histórico', use_container_width=True, key='side_menu_hist', disabled=st.session_state.pagina == 'historico'):
            nav('historico')

        if st.session_state.get('user_role') == 'Admin':
            st.markdown('<div class="gaia-side-section">Administração</div>', unsafe_allow_html=True)
            if st.button('👥  Usuários', use_container_width=True, key='side_menu_users', disabled=st.session_state.pagina == 'gestao_usuarios'):
                nav('gestao_usuarios')
            if st.button('🔑  Minha senha', use_container_width=True, key='side_menu_senha', disabled=st.session_state.pagina == 'alterar_senha'):
                nav('alterar_senha')

        st.markdown('<div class="gaia-side-section">Sessão</div>', unsafe_allow_html=True)
        if st.button('⎋  Sair', use_container_width=True, key='side_menu_sair'):
            for key in ['user_id', 'user_role', 'login_usuario', 'empresa_id', 'socio_id']:
                st.session_state[key] = None
            st.session_state['_alterar_senha_origem_aviso'] = False
            nav('login')



def aquecer_ollama_em_segundo_plano():
    """Pré-aquece a IA local de embeddings sem travar a interface.

    Ao abrir o GAIA, esta função tenta carregar o modelo de embedding e,
    principalmente, iniciar a criação do índice semântico da base CNAE em
    segundo plano. Assim, quando o usuário chegar ao evento 244, a busca
    inteligente já pode estar pronta ou em preparação.
    """
    def _worker_aquecimento_leve():
        try:
            # Simula uma primeira solicitação real ao modelo de embeddings.
            # Essas consultas são genéricas e servem só para carregar o modelo
            # e validar o caminho de busca; não viram regra de CNAE.
            _chamar_ollama_embed(["comércio varejista", "prestação de serviço"], timeout=8)
            executar_buscas_padrao_aquecimento_cnae()
        except Exception:
            pass

    try:
        threading.Thread(target=_worker_aquecimento_leve, daemon=True).start()
    except Exception:
        pass

    preaquecer_indice_semantico_cnae_em_segundo_plano()

# ==========================================
# ROTEADOR DE PÁGINAS
# ==========================================
paginas = {
    'login': tela_login, 'registrar': tela_registrar,
    'aviso_senha_padrao': tela_aviso_senha_padrao, 'alterar_senha': tela_alterar_senha,
    'menu': tela_menu,
    'gestao_usuarios': tela_gestao_usuarios, 'cad_empresa': tela_cad_empresa,
    'editar_empresa': tela_editar_empresa, 'buscar_empresa': tela_buscar_empresa,
    'lista_socios': tela_socios, 'cad_socio': tela_cad_socio,
    'editar_socio': tela_editar_socio, 'orquestrador': tela_orquestrador,
    'orquestrador_dados': tela_orquestrador_dados,
    'historico': tela_historico, 'historico_empresa': tela_historico_empresa,
}

aquecer_ollama_em_segundo_plano()
renderizar_sidebar()
paginas.get(st.session_state.pagina, tela_login)()
