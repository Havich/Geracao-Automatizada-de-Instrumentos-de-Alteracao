import streamlit as st
import streamlit.components.v1 as components
import sqlite3
import hashlib
import io
import copy
import requests
import tempfile
import subprocess
import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ==========================================
# CONFIGURAÇÃO DE LAYOUT E ESTILOS
# ==========================================
st.set_page_config(page_title="GAIA", page_icon="logo.png", layout="centered")

st.markdown("""
    <style>
    header[data-testid="stHeader"] { display: none !important; }
    .block-container { padding-top: 5rem; padding-bottom: 1rem; }
    .stApp { background-color: #FFFFFF; }
    h1, h2, h3 { color: #006400 !important; text-align: center; font-family: 'Arial'; margin-top: 0px; margin-bottom: 5px;}
    [data-testid="stImage"] { display: flex; justify-content: center; }
    
    /* Forçar cursor de pointer no selectbox */
    div[data-baseweb="select"] > div, div[data-baseweb="select"] input {
        cursor: pointer !important;
    }

    button[kind="primary"], button[data-testid="baseButton-primary"],
    button[kind="primaryFormSubmit"], button[data-testid="baseButton-primaryFormSubmit"] {
        background-color: #28a745 !important; border-color: #28a745 !important; color: white !important;
    }
    button[kind="secondary"], button[data-testid="baseButton-secondary"],
    button[kind="secondaryFormSubmit"], button[data-testid="baseButton-secondaryFormSubmit"] {
        background-color: #6c757d !important; border-color: #6c757d !important; color: white !important;
    }
    .stButton>button, .stFormSubmitButton>button {
        border-radius: 8px; width: 100%; height: 38px; font-weight: bold; transition: 0.3s;
    }
    hr { margin: 8px 0px; border-color: #e0e0e0; }
    .title-separator { border-top: 2px solid #006400; margin-top: 20px; margin-bottom: 10px;}
    
    input:disabled, textarea:disabled { background-color: #f0f2f6 !important; color: #495057 !important; }
    </style>
""", unsafe_allow_html=True)

def scroll_to_top():
    """Força a página a rolar para o topo, para que mensagens de erro/sucesso fiquem visíveis."""
    components.html(
        """
        <script>
        setTimeout(function() {
            window.parent.scrollTo({top: 0, behavior: 'smooth'});
        }, 80);
        </script>
        """,
        height=0,
        width=0,
    )


def set_flash(tipo, texto):
    """Guarda mensagem para aparecer após st.rerun()/troca de tela."""
    st.session_state['_flash_msg'] = {'tipo': tipo, 'texto': texto}


def show_flash():
    """Mostra mensagens persistentes no topo da página.

    Corrige o problema de mensagens que desapareciam porque o sistema chamava
    st.rerun() logo após st.success()/st.error().
    """
    flash = st.session_state.pop('_flash_msg', None)
    if not flash:
        return
    tipo = flash.get('tipo', 'info')
    texto = flash.get('texto', '')
    if tipo == 'success':
        st.success(texto)
        st.toast(texto, icon='✅')
    elif tipo == 'error':
        st.error(texto)
        st.toast(texto, icon='❌')
    elif tipo == 'warning':
        st.warning(texto)
        st.toast(texto, icon='⚠️')
    else:
        st.info(texto)
        st.toast(texto, icon='ℹ️')
    scroll_to_top()


def msg_success(texto):
    st.success(texto)
    st.toast(texto, icon='✅')
    scroll_to_top()


def msg_error(texto):
    st.error(texto)
    st.toast(texto, icon='❌')
    scroll_to_top()


def msg_warning(texto):
    st.warning(texto)
    st.toast(texto, icon='⚠️')
    scroll_to_top()


def limpar_api(prefixo):
    for k in list(st.session_state.keys()):
        if k.startswith(prefixo):
            del st.session_state[k]


def cabecalho():
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        try:
            st.image("logo.png", width=400)
        except:
            st.markdown("<h2>🌿 GAIA</h2>", unsafe_allow_html=True)
    show_flash()

# ==========================================
# SPRINT 4 - Modelagem e Banco de Dados
# ==========================================
def get_db():
    conn = sqlite3.connect('gaia.db')
    conn.execute("PRAGMA foreign_keys = ON;")
    return conn

def inicializar_banco():
    with get_db() as conn:
        conn.execute("CREATE TABLE IF NOT EXISTS tb_perfil_acesso (id INTEGER PRIMARY KEY, nome_perfil TEXT UNIQUE NOT NULL)")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_usuario (
            id INTEGER PRIMARY KEY AUTOINCREMENT, nome TEXT, login TEXT UNIQUE,
            senha_hash TEXT, id_perfil INTEGER,
            FOREIGN KEY (id_perfil) REFERENCES tb_perfil_acesso (id))""")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_empresa (
            id INTEGER PRIMARY KEY AUTOINCREMENT, cnpj TEXT UNIQUE, nome_empresarial TEXT,
            nome_fantasia TEXT, nire TEXT, cidade TEXT, uf TEXT, cep TEXT, logradouro TEXT, numero TEXT,
            bairro TEXT, complemento TEXT)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_socio (
            id INTEGER PRIMARY KEY AUTOINCREMENT, id_empresa INTEGER, nome TEXT,
            nacionalidade TEXT, cpf TEXT, rg TEXT, orgao_emissor TEXT,
            estado_civil TEXT, data_nasc TEXT, profissao TEXT, cep TEXT,
            logradouro TEXT, numero TEXT, complemento TEXT, bairro TEXT, cidade TEXT, uf TEXT,
            FOREIGN KEY (id_empresa) REFERENCES tb_empresa (id) ON DELETE CASCADE)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS tb_documento (
            id INTEGER PRIMARY KEY AUTOINCREMENT, id_empresa INTEGER, id_usuario INTEGER,
            data_geracao DATETIME DEFAULT CURRENT_TIMESTAMP, tipo_documento TEXT, status TEXT DEFAULT 'Gerado',
            FOREIGN KEY (id_empresa) REFERENCES tb_empresa (id) ON DELETE CASCADE,
            FOREIGN KEY (id_usuario) REFERENCES tb_usuario (id))""")
        conn.commit()

        if conn.execute("SELECT COUNT(*) FROM tb_perfil_acesso").fetchone()[0] == 0:
            conn.execute("INSERT INTO tb_perfil_acesso (nome_perfil) VALUES ('Admin'), ('Operador')")
            conn.commit()

        if conn.execute("SELECT COUNT(*) FROM tb_usuario").fetchone()[0] == 0:
            h = hashlib.sha256("1234".encode()).hexdigest()
            conn.execute("INSERT INTO tb_usuario (nome, login, senha_hash, id_perfil) VALUES (?,?,?,1)", ("Administrador", "admin", h))
            conn.commit()

inicializar_banco()

# ==========================================
# ESTADO E NAVEGAÇÃO
# ==========================================
defaults = {
    'pagina': 'login', 'empresa_id': None, 'socio_id': None, 'user_id': None,
    'nome_usuario': None, 'user_role': None, 'eventos_selecionados_lista': [],
    'mostrar_confirmacao_repetidos': False, 'msg_aviso_texto': ''
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v

def gerar_hash(s): return hashlib.sha256(s.encode()).hexdigest()

def nav(p):
    st.session_state.pagina = p
    scroll_to_top()
    st.rerun()

def _form_socio(titulo, defaults=None):
    d = defaults or {}
    estados_civis = ["", "Solteiro(a)", "Casado(a)", "Divorciado(a)", "Viúvo(a)", "União Estável"]
    cabecalho()
    st.subheader(titulo)

    with st.form("f_socio"):
        st.markdown("**Dados Pessoais**")
        nome = st.text_input("NOME COMPLETO*", value=d.get('nome') or '')       
        c1, c2 = st.columns(2)
        nac = c1.text_input("NACIONALIDADE*", value=d.get('nacionalidade') or 'BRASILEIRO(A)')
        profissao = c2.text_input("PROFISSÃO*", value=d.get('profissao') or 'EMPRESÁRIO')
        
        c3, c4 = st.columns(2)
        cpf = c3.text_input("CPF*", value=d.get('cpf') or '')
        data_nasc = c4.text_input("DATA NASCIMENTO (DD/MM/AAAA)*", value=d.get('data_nasc') or '')
        
        c5, c6 = st.columns(2)
        rg = c5.text_input("RG*", value=d.get('rg') or '')
        orgao = c6.text_input("ÓRGÃO EMISSOR*", value=d.get('orgao_emissor') or '')
        
        ec_valor = d.get('estado_civil') or ""
        ec_idx = estados_civis.index(ec_valor) if ec_valor in estados_civis else 0
        estado_civil = st.selectbox("ESTADO CIVIL*", estados_civis, index=ec_idx)
        
        st.markdown("<hr>", unsafe_allow_html=True)
        st.markdown("**Morada**")
        
        c_cep, c_btn_cep = st.columns([3, 1], vertical_alignment="bottom")
        cep_valor = st.session_state.get('api_socio_cep', d.get('cep') or '')
        cep = c_cep.text_input("CEP*", value=cep_valor)
        btn_buscar_cep = c_btn_cep.form_submit_button("BUSCAR", type="secondary", use_container_width=True)
        
        if btn_buscar_cep:
            cep_limpo = ''.join(filter(str.isdigit, cep))
            if len(cep_limpo) == 8:
                try:
                    resp = requests.get(f"https://viacep.com.br/ws/{cep_limpo}/json/", timeout=5)
                    if resp.status_code == 200:
                        dados_cep = resp.json()
                        if "erro" not in dados_cep:
                            st.session_state['api_socio_cep'] = dados_cep.get('cep', '')
                            st.session_state['api_socio_logradouro'] = dados_cep.get('logradouro', '').upper()
                            st.session_state['api_socio_bairro'] = dados_cep.get('bairro', '').upper()
                            st.session_state['api_socio_cidade'] = dados_cep.get('localidade', '').upper()
                            st.session_state['api_socio_uf'] = dados_cep.get('uf', '').upper()
                            st.rerun()
                        else:
                            msg_error("CEP não encontrado.")
                except Exception:
                    msg_error("Erro ao consultar a API de CEP.")
            else:
                msg_warning("O CEP deve conter 8 dígitos.")

        c7, c8 = st.columns([3, 1])
        v_log = st.session_state.get('api_socio_logradouro', d.get('logradouro') or '')
        logradouro = c7.text_input("LOGRADOURO*", value=v_log)
        numero = c8.text_input("NÚMERO*", value=d.get('numero') or '')
        
        c9, c10 = st.columns(2)
        complemento = c9.text_input("COMPLEMENTO", value=d.get('complemento') or '')
        v_bai = st.session_state.get('api_socio_bairro', d.get('bairro') or '')
        bairro = c10.text_input("BAIRRO*", value=v_bai)
        
        c11, c12 = st.columns([3, 1])
        v_cid = st.session_state.get('api_socio_cidade', d.get('cidade') or '')
        cidade = c11.text_input("CIDADE*", value=v_cid)
        v_uf = st.session_state.get('api_socio_uf', d.get('uf') or '')
        uf = c12.text_input("UF*", value=v_uf)
        
        st.markdown("<br>", unsafe_allow_html=True)
        b1, b2 = st.columns(2)
        btn_v = b1.form_submit_button("⬅️ CANCELAR", type="secondary", use_container_width=True)
        btn_ok = b2.form_submit_button("💾 CONFIRMAR", type="primary", use_container_width=True)

    if btn_v or btn_ok:
        for k in list(st.session_state.keys()):
            if k.startswith('api_socio_'):
                del st.session_state[k]

    return btn_v, btn_ok, dict(
        nome=(nome or "").upper(), nacionalidade=(nac or "").upper(), cpf=cpf, rg=rg,
        orgao_emissor=(orgao or "").upper(), estado_civil=estado_civil, data_nasc=data_nasc,
        profissao=(profissao or "").upper(), logradouro=(logradouro or "").upper(),
        numero=numero, complemento=(complemento or "").upper(), bairro=(bairro or "").upper(),
        cidade=(cidade or "").upper(), uf=(uf or "").upper(), cep=cep
    )

# ==========================================
# TELAS DO SISTEMA
# ==========================================
def tela_login():
    cabecalho()
    st.subheader("Acesso ao sistema")
    with st.form("form_login"):
        login = st.text_input("Login")
        senha = st.text_input("Senha", type="password")
        btn_e = st.form_submit_button("ENTRAR", type="primary", use_container_width=True)
        btn_cad = st.form_submit_button("Ainda não tem conta? CADASTRE-SE", type="secondary", use_container_width=True)
    
    if btn_e:
        with get_db() as conn:
            user = conn.execute("""
                SELECT u.id, u.nome, p.nome_perfil FROM tb_usuario u
                JOIN tb_perfil_acesso p ON u.id_perfil = p.id
                WHERE u.login=? AND u.senha_hash=?
            """, (login, gerar_hash(senha))).fetchone()
        if user:
            st.session_state.user_id = user[0]
            st.session_state.nome_usuario = user[1]
            st.session_state.user_role = user[2]
            set_flash('success', f'Bem-vindo(a), {user[1]}!')
            nav('menu')
        else:
            msg_error("Login ou senha incorretos.")
    if btn_cad: nav('registrar')

def tela_registrar():
    cabecalho()
    st.subheader("Criar nova conta")
    with st.form("form_reg"):
        nome_completo = st.text_input("Nome Completo (Opcional)")
        u = st.text_input("Usuário (Login)*")
        s1 = st.text_input("Senha*", type="password")
        s2 = st.text_input("Confirmar senha*", type="password")
        c1, c2 = st.columns(2)
        btn_voltar = c1.form_submit_button("⬅️ CANCELAR", type="secondary")
        btn_confirmar = c2.form_submit_button("CONFIRMAR", type="primary")

    if btn_voltar: nav('login')
    if btn_confirmar:
        if not u: msg_error("Informe um nome de utilizador.")
        elif s1 != s2: msg_error("As senhas não conferem.")
        elif len(s1) < 4: msg_error("A senha deve ter pelo menos 4 caracteres.")
        else:
            nome_final = nome_completo.strip() if nome_completo.strip() else u
            try:
                with get_db() as conn:
                    conn.execute("INSERT INTO tb_usuario (nome, login, senha_hash, id_perfil) VALUES (?,?,?,2)", (nome_final, u, gerar_hash(s1)))
                    conn.commit()
                set_flash("success", f"Utilizador **{u}** criado com sucesso! Faça login para continuar.")
                nav('login')
            except sqlite3.IntegrityError:
                msg_error(f"O utilizador '{u}' já existe. Escolha outro nome.")

def tela_menu():
    cabecalho()
    st.markdown(f"<h3 style='text-align: center; color: #006400;'>Olá, {st.session_state.nome_usuario}!</h3>", unsafe_allow_html=True)
    st.markdown(f"<p style='text-align: center; color: #6c757d; margin-top:-10px;'>Perfil: {st.session_state.user_role}</p>", unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    
    if st.button("➕ Nova empresa / Alteração", type="primary", use_container_width=True): nav('cad_empresa')
    if st.button("🔍 Buscar empresa cadastrada", type="primary", use_container_width=True): nav('buscar_empresa')
    if st.button("📄 Histórico geral de documentos", type="secondary", use_container_width=True): nav('historico')
        
    if st.session_state.user_role == 'Admin':
        st.markdown("<hr style='margin: 15px 0px;'>", unsafe_allow_html=True)
        if st.button("👥 Gestão de Utilizadores", type="secondary", use_container_width=True): nav('gestao_usuarios')
            
    st.divider()
    if st.button("⬅️ Sair da conta", type="secondary", use_container_width=True):
        for key in ['user_id', 'user_role', 'empresa_id', 'socio_id']: st.session_state[key] = None
        nav('login')

def tela_gestao_usuarios():
    if st.session_state.user_role != 'Admin':
        msg_error("Acesso Negado.")
        nav('menu')
        return

    cabecalho()
    st.subheader("👥 Gestão de Utilizadores")
    with get_db() as conn:
        usuarios = conn.execute("SELECT u.id, u.nome, u.login, p.nome_perfil, u.id_perfil FROM tb_usuario u JOIN tb_perfil_acesso p ON u.id_perfil = p.id").fetchall()

    for u in usuarios:
        user_id, nome, login, perfil_nome, perfil_id = u
        with st.container():
            col1, col2, col3 = st.columns([0.5, 0.3, 0.2], vertical_alignment="center")
            col1.markdown(f"**{nome}**<br><span style='color:gray; font-size:0.8em;'>Login: {login}</span>", unsafe_allow_html=True)
            if user_id == st.session_state.user_id:
                col2.markdown(f"*(Você)* - {perfil_nome}")
            else:
                acao = "Rebaixar a Operador" if perfil_nome == 'Admin' else "Promover a Admin"
                novo_perfil_id = 2 if perfil_nome == 'Admin' else 1
                if col2.button(acao, key=f"prom_{user_id}", type="secondary"):
                    with get_db() as conn_upd:
                        conn_upd.execute("UPDATE tb_usuario SET id_perfil=? WHERE id=?", (novo_perfil_id, user_id))
                        conn_upd.commit()
                    st.rerun()
                if col3.button("🗑️", key=f"del_{user_id}", type="secondary"):
                    with get_db() as conn_del:
                        conn_del.execute("DELETE FROM tb_usuario WHERE id=?", (user_id,))
                        conn_del.commit()
                    st.rerun()
            st.divider()
    if st.button("⬅️ VOLTAR AO MENU", type="secondary"): nav('menu')

def tela_cad_empresa():
    cabecalho()
    st.subheader("DADOS DA EMPRESA")
    
    st.markdown("<h4 style='color:#6c757d; font-size:16px;'>🔍 Preenchimento Automático</h4>", unsafe_allow_html=True)
    c_cnpj, c_btn = st.columns([3, 1], vertical_alignment="bottom")
    cnpj_busca = c_cnpj.text_input("Digite o CNPJ para buscar na Receita:")
    
    if c_btn.button("BUSCAR", use_container_width=True):
        if cnpj_busca:
            cnpj_limpo = ''.join(filter(str.isdigit, cnpj_busca))
            if len(cnpj_limpo) == 14:
                with st.spinner("A consultar dados oficiais..."):
                    try:
                        resp = requests.get(f"https://brasilapi.com.br/api/cnpj/v1/{cnpj_limpo}", timeout=10)
                        if resp.status_code == 200:
                            d = resp.json()
                            st.session_state['api_cnpj'] = cnpj_busca
                            st.session_state['api_nome'] = d.get('razao_social') or ''
                            
                            tipo_log = d.get('descricao_tipo_de_logradouro') or ''
                            log = d.get('logradouro') or ''
                            st.session_state['api_logradouro'] = f"{tipo_log} {log}".strip()
                            
                            st.session_state['api_numero'] = d.get('numero') or ''
                            st.session_state['api_complemento'] = d.get('complemento') or ''
                            st.session_state['api_bairro'] = d.get('bairro') or ''
                            st.session_state['api_cidade'] = d.get('municipio') or ''
                            st.session_state['api_uf'] = d.get('uf') or ''
                            
                            cep_api = str(d.get('cep') or '')
                            if len(cep_api) == 8: cep_api = f"{cep_api[:5]}-{cep_api[5:]}"
                            st.session_state['api_cep'] = cep_api
                            
                            qsa = d.get('qsa') or []
                            socios_encontrados = []
                            for s in qsa:
                                nome_socio = s.get("nome_socio")
                                if nome_socio:
                                    socios_encontrados.append({
                                        "nome": nome_socio,
                                        "cpf": s.get("cnpj_cpf_do_socio") or "" 
                                    })
                            st.session_state['api_qsa'] = socios_encontrados
                            
                            if socios_encontrados:
                                set_flash("success", f"✅ Dados e {len(socios_encontrados)} sócio(s) encontrados!")
                            else:
                                set_flash("success", "✅ Dados encontrados! (Nota: Esta empresa não possui quadro de sócios na Receita).")
                            st.rerun()
                        else:
                            msg_error("CNPJ não encontrado na base de dados.")
                    except Exception as e:
                        msg_error(f"Ocorreu um erro ao ligar à API: {e}")
            else:
                msg_warning("O CNPJ deve ter 14 dígitos.")
        else:
            msg_warning("Digite um CNPJ válido antes de buscar.")

    st.divider()
    
    with st.form("f_emp"):
        cnpj = st.text_input("CNPJ*", value=st.session_state.get('api_cnpj', ''))
        nome = st.text_input("NOME EMPRESARIAL*", value=st.session_state.get('api_nome', ''))
        nire = st.text_input("NIRE (Apenas Manual)")
        st.markdown("**Morada da Sede***")
        
        c_cep_emp, c_btn_cep_emp = st.columns([3, 1], vertical_alignment="bottom")
        cep = c_cep_emp.text_input("CEP*", value=st.session_state.get('api_cep', ''))
        btn_buscar_cep_emp = c_btn_cep_emp.form_submit_button("BUSCAR", type="secondary", use_container_width=True)
        
        if btn_buscar_cep_emp:
             cep_limpo = ''.join(filter(str.isdigit, cep))
             if len(cep_limpo) == 8:
                 try:
                     resp = requests.get(f"https://viacep.com.br/ws/{cep_limpo}/json/", timeout=5)
                     if resp.status_code == 200:
                         dados_cep = resp.json()
                         if "erro" not in dados_cep:
                             st.session_state['api_cep'] = dados_cep.get('cep', '')
                             st.session_state['api_logradouro'] = dados_cep.get('logradouro', '').upper()
                             st.session_state['api_bairro'] = dados_cep.get('bairro', '').upper()
                             st.session_state['api_cidade'] = dados_cep.get('localidade', '').upper()
                             st.session_state['api_uf'] = dados_cep.get('uf', '').upper()
                             st.rerun()
                         else:
                             msg_error("CEP não encontrado.")
                 except Exception:
                     msg_error("Erro ao consultar a API de CEP.")
             else:
                 msg_warning("O CEP deve conter 8 dígitos.")
        
        c7, c8 = st.columns([3, 1])
        logradouro = c7.text_input("LOGRADOURO*", value=st.session_state.get('api_logradouro', ''))
        numero = c8.text_input("NÚMERO*", value=st.session_state.get('api_numero', ''))
        c9, c10 = st.columns(2)
        complemento = c9.text_input("COMPLEMENTO", value=st.session_state.get('api_complemento', ''))
        bairro = c10.text_input("BAIRRO*", value=st.session_state.get('api_bairro', ''))
        c11, c12 = st.columns([3, 1])
        cidade = c11.text_input("CIDADE*", value=st.session_state.get('api_cidade', ''))
        uf = c12.text_input("UF*", value=st.session_state.get('api_uf', ''))
        
        c1, c2 = st.columns(2)
        btn_voltar = c1.form_submit_button("⬅️ VOLTAR", type="secondary")
        btn_seguir = c2.form_submit_button("SEGUIR E SALVAR ➡️", type="primary")

    if btn_voltar:
        for k in list(st.session_state.keys()):
            if k.startswith('api_'): del st.session_state[k]
        nav('menu')
        
    if btn_seguir:
        erros = [c for c, v in zip(["CNPJ", "Nome", "Logradouro", "Número", "Bairro", "Cidade", "UF", "CEP"], [cnpj, nome, logradouro, numero, bairro, cidade, uf, cep]) if not v]
        if erros: 
            msg_error(f"Preencha os campos obrigatórios: {', '.join(erros)}.")
        else:
            try:
                with get_db() as conn:
                    cur = conn.cursor()
                    cur.execute("INSERT INTO tb_empresa (cnpj, nome_empresarial, nire, logradouro, numero, complemento, bairro, cidade, uf, cep) VALUES (?,?,?,?,?,?,?,?,?,?)",
                        (cnpj, nome.upper(), nire, logradouro.upper(), numero, complemento.upper(), bairro.upper(), cidade.upper(), uf.upper(), cep))
                    st.session_state.empresa_id = cur.lastrowid
                    
                    qsa_api = st.session_state.get('api_qsa', [])
                    for s in qsa_api:
                        if s['nome']:
                            cur.execute("""INSERT INTO tb_socio 
                                (id_empresa, nome, cpf, nacionalidade, estado_civil, logradouro, numero, bairro, cidade, uf, cep) 
                                VALUES (?,?,?, 'BRASILEIRO(A)', '', '', '', '', '', '', '')""",
                                (st.session_state.empresa_id, s['nome'].upper(), s['cpf']))
                    
                    conn.commit()
                
                for k in list(st.session_state.keys()):
                    if k.startswith('api_'): del st.session_state[k]
                    
                nav('lista_socios')
            except sqlite3.IntegrityError:
                msg_error(f"Já existe uma empresa com o CNPJ '{cnpj}'. Use 'Buscar empresa' para aceder à mesma.")

def tela_buscar_empresa():
    cabecalho()
    st.subheader("SELECIONE UMA EMPRESA")
    with get_db() as conn:
        empresas = conn.execute("SELECT id, cnpj, nome_empresarial FROM tb_empresa").fetchall()
    
    if not empresas:
        st.info("Nenhuma empresa cadastrada.")
        if st.button("⬅️ VOLTAR AO MENU", type="secondary"): nav('menu')
    else:
        opcoes = {f"{e[1]} - {e[2]}": e[0] for e in empresas}
        
        escolha = st.selectbox("Empresas:", list(opcoes.keys()), index=None, placeholder="Selecione uma empresa na lista...")
        
        if st.button("ACESSAR EMPRESA", type="primary", use_container_width=True):
            if escolha is None:
                msg_warning("⚠️ Selecione uma empresa antes de avançar.")
            else:
                st.session_state.empresa_id = opcoes[escolha]
                nav('lista_socios')
                
        if st.button("⬅️ VOLTAR AO MENU", type="secondary", use_container_width=True): nav('menu')

def tela_editar_empresa():
    cabecalho()
    st.subheader("EDITAR EMPRESA")
    is_admin = st.session_state.user_role == 'Admin'
    with get_db() as conn:
        e = conn.execute("SELECT cnpj, nome_empresarial, nire, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,)).fetchone()

    if not is_admin: st.info("💡 Apenas Administradores podem forçar a edição manual dos dados da empresa.")

    with st.form("f_edit_emp"):
        cnpj = st.text_input("CNPJ*", value=e[0] or '', disabled=not is_admin)
        nome = st.text_input("NOME EMPRESARIAL*", value=e[1] or '', disabled=not is_admin)
        nire = st.text_input("NIRE", value=e[2] or '', disabled=not is_admin)
        st.markdown("**Endereço**")
        c7, c8 = st.columns([3, 1])
        logradouro = c7.text_input("LOGRADOURO*", value=e[3] or '', disabled=not is_admin)
        numero = c8.text_input("NÚMERO*", value=e[4] or '', disabled=not is_admin)
        c9, c10 = st.columns(2)
        complemento = c9.text_input("COMPLEMENTO", value=e[5] or '', disabled=not is_admin)
        bairro = c10.text_input("BAIRRO*", value=e[6] or '', disabled=not is_admin)
        c11, c12, c13 = st.columns([3, 1, 2])
        cidade = c11.text_input("CIDADE*", value=e[7] or '', disabled=not is_admin)
        uf = c12.text_input("UF*", value=e[8] or '', disabled=not is_admin)
        cep = c13.text_input("CEP*", value=e[9] or '', disabled=not is_admin)
        c1, c2 = st.columns(2)
        btn_v = c1.form_submit_button("⬅️ CANCELAR", type="secondary")
        btn_ok = c2.form_submit_button("💾 SALVAR", type="primary") if is_admin else False 

    if btn_v: nav('lista_socios')
    if btn_ok and is_admin:
        with get_db() as conn:
            conn.execute("UPDATE tb_empresa SET cnpj=?, nome_empresarial=?, nire=?, logradouro=?, numero=?, complemento=?, bairro=?, cidade=?, uf=?, cep=? WHERE id=?",
                (cnpj, nome.upper(), nire, logradouro.upper(), numero, complemento.upper(), bairro.upper(), cidade.upper(), uf.upper(), cep, st.session_state.empresa_id))
            conn.commit()
        set_flash("success", "Empresa atualizada com sucesso!")
        nav('lista_socios')

def tela_socios():
    cabecalho()
    is_admin = st.session_state.user_role == 'Admin'

    with get_db() as conn:
        emp = conn.execute("SELECT nome_empresarial FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,)).fetchone()
        socios = conn.execute("SELECT id, nome, cpf FROM tb_socio WHERE id_empresa=?", (st.session_state.empresa_id,)).fetchall()

    if st.session_state.get('_confirm_excluir_empresa') and is_admin:
        st.warning(f"⚠️ Tem certeza que deseja excluir a empresa **{emp[0]}** e todos os seus sócios? Esta ação não pode ser desfeita.")
        ca, cb = st.columns(2)
        if ca.button("❌ Cancelar", type="secondary", use_container_width=True):
            st.session_state.pop('_confirm_excluir_empresa', None)
            st.rerun()
        if cb.button("🗑️ Sim, excluir empresa", type="primary", use_container_width=True):
            with get_db() as conn:
                conn.execute("DELETE FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,))
                conn.commit()
            st.session_state.empresa_id = None
            st.session_state.pop('_confirm_excluir_empresa', None)
            nav('menu')
        return

    if st.session_state.get('_confirm_excluir_socio') and is_admin:
        sid, snome = st.session_state['_confirm_excluir_socio']
        st.warning(f"⚠️ Tem certeza que deseja excluir o sócio **{snome}**?")
        ca, cb = st.columns(2)
        if ca.button("❌ Cancelar", type="secondary", use_container_width=True):
            st.session_state.pop('_confirm_excluir_socio', None)
            st.rerun()
        if cb.button("🗑️ Sim, excluir sócio", type="primary", use_container_width=True):
            with get_db() as conn:
                conn.execute("DELETE FROM tb_socio WHERE id=?", (sid,))
                conn.commit()
            st.session_state.pop('_confirm_excluir_socio', None)
            st.rerun()
        return

    if is_admin:
        col_n, col_ed, col_ex = st.columns([0.6, 0.2, 0.2], vertical_alignment="center")
        col_n.markdown(f"### {emp[0]}")
        if col_ed.button("✏️ Editar", key="ed_emp"): nav('editar_empresa')
        if col_ex.button("🗑️ Excluir", key="ex_emp"):
            st.session_state['_confirm_excluir_empresa'] = True
            st.rerun()
    else:
        col_n, col_ed = st.columns([0.8, 0.2], vertical_alignment="center")
        col_n.markdown(f"### {emp[0]}")
        if col_ed.button("👁️‍🗨️ Ver Dados", key="ed_emp"): nav('editar_empresa')

    st.markdown("<div class='title-separator'></div>", unsafe_allow_html=True)
    st.markdown("<h4 style='color:#006400;text-align:center;'>Quadro Societário</h4>", unsafe_allow_html=True)

    for s in socios:
        if is_admin:
            c1, c2, c3 = st.columns([0.6, 0.2, 0.2], vertical_alignment="center")
            c1.markdown(f"☑ **{s[1]}** *(CPF: {s[2]})*")
            if c2.button("✏️ Editar", key=f"ed_s_{s[0]}", type="secondary"):
                st.session_state.socio_id = s[0]
                nav('editar_socio')
            if c3.button("🗑️ Excluir", key=f"ex_s_{s[0]}", type="secondary"):
                st.session_state['_confirm_excluir_socio'] = (s[0], s[1])
                st.rerun()
        else:
            c1, c2 = st.columns([0.8, 0.2], vertical_alignment="center")
            c1.markdown(f"☑ **{s[1]}** *(CPF: {s[2]})*")
            if c2.button("✏️ Editar", key=f"ed_s_{s[0]}", type="secondary"):
                st.session_state.socio_id = s[0]
                nav('editar_socio')

    st.markdown("<br>", unsafe_allow_html=True)
    b1, b2, b3 = st.columns([1, 2, 1])
    with b2:
        if st.button("➕ Novo Sócio", type="primary", use_container_width=True): nav('cad_socio')

    st.markdown("<br>", unsafe_allow_html=True)
    if st.button("📄 Histórico de Documentos desta Empresa", type="secondary", use_container_width=True): nav('historico_empresa')
    st.markdown("<hr>", unsafe_allow_html=True)
    n1, n2 = st.columns(2)
    if n1.button("⬅️ VOLTAR AO MENU", type="secondary"):
        st.session_state.empresa_id = None
        nav('menu')
    if n2.button("AVANÇAR PARA EVENTOS ➡️", type="primary"):
        nav('orquestrador')

def tela_cad_socio():
    btn_v, btn_ok, dados = _form_socio("CADASTRAR NOVO SÓCIO")
    if btn_v: nav('lista_socios')
    if btn_ok:
        with get_db() as conn:
            conn.execute("""INSERT INTO tb_socio (id_empresa, nome, nacionalidade, cpf, rg, orgao_emissor, estado_civil, data_nasc, profissao, logradouro, numero, complemento, bairro, cidade, uf, cep) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (st.session_state.empresa_id, dados['nome'], dados['nacionalidade'], dados['cpf'], dados['rg'], dados['orgao_emissor'], dados['estado_civil'], dados['data_nasc'], dados['profissao'], dados['logradouro'], dados['numero'], dados['complemento'], dados['bairro'], dados['cidade'], dados['uf'], dados['cep']))
            conn.commit()
        set_flash("success", f"Sócio **{dados['nome']}** cadastrado com sucesso!")
        nav('lista_socios')

def tela_editar_socio():
    with get_db() as conn:
        s = conn.execute("SELECT nome, nacionalidade, cpf, rg, orgao_emissor, estado_civil, data_nasc, profissao, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_socio WHERE id=?", (st.session_state.socio_id,)).fetchone()
    defaults = { 'nome': s[0], 'nacionalidade': s[1], 'cpf': s[2], 'rg': s[3], 'orgao_emissor': s[4], 'estado_civil': s[5], 'data_nasc': s[6], 'profissao': s[7], 'logradouro': s[8], 'numero': s[9], 'complemento': s[10], 'bairro': s[11], 'cidade': s[12], 'uf': s[13], 'cep': s[14] }
    btn_v, btn_ok, dados = _form_socio("EDITAR SÓCIO", defaults)
    if btn_v: nav('lista_socios')
    if btn_ok:
        with get_db() as conn:
            conn.execute("""UPDATE tb_socio SET nome=?, nacionalidade=?, cpf=?, rg=?, orgao_emissor=?, estado_civil=?, data_nasc=?, profissao=?, logradouro=?, numero=?, complemento=?, bairro=?, cidade=?, uf=?, cep=? WHERE id=?""",
                (dados['nome'], dados['nacionalidade'], dados['cpf'], dados['rg'], dados['orgao_emissor'], dados['estado_civil'], dados['data_nasc'], dados['profissao'], dados['logradouro'], dados['numero'], dados['complemento'], dados['bairro'], dados['cidade'], dados['uf'], dados['cep'], st.session_state.socio_id))
            conn.commit()
        set_flash("success", "Sócio atualizado com sucesso!")
        nav('lista_socios')

def tela_historico():
    cabecalho()
    st.subheader("📄 Histórico Geral de Documentos")
    with get_db() as conn:
        docs = conn.execute("""
            SELECT d.data_geracao, d.tipo_documento, e.nome_empresarial, u.nome, d.status
            FROM tb_documento d JOIN tb_empresa e ON d.id_empresa = e.id JOIN tb_usuario u ON d.id_usuario = u.id ORDER BY d.data_geracao DESC
        """).fetchall()
    if not docs: st.info("Nenhum documento gerado ainda.")
    else:
        for doc in docs:
            st.markdown(f"**{doc[2]}** — {doc[1]} | Gerado por: {doc[3]} em {doc[0]} | Status: {doc[4]}")
            st.divider()
    if st.button("⬅️ VOLTAR AO MENU", type="secondary"): nav('menu')

def tela_historico_empresa():
    cabecalho()
    st.subheader("📄 Histórico da Empresa")
    with get_db() as conn:
        emp = conn.execute("SELECT nome_empresarial FROM tb_empresa WHERE id=?", (st.session_state.empresa_id,)).fetchone()
        docs = conn.execute("""
            SELECT d.data_geracao, d.tipo_documento, u.nome, d.status
            FROM tb_documento d JOIN tb_usuario u ON d.id_usuario = u.id WHERE d.id_empresa=? ORDER BY d.data_geracao DESC
        """, (st.session_state.empresa_id,)).fetchall()
    st.markdown(f"**Empresa:** {emp[0]}")
    if not docs: st.info("Nenhum documento gerado para esta empresa.")
    else:
        for doc in docs:
            st.markdown(f"{doc[1]} | Gerado por: {doc[2]} em {doc[0]} | Status: {doc[3]}")
            st.divider()
    if st.button("⬅️ VOLTAR", type="secondary"): nav('lista_socios')

# ======================================================================
# ARQUITETURA PROTOTYPE + CÓDIGOS DBE/JUCEC
# ======================================================================
class ClausulaPrototipo:
    def clone(self): return copy.deepcopy(self)
    def aplicar_dados(self, dados): pass

class ClausulaTransformacao(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Transformar o tipo jurídico para Sociedade Empresária Limitada, adotando o nome empresarial {novo_nome}."
    def aplicar_dados(self, dados): return self.texto_base.format(novo_nome=dados.get("novo_nome", ""))

class ClausulaNome(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passará a denominar-se: {novo_nome}."
    def aplicar_dados(self, dados): return self.texto_base.format(novo_nome=dados.get("novo_nome", ""))

class ClausulaNomeFantasia(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passará a ter o TÍTULO DO ESTABELECIMENTO (NOME FANTASIA): {nome_fantasia}."
    def aplicar_dados(self, dados): return self.texto_base.format(nome_fantasia=dados.get("nome_fantasia", ""))

class ClausulaEndereco(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Alterar o endereço da sociedade, que passa a localizar-se na(o) {logradouro}, número {numero}{complemento}, bairro {bairro}, município {cidade} - {uf}, CEP: {cep}."
    def aplicar_dados(self, dados):
        comp = f", complemento {dados.get('complemento')}" if dados.get("complemento") else ""
        return self.texto_base.format(
            logradouro=dados.get("logradouro", ""), numero=dados.get("numero", ""),
            complemento=comp, bairro=dados.get("bairro", ""),
            cidade=dados.get("cidade", ""), uf=dados.get("uf", ""), cep=dados.get("cep", "")
        )

class ClausulaTipoUnidade(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Fica alterado o tipo de unidade do estabelecimento, passando a ser classificado como: {tipo_unidade}."
    def aplicar_dados(self, dados): return self.texto_base.format(tipo_unidade=dados.get("tipo_unidade", ""))

class ClausulaAtividadesEcono(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passa a ter por objeto, o exercício das seguintes atividades econômicas: {objeto_descricao}. \nA sociedade passa a ter as seguintes atividades econômicas (CNAES) vinculadas ao seu objeto social: {cnaes_lista}."
    def aplicar_dados(self, dados):
        return self.texto_base.format(
            objeto_descricao=dados.get("objeto_descricao", ""),
            cnaes_lista=dados.get("cnaes_lista", "")
        )

class ClausulaFormaAtuacao(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade passará a exercer suas atividades sob a seguinte forma de atuação: {forma_atuacao}."
    def aplicar_dados(self, dados): return self.texto_base.format(forma_atuacao=dados.get("forma_atuacao", ""))

class ClausulaReativacao(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "A sociedade, que se encontrava com seu registro inativo, reativa neste ato as suas atividades, em conformidade com o Artigo 60 da Lei 8.934/94."
    def aplicar_dados(self, dados): return self.texto_base

class ClausulaLicenciamento(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "Fica registrado o licenciamento do estabelecimento anteriormente registrado (Legado), para todos os efeitos legais e administrativos inerentes."
    def aplicar_dados(self, dados): return self.texto_base

class ClausulaCapital(ClausulaPrototipo):
    def __init__(self):
        self.texto_base = "O capital social é de R$ {novo_capital}, dividido em quotas de valor nominal de R$ 1,00 cada, totalmente subscrito e integralizado neste ato."
    def aplicar_dados(self, dados): return self.texto_base.format(novo_capital=dados.get("novo_capital", ""))

CATALOGO_CLAUSULAS = {
    "225": ClausulaTransformacao(), "220": ClausulaNome(), "221": ClausulaNomeFantasia(),
    "211": ClausulaEndereco(), "209": ClausulaEndereco(), "210": ClausulaEndereco(),
    "248": ClausulaTipoUnidade(), "244": ClausulaAtividadesEcono(), "249": ClausulaFormaAtuacao(),
    "052": ClausulaReativacao(), "999": ClausulaLicenciamento(), "capital": ClausulaCapital()
}

# ======================================================================
# APOIO DE IA PARA CNAE - PROTÓTIPO LOCAL
# ======================================================================
CATALOGO_CNAE_IA = [
    {
        "codigo": "4729-6/99",
        "descricao": "Comércio varejista de produtos alimentícios em geral ou especializado em produtos alimentícios não especificados anteriormente",
        "palavras": ["salgado", "salgados", "alimento", "alimenticio", "alimentícios", "comida", "produtos alimenticios", "produtos alimentícios"],
    },
    {
        "codigo": "5611-2/03",
        "descricao": "Lanchonetes, casas de chá, de sucos e similares",
        "palavras": ["lanchonete", "lanche", "coxinha", "pastel", "salgado pronto", "suco", "café", "cafe"],
    },
    {
        "codigo": "5620-1/04",
        "descricao": "Fornecimento de alimentos preparados preponderantemente para consumo domiciliar",
        "palavras": ["delivery", "entrega", "marmita", "comida pronta", "consumo domiciliar"],
    },
    {
        "codigo": "1091-1/02",
        "descricao": "Fabricação de produtos de padaria e confeitaria com predominância de produção própria",
        "palavras": ["fabricacao", "fabricação", "produzir", "produção", "padaria", "confeitaria", "bolo", "pão", "pao"],
    },
    {
        "codigo": "4693-1/00",
        "descricao": "Comércio atacadista de mercadorias em geral, sem predominância de alimentos ou de insumos agropecuários",
        "palavras": ["atacado", "atacadista", "distribuidora", "distribuição", "distribuicao"],
    },
]

def _normalizar_busca(txt):
    txt = (txt or "").lower()
    trocas = {"á":"a", "à":"a", "â":"a", "ã":"a", "é":"e", "ê":"e", "í":"i", "ó":"o", "ô":"o", "õ":"o", "ú":"u", "ç":"c"}
    for a, b in trocas.items():
        txt = txt.replace(a, b)
    return txt

def sugerir_cnaes_por_texto(texto):
    """Sugere CNAEs a partir de texto livre.

    Para o TCC, esta função simula a camada de IA. Em produção, pode ser trocada
    por integração com API de IA mantendo a mesma entrada e saída.
    """
    texto_busca = _normalizar_busca(texto)
    resultados = []
    for item in CATALOGO_CNAE_IA:
        score = 0
        for palavra in item["palavras"]:
            if _normalizar_busca(palavra) in texto_busca:
                score += 1
        if score:
            resultados.append((score, item))
    resultados.sort(key=lambda x: x[0], reverse=True)
    return [item for _, item in resultados[:5]]

def montar_objeto_social_sugerido(descricao, sugestoes):
    if not descricao:
        return ""
    if not sugestoes:
        return (descricao or "").upper()
    cnaes = "; ".join([f"{i['codigo']} - {i['descricao']}" for i in sugestoes])
    return (
        f"A sociedade terá por objeto social o exercício de atividades relacionadas a {descricao}, "
        f"observados os seguintes CNAEs compatíveis: {cnaes}."
    ).upper()

def _ordinal(n):
    nomes = {1: "Primeira", 2: "Segunda", 3: "Terceira", 4: "Quarta", 5: "Quinta", 6: "Sexta", 7: "Sétima", 8: "Oitava", 9: "Nona", 10: "Décima"}
    return nomes.get(n, str(n) + "ª")

def processar_clausulas_prototype(eventos_selecionados):
    clausulas_prontas = []
    for i, evento in enumerate(eventos_selecionados, start=1):
        prototipo = CATALOGO_CLAUSULAS[evento["tipo"]].clone()
        texto_preenchido = prototipo.aplicar_dados(evento["dados"])
        clausulas_prontas.append(f"Cláusula {_ordinal(i)} – {texto_preenchido}")
    return clausulas_prontas

# ======================================================================
# GERAÇÃO DO DOCUMENTO DOCX
# ======================================================================
def gerar_docx_modular(empresa, socios, clausulas_prontas):
    doc = Document()
    for section in doc.sections:
        section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = Cm(2.5), Cm(2.5), Cm(3), Cm(3)

    def add_paragrafo(texto, bold=False, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        run = p.add_run(texto)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.name = 'Arial'
        return p

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_after = Pt(12)
    run_titulo = titulo.add_run(f"ALTERAÇÃO CONTRATUAL {empresa.get('nome_empresarial', '')}")
    run_titulo.bold, run_titulo.font.size, run_titulo.font.name = True, Pt(14), 'Arial'

    for s in socios:
        partes = [s.get('nome', '')]
        if s.get('nacionalidade'): partes.append(f"nacionalidade {s['nacionalidade']}")
        if s.get('estado_civil'): partes.append(s['estado_civil'])
        if s.get('data_nasc'): partes.append(f"nascido em {s['data_nasc']}")
        if s.get('profissao'): partes.append(f"profissão: {s['profissao']}")
        if s.get('cpf'): partes.append(f"nº do CPF: {s['cpf']}")
        if s.get('rg'): partes.append(f"identidade: {s['rg']}")
        if s.get('orgao_emissor'): partes.append(f"órgão expedidor: {s['orgao_emissor']}")
        
        end_partes = []
        if s.get('logradouro'): end_partes.append(s['logradouro'])
        if s.get('numero'): end_partes.append(f"número {s['numero']}")
        if s.get('complemento'): end_partes.append(s['complemento'])
        if s.get('bairro'): end_partes.append(f"bairro {s['bairro']}")
        if s.get('cidade') and s.get('uf'): end_partes.append(f"município {s['cidade']} - {s['uf']}")
        if s.get('cep'): end_partes.append(f"CEP: {s['cep']}")
        
        if end_partes: partes.append(f"RESIDENTE E DOMICILIADO no(a): {', '.join(end_partes)}")
        add_paragrafo(', '.join(partes) + ',')

    emp_partes = [f"Sócio(s) da sociedade {empresa.get('nome_empresarial', '')}"]
    end_emp = []
    if empresa.get('logradouro'): end_emp.append(empresa['logradouro'])
    if empresa.get('numero'): end_emp.append(f"número {empresa['numero']}")
    if empresa.get('complemento'): end_emp.append(empresa['complemento'])
    if empresa.get('bairro'): end_emp.append(f"bairro {empresa['bairro']}")
    if empresa.get('cidade') and empresa.get('uf'): end_emp.append(f"município {empresa['cidade']} - {empresa['uf']}")
    if empresa.get('cep'): end_emp.append(f"CEP: {empresa['cep']}")
    if end_emp: emp_partes.append(f"sediada na {', '.join(end_emp)}")
    if empresa.get('nire'): emp_partes.append(f"com seu contrato social arquivado nessa Junta Comercial")
    if empresa.get('cnpj'): emp_partes.append(f"devidamente inscrita no CNPJ sob o nº {empresa['cnpj']}")
    emp_partes.append("resolvem:")

    add_paragrafo(', '.join(emp_partes))

    for idx, clausula in enumerate(clausulas_prontas):
        add_paragrafo(clausula, bold=False) 
        
    add_paragrafo(f"Cláusula {_ordinal(len(clausulas_prontas) + 1)} – Permanecem em vigor todas as demais cláusulas e condições não alteradas pelo presente instrumento.")
    add_paragrafo("E, por estarem assim justos e acordados, assinam o presente instrumento para todos os efeitos legais.")
    add_paragrafo("")

    cidade_empresa = empresa.get('cidade', '_____________')
    add_paragrafo(f"{cidade_empresa}, _____ de ________________ de _______.", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragrafo("")

    for s in socios:
        add_paragrafo("_" * 50, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        add_paragrafo(f"{s.get('nome', '')}: Sócio/Administrador", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def converter_docx_para_pdf(docx_bytes):
    """Converte DOCX para PDF usando LibreOffice/soffice, quando disponível.

    Em ambiente local, instale o LibreOffice para habilitar essa função.
    Retorna (pdf_bytes, erro).
    """
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        docx_path = tmp / "documento.docx"
        pdf_path = tmp / "documento.pdf"

        with open(docx_path, "wb") as f:
            f.write(docx_bytes.getvalue())

        for comando in ("libreoffice", "soffice"):
            try:
                subprocess.run(
                    [comando, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(docx_path)],
                    check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30
                )
                if pdf_path.exists():
                    pdf_bytes = io.BytesIO(pdf_path.read_bytes())
                    pdf_bytes.seek(0)
                    return pdf_bytes, None
            except FileNotFoundError:
                continue
            except Exception as e:
                return None, f"Erro ao converter para PDF: {e}"

    return None, "PDF não gerado porque LibreOffice/soffice não está instalado no ambiente."


# ======================================================================
# ORQUESTRADOR STREAMLIT (FLUXO WIZARD - 2 TELAS)
# ======================================================================
def tela_orquestrador():
    cabecalho()
    st.subheader("🛒 Orquestrador de Eventos - Passo 1: Seleção")
    st.markdown("Selecione todos os eventos desejados e clique em Avançar.")
    
    with st.form("form_eventos_selecao"):
        st.markdown("<h4 style='color:#6c757d;'>📝 Identificação da Empresa</h4>", unsafe_allow_html=True)
        alt_220 = st.checkbox("220 - Alteração do Nome Empresarial")
        alt_221 = st.checkbox("221 - Alteração do Título do Estabelecimento (Nome Fantasia)")
        
        st.markdown("<h4 style='color:#6c757d;'>🏠 Localização & Unidade</h4>", unsafe_allow_html=True)
        alt_endereco = st.radio("Alteração de Endereço:", ["Nenhuma", "211 - Mesmo Município", "209 - Entre Municípios", "210 - Entre Estados"], horizontal=True)
        alt_248 = st.checkbox("248 - Alteração do Tipo de Unidade")
        
        st.markdown("<h4 style='color:#6c757d;'>💼 Atividades e Atuação</h4>", unsafe_allow_html=True)
        alt_244 = st.checkbox("244 - Alteração de Atividades Econômicas / Objeto Social")
        alt_249 = st.checkbox("249 - Alteração da Forma de Atuação")
        
        st.markdown("<h4 style='color:#6c757d;'>💰 Estrutura Jurídica e Status</h4>", unsafe_allow_html=True)
        alt_225 = st.checkbox("225 - Alteração da Natureza Jurídica (Transformação)")
        alt_capital = st.checkbox("Alteração de Capital Social")
        alt_052 = st.checkbox("052 - Reativação (Artigo 60 Lei 8.934/94)")
        alt_999 = st.checkbox("999 - Licenciamento de Estabelecimento (Legado)")
        
        st.divider()
        c1, c2 = st.columns(2)
        btn_voltar = c1.form_submit_button("⬅️ VOLTAR AOS SÓCIOS", type="secondary")
        btn_avancar = c2.form_submit_button("AVANÇAR PARA PREENCHIMENTO ➡️", type="primary")
        
    if btn_voltar:
        st.session_state.mostrar_confirmacao_repetidos = False
        nav('lista_socios')
        
    if btn_avancar:
        selecionados = []
        if alt_220: selecionados.append("220")
        if alt_221: selecionados.append("221")
        if alt_endereco != "Nenhuma": selecionados.append(alt_endereco[:3])
        if alt_248: selecionados.append("248")
        if alt_244: selecionados.append("244")
        if alt_249: selecionados.append("249")
        if alt_225: selecionados.append("225")
        if alt_capital: selecionados.append("capital")
        if alt_052: selecionados.append("052")
        if alt_999: selecionados.append("999")
        
        if not selecionados:
            msg_error("❌ **Erro:** Selecione pelo menos um evento para continuar.")
        else:
            st.session_state.eventos_selecionados_lista = selecionados
            st.session_state.mostrar_confirmacao_repetidos = False
            nav('orquestrador_dados')

def tela_orquestrador_dados():
    cabecalho()
    st.subheader("📝 Orquestrador - Passo 2: Preenchimento")
    
    selecionados = st.session_state.get('eventos_selecionados_lista', [])
    if not selecionados:
        msg_error("Nenhum evento selecionado.")
        if st.button("Voltar", type="secondary"): nav('orquestrador')
        return

    with get_db() as conn:
        emp_row = conn.execute(
            "SELECT cnpj, nome_empresarial, nome_fantasia, nire, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_empresa WHERE id=?",
            (st.session_state.empresa_id,)
        ).fetchone()
        
    if "244" in selecionados:
        st.markdown("### 🤖 Apoio de IA para CNAE")
        atividade_ia = st.text_input(
            "Descreva a atividade de forma simples",
            placeholder="Ex.: venda de salgados, entrega de marmitas, lanchonete...",
            key="atividade_ia_cnae"
        )
        if st.button("SUGERIR CNAEs COMPATÍVEIS", type="secondary", use_container_width=True):
            sugestoes = sugerir_cnaes_por_texto(atividade_ia)
            st.session_state["ia_cnaes_sugeridos"] = sugestoes
            st.session_state["ia_objeto_sugerido"] = montar_objeto_social_sugerido(atividade_ia, sugestoes)
            if sugestoes:
                set_flash("success", "CNAEs sugeridos com sucesso. Revise antes de gerar o documento.")
            else:
                set_flash("warning", "Nenhuma sugestão encontrada para esse texto. Tente descrever a atividade de outra forma.")
            st.rerun()

        if st.session_state.get("ia_cnaes_sugeridos"):
            st.markdown("**Sugestões encontradas:**")
            for item in st.session_state["ia_cnaes_sugeridos"]:
                st.markdown(f"- **{item['codigo']}** — {item['descricao']}")

    with st.form("form_eventos_dados"):
        dados = {}
        
        if "220" in selecionados:
            st.markdown("**220 - Alteração do Nome Empresarial**")
            dados['220'] = st.text_input("Novo Nome Empresarial*").upper()
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
            
        if "221" in selecionados:
            st.markdown("**221 - Alteração do Título do Estabelecimento (Nome Fantasia)**")
            dados['221'] = st.text_input("Novo Nome Fantasia*").upper()
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)
            
        end_code = next((code for code in ["211", "209", "210"] if code in selecionados), None)
        if end_code:
            st.markdown(f"**{end_code} - Alteração de Endereço**")
            c1, c2 = st.columns([3, 1])
            dados['end_log'] = c1.text_input("Novo Logradouro*").upper()
            dados['end_num'] = c2.text_input("Número*")
            c3, c4, c5 = st.columns([2, 2, 1])
            dados['end_comp'] = c3.text_input("Complemento").upper()
            dados['end_bai'] = c4.text_input("Bairro*").upper()
            dados['end_cid'] = c5.text_input("Cidade*").upper()
            c6, c7 = st.columns([1, 2])
            dados['end_uf'] = c6.text_input("UF*").upper()
            dados['end_cep'] = c7.text_input("CEP*")
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "244" in selecionados:
            st.markdown("**244 - Alteração de Atividades Econômicas**")
            objeto_default = st.session_state.get("ia_objeto_sugerido", "")
            cnae_default = ""
            if st.session_state.get("ia_cnaes_sugeridos"):
                cnae_default = "\n".join([f"{i['codigo']} - {i['descricao']}" for i in st.session_state["ia_cnaes_sugeridos"]])
            dados['244_obj'] = st.text_area("Descreva o novo Objeto Social de forma corrida*", value=objeto_default).upper()
            dados['244_cnae'] = st.text_area("Liste os novos CNAEs (Ex: 4729-6/99 - Comércio Varejista...)*", value=cnae_default).upper()
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "248" in selecionados:
            st.markdown("**248 - Alteração do Tipo de Unidade**")
            dados['248'] = st.text_input("Novo Tipo de Unidade (Ex: Sede, Filial)*").upper()
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "249" in selecionados:
            st.markdown("**249 - Alteração da Forma de Atuação**")
            dados['249'] = st.text_input("Nova Forma de Atuação (Ex: Estabelecimento Fixo, Internet)*").upper()
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "225" in selecionados:
            st.markdown("**225 - Alteração da Natureza Jurídica (Transformação)**")
            dados['225'] = st.text_input("Nome da nova Sociedade Limitada*").upper()
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "capital" in selecionados:
            st.markdown("**Alteração de Capital Social**")
            dados['cap'] = st.text_input("Novo Capital Social (R$)*")
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "052" in selecionados:
            st.markdown("**052 - Reativação (Artigo 60 Lei 8.934/94)** *(Processado automaticamente)*")
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        if "999" in selecionados:
            st.markdown("**999 - Licenciamento de Estabelecimento (Legado)** *(Processado automaticamente)*")
            st.markdown("<hr style='margin: 10px 0;'>", unsafe_allow_html=True)

        # Aviso Condicional de Dados Iguais
        mostrar_confirmacao = st.session_state.get('mostrar_confirmacao_repetidos', False)
        if mostrar_confirmacao:
            st.warning(st.session_state.get('msg_aviso_texto', ''))
            confirmacao_forcar = st.checkbox("⚠️ Confirmo a geração do documento mesmo com os dados idênticos.")
        else:
            confirmacao_forcar = False

        c1, c2 = st.columns(2)
        btn_voltar = c1.form_submit_button("⬅️ VOLTAR À SELEÇÃO", type="secondary")
        btn_gerar = c2.form_submit_button("📄 PROCESSAR E GERAR", type="primary")
        
    if btn_voltar:
        st.session_state.mostrar_confirmacao_repetidos = False
        nav('orquestrador')
        
    if btn_gerar:
        eventos_para_processar = []
        campos_faltantes = []
        repetidos = []
        
        if "220" in selecionados:
            v = dados.get('220', '')
            if not v: campos_faltantes.append("220 - Nome Empresarial")
            else:
                eventos_para_processar.append({"tipo": "220", "dados": {"novo_nome": v}})
                if v == (emp_row[1] or ""): repetidos.append("Nome Empresarial")
                
        if "221" in selecionados:
            v = dados.get('221', '')
            if not v: campos_faltantes.append("221 - Nome Fantasia")
            else:
                eventos_para_processar.append({"tipo": "221", "dados": {"nome_fantasia": v}})
                if v == (emp_row[2] or ""): repetidos.append("Nome Fantasia")
                
        if end_code:
            if not dados['end_log'] or not dados['end_num'] or not dados['end_bai'] or not dados['end_cid'] or not dados['end_uf'] or not dados['end_cep']:
                campos_faltantes.append("Endereço Completo (Todos os campos com * são obrigatórios)")
            else:
                eventos_para_processar.append({"tipo": end_code, "dados": {
                    "logradouro": dados['end_log'], "numero": dados['end_num'], 
                    "complemento": dados['end_comp'], "bairro": dados['end_bai'], 
                    "cidade": dados['end_cid'], "uf": dados['end_uf'], "cep": dados['end_cep']
                }})
                if (dados['end_log'] == (emp_row[4] or "") and
                    str(dados['end_num']) == str(emp_row[5] or "") and
                    dados['end_comp'] == (emp_row[6] or "") and
                    dados['end_bai'] == (emp_row[7] or "") and
                    dados['end_cid'] == (emp_row[8] or "") and
                    dados['end_uf'] == (emp_row[9] or "") and
                    str(dados['end_cep']) == str(emp_row[10] or "")):
                    repetidos.append("Endereço")

        if "244" in selecionados:
            v_obj = dados.get('244_obj', '')
            v_cnae = dados.get('244_cnae', '')
            if not v_obj or not v_cnae: campos_faltantes.append("244 - Objeto Social e CNAEs")
            else: eventos_para_processar.append({"tipo": "244", "dados": {"objeto_descricao": v_obj, "cnaes_lista": v_cnae}})
            
        if "248" in selecionados:
            v = dados.get('248', '')
            if not v: campos_faltantes.append("248 - Tipo de Unidade")
            else: eventos_para_processar.append({"tipo": "248", "dados": {"tipo_unidade": v}})

        if "249" in selecionados:
            v = dados.get('249', '')
            if not v: campos_faltantes.append("249 - Forma de Atuação")
            else: eventos_para_processar.append({"tipo": "249", "dados": {"forma_atuacao": v}})
            
        if "225" in selecionados:
            v = dados.get('225', '')
            if not v: campos_faltantes.append("225 - Nova Sociedade Limitada")
            else: eventos_para_processar.append({"tipo": "225", "dados": {"novo_nome": v}})
            
        if "capital" in selecionados:
            v = dados.get('cap', '')
            if not v: campos_faltantes.append("Alteração de Capital Social")
            else: eventos_para_processar.append({"tipo": "capital", "dados": {"novo_capital": v}})

        if "052" in selecionados: eventos_para_processar.append({"tipo": "052", "dados": {}})
        if "999" in selecionados: eventos_para_processar.append({"tipo": "999", "dados": {}})

        # Avaliação de Erros e Avisos
        if campos_faltantes:
            st.session_state.mostrar_confirmacao_repetidos = False
            st.error(f"❌ **Erro:** Preencha: {', '.join(campos_faltantes)}.")
            return
            
        if repetidos and not confirmacao_forcar:
            st.session_state.mostrar_confirmacao_repetidos = True
            st.session_state.msg_aviso_texto = f"⚠️ **Aviso:** Os novos dados para [{', '.join(repetidos)}] são idênticos aos já registrados. Marque a caixa de confirmação abaixo se deseja prosseguir."
            st.rerun()

        st.session_state.mostrar_confirmacao_repetidos = False

        clausulas_prontas = processar_clausulas_prototype(eventos_para_processar)
        msg_success("✅ Documento processado com sucesso!")
        
        with st.expander("Ver Cláusulas Geradas", expanded=True):
            for c in clausulas_prontas: st.info(c)
        
        with get_db() as conn:
            socios_rows = conn.execute(
                "SELECT nome, nacionalidade, cpf, rg, orgao_emissor, estado_civil, data_nasc, profissao, logradouro, numero, complemento, bairro, cidade, uf, cep FROM tb_socio WHERE id_empresa=?",
                (st.session_state.empresa_id,)
            ).fetchall()

        empresa = {
            'cnpj': emp_row[0] if emp_row else '', 'nome_empresarial': emp_row[1] if emp_row else '', 
            'nome_fantasia': emp_row[2] if emp_row else '', 'nire': emp_row[3] if emp_row else '', 
            'logradouro': emp_row[4] if emp_row else '', 'numero': emp_row[5] if emp_row else '',
            'complemento': emp_row[6] if emp_row else '', 'bairro': emp_row[7] if emp_row else '', 
            'cidade': emp_row[8] if emp_row else '', 'uf': emp_row[9] if emp_row else '', 
            'cep': emp_row[10] if emp_row else ''
        }
        
        socios = [{'nome': r[0], 'nacionalidade': r[1], 'cpf': r[2], 'rg': r[3], 'orgao_emissor': r[4], 'estado_civil': r[5], 'data_nasc': r[6], 'profissao': r[7], 'logradouro': r[8], 'numero': r[9], 'complemento': r[10], 'bairro': r[11], 'cidade': r[12], 'uf': r[13], 'cep': r[14]} for r in socios_rows]

        doc_bytes = gerar_docx_modular(empresa, socios, clausulas_prontas)
        nome_emp = empresa.get('nome_empresarial', 'EMPRESA').replace(' ', '_')
        
        st.download_button(
            label="⬇️ BAIXAR DOCUMENTO .DOCX",
            data=doc_bytes,
            file_name=f"Alteracao_DBE_{nome_emp}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
            use_container_width=True
        )

        pdf_bytes, erro_pdf = converter_docx_para_pdf(doc_bytes)
        if pdf_bytes:
            st.download_button(
                label="⬇️ BAIXAR DOCUMENTO .PDF",
                data=pdf_bytes,
                file_name=f"Alteracao_DBE_{nome_emp}.pdf",
                mime="application/pdf",
                type="primary",
                use_container_width=True
            )
        else:
            st.warning(erro_pdf)

        with get_db() as conn:
            conn.execute("INSERT INTO tb_documento (id_empresa, id_usuario, tipo_documento) VALUES (?,?,?)",
                (st.session_state.empresa_id, st.session_state.user_id, "Alteração Contratual Múltipla"))
            conn.commit()

# ==========================================
# ROTEADOR DE PÁGINAS
# ==========================================
paginas = {
    'login': tela_login, 'registrar': tela_registrar, 'menu': tela_menu,
    'gestao_usuarios': tela_gestao_usuarios, 'cad_empresa': tela_cad_empresa,
    'editar_empresa': tela_editar_empresa, 'buscar_empresa': tela_buscar_empresa,
    'lista_socios': tela_socios, 'cad_socio': tela_cad_socio,
    'editar_socio': tela_editar_socio, 'orquestrador': tela_orquestrador,
    'orquestrador_dados': tela_orquestrador_dados,
    'historico': tela_historico, 'historico_empresa': tela_historico_empresa,
}

paginas.get(st.session_state.pagina, tela_login)()

def aplicar_cores():
    components.html("""
    <script>
    const botoes = window.parent.document.querySelectorAll('button');
    botoes.forEach(b => {
        const t = b.innerText;
        if (t.includes('✏️')) {
            b.style.setProperty('background-color', '#ffc107', 'important');
            b.style.setProperty('color', '#212529', 'important');
            b.style.setProperty('border-color', '#ffc107', 'important');
        } else if (t.includes('🗑️')) {
            b.style.setProperty('background-color', '#dc3545', 'important');
            b.style.setProperty('border-color', '#dc3545', 'important');
            b.style.setProperty('color', 'white', 'important');
        } else if (t.includes('👁️‍🗨️')) {
            b.style.setProperty('background-color', '#17a2b8', 'important');
            b.style.setProperty('border-color', '#17a2b8', 'important');
            b.style.setProperty('color', 'white', 'important');
        }
    });
    </script>
    """, height=0, width=0)

aplicar_cores()
